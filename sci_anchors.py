"""sci_anchors.py -- verbatim excerpts ("anchors") for the Recourse site: build, check and hand over.

An anchor is a short passage of a Supreme Court judgment that the site hands the language model for a whole kind of question (today: delay in a criminal trial, Article 21). The text must be the
judgment's own words, so this tool checks it by code against the Supreme Court's own archive PDF (and, where the index worktree is at hand, against later judgments that quote the same words).
It does NOT decide what the law is, and it never says a case is good law. A person confirms the legal reading: until then the anchor's status stays "machine_checked" and the site words it as
"not yet confirmed by our reviewer".

    python sci_anchors.py verify  [--write] [--cache DIR] [--index DIR]     # fetch each PDF, check every segment word for word, record the result
    python sci_anchors.py export  --out DIR                                  # writes excerpt_anchors.json for the Recourse checkout (only anchors that passed the check)
    python sci_anchors.py sheet   --out DIR                                  # the sheet for the reviewer (Word and Markdown)
    python sci_anchors.py confirm ID --by NAME [--on YYYY-MM-DD]             # records the reviewer's confirmation (status -> reviewer_confirmed)

Data: anchors/anchors.json. The Recourse module excerpt_anchors.py reads the exported copy; its tests cover the rest.

TOPICS. Every anchor belongs to a topic (a kind of question): "trial_delay" (Article 21, ANCHORS switch in Recourse), "circumstantial_evidence" (the landmark case fix of 2026-10-29, LANDMARKS switch). The reviewer sheet is
written per topic (`sheet --topic ID`, or one sheet for every topic when none is named). To add a topic: add its anchors to anchors/anchors.json, an entry to TOPIC_INFO below, and (in Recourse) its question patterns
and prompt wording to excerpt_anchors.py, tested on questions written BEFORE the patterns (see anchors/README.md).

OLD PRINTED REPORTS. The Supreme Court's archive copy of an old judgment is a scan of the printed report: the printer's page numbers sit on lines of their own, footnote marks follow words, a footnote can fall in the
middle of a sentence and a word can be split with a hyphen at the end of a line. An anchor may say what its check should look past, in "verify_options" (default: nothing, which is how the first two anchors were checked):
  ignore_page_numbers    lines that hold only a number of one to four digits are dropped from the archive text;
  ignore_footnote_marks  a bracketed one- or two-digit number (or the OCR's "(l)") stuck to the end of a word, such as "Pradesh.(1)", is also read as absent from the archive text; the text is read both ways, so a
                         segment is never altered and a segment that holds such a bracket of its own must match it exactly;
  join_hyphens           a word split by a hyphen at the end of a line is also read joined ("estab-" "lished"); the segment may match either reading;
  ignore_lines           footnote lines (exact text) that the page break put inside a sentence.
Only furniture is ignored; a word that differs is still a failure. What was ignored is written into the check record.
"""
import argparse
import datetime
import difflib
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "anchors", "anchors.json")
SCHEMA = "recourse-anchors-1"
STATUSES = ("machine_checked", "reviewer_confirmed")
_HEADER = re.compile(r"http://JUDIS\.NIC\.IN\s+SUPREME COURT OF INDIA\s+Page \d+ of \d+", re.I)
METHOD = ("every segment, with quotation marks, dashes and other punctuation, capitals and line breaks ignored, was found word for word in the text of the Supreme Court's archive PDF "
          "(pdftotext); page headers removed")


def norm(text):
    """Words only: lower case, typographic quotes made plain, page headers and every non-letter/digit removed. Two texts that say the same words in the same order give the same string."""
    t = _HEADER.sub(" ", text or "")
    t = t.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    t = re.sub(r"[^a-z0-9]+", " ", t.lower())
    return " ".join(t.split())


_PAGE_NUMBER_LINE = re.compile(r"(?m)^[ \t]*\d{1,4}[ \t]*$")
_FOOTNOTE_MARK = re.compile(r"(?<=[^\s(])\(\s*(?:\d{1,2}|[lI])\s*\)")
_LINE_END_HYPHEN = re.compile(r"(?<=[A-Za-z])-[ \t]*\n[ \t]*(?=[a-z])")
OPTION_KEYS = ("ignore_page_numbers", "ignore_footnote_marks", "join_hyphens", "ignore_lines")


def source_variants(raw, opts=None):
    """The archive text, once for each reading the anchor's options allow, with the furniture the options name removed. Plain strings, not yet normalised. A segment is never altered: with ignore_footnote_marks
    the text is read both with the marks dropped and with them kept, so a segment without a mark matches the first and a segment that holds a bracketed number of its own ("section 313(2)") must match it exactly
    (a wrong number in such a bracket is a failure, not something the check looks past); with join_hyphens each of those readings is also read with split words joined."""
    opts = opts or {}
    text = raw or ""
    drop = {x.strip() for x in opts.get("ignore_lines") or [] if x and x.strip()}
    if drop:
        text = "\n".join(l for l in text.split("\n") if l.strip() not in drop)
    if opts.get("ignore_page_numbers"):
        text = _PAGE_NUMBER_LINE.sub("", text)
    readings = [text]
    if opts.get("ignore_footnote_marks"):
        dropped = _FOOTNOTE_MARK.sub(" ", text)
        if dropped != text:
            readings.insert(0, dropped)
    if opts.get("join_hyphens"):
        for r in list(readings):
            joined = _LINE_END_HYPHEN.sub("", r)
            if joined != r:
                readings.append(joined)
    return readings


def describe_options(opts):
    """A sentence for the check record and the reviewer's sheet; '' when the anchor needs no allowance."""
    opts = opts or {}
    bits = []
    if opts.get("ignore_page_numbers"):
        bits.append("lines holding only a page number")
    if opts.get("ignore_footnote_marks"):
        bits.append("footnote marks such as (1) stuck to the end of a word")
    if opts.get("join_hyphens"):
        bits.append("hyphens at the ends of lines (a split word was also read joined)")
    if opts.get("ignore_lines"):
        bits.append("the footnote line(s) " + "; ".join(repr(x) for x in opts["ignore_lines"]) + " that the page break puts inside a sentence")
    return ("In this old printed report the check looked past: " + ", ".join(bits) + ". No word was ignored.") if bits else ""


def find_segment(segment, source_norm):
    """True when the segment's words occur in the source, in order and whole words."""
    s = norm(segment)
    return bool(s) and (" " + s + " ") in (" " + source_norm + " ")


def nearest(segment, source_norm, width=None):
    """The stretch of the source that looks most like the segment (for the message when a segment is not found)."""
    s = norm(segment).split()
    w = source_norm.split()
    n = width or len(s)
    best, best_r = None, 0.0
    first = s[0] if s else ""
    for i in range(0, max(len(w) - n + 1, 1)):
        if w[i] != first and (not first or difflib.SequenceMatcher(None, w[i], first).ratio() < 0.7):
            continue
        cand = " ".join(w[i:i + n])
        r = difflib.SequenceMatcher(None, cand, " ".join(s)).ratio()
        if r > best_r:
            best, best_r = cand, r
    return best, round(best_r, 3)


def load(path=None):
    doc = json.load(open(path or DATA, encoding="utf-8"))
    if doc.get("schema") != SCHEMA:
        raise ValueError(f"unexpected schema {doc.get('schema')!r}")
    return doc


def save(doc, path=None):
    with open(path or DATA, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, ensure_ascii=False, indent=1)
        fh.write("\n")


def excerpt_text(ex):
    """What the site shows for one excerpt: its segments joined by an ellipsis."""
    return " ... ".join(s.strip() for s in ex["segments"])


def fetch_pdf(url, cache_dir):
    """The PDF at `url`, downloaded once into `cache_dir` (a directory of its own: downloaded files are untrusted data and are only ever handed to pdftotext)."""
    os.makedirs(cache_dir, exist_ok=True)
    path = os.path.join(cache_dir, hashlib.sha1(url.encode()).hexdigest()[:12] + ".pdf")
    if not os.path.exists(path) or os.path.getsize(path) < 1000:
        req = urllib.request.Request(url, headers={"User-Agent": "recourse-sci-anchors/1"})
        with urllib.request.urlopen(req, timeout=60) as r:
            data = r.read()
        if b"%PDF" not in data[:1024]:
            raise RuntimeError(f"{url} did not return a PDF")
        with open(path, "wb") as fh:
            fh.write(data)
    return path


def pdf_text(path):
    out = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, timeout=120)
    if out.returncode != 0:
        raise RuntimeError("pdftotext failed: " + out.stderr[:200])
    return out.stdout


def index_passages(index_dir, doc):
    """The passages of one judgment from the index worktree (claude/sci-index), in order; [] when the index is not at hand."""
    import glob
    import gzip
    for f in sorted(glob.glob(os.path.join(index_dir, "*.meta.jsonl.gz"))):
        m = re.match(r"(\d+)-(\d+)\.", os.path.basename(f))
        if not m or not (int(m.group(1)) <= doc <= int(m.group(2))):
            continue
        out = []
        with gzip.open(f, "rt", encoding="utf-8") as fh:
            for line in fh:
                if f'"doc": {doc},' in line[:40]:
                    d = json.loads(line)
                    if d["doc"] == doc:
                        out.append(d)
        return sorted(out, key=lambda d: d["i"])
    return []


def verify_anchor(a, cache_dir, index_dir=None, fetch=fetch_pdf, text_of=pdf_text):
    """Check one anchor. Returns {'ok', 'failures': [...], 'check': record}. Every segment of every excerpt must be in the archive PDF (looking past only what the anchor's verify_options name); every 'also_in' phrase
    must be in that later judgment."""
    path = fetch(a["link"], cache_dir)
    raw = text_of(path)
    opts = a.get("verify_options") or {}
    srcs = [norm(v) for v in source_variants(raw, opts)]
    failures, n = [], 0
    for ex in a["excerpts"]:
        for seg in ex["segments"]:
            n += 1
            if not any(find_segment(seg, src) for src in srcs):
                near, ratio = max((nearest(seg, src) for src in srcs), key=lambda t: t[1])
                failures.append({"excerpt": ex["label"], "segment": seg[:120], "nearest": near, "ratio": ratio})
    also, also_ok = [], True
    for rec in a.get("also_in", []):
        res = {"doc": rec["doc"], "name": rec.get("name"), "found": None}
        if index_dir:
            ps = index_passages(index_dir, rec["doc"])
            joined = norm(" ".join(p["text"] for p in ps))
            res["found"] = bool(ps) and all(find_segment(ph, joined) for ph in rec.get("must_contain", []))
            also_ok = also_ok and res["found"]
        also.append(res)
    check = {"checked_on": datetime.date.today().isoformat(), "source": "Supreme Court archive PDF", "url": a["link"], "sha256": hashlib.sha256(open(path, "rb").read()).hexdigest(),
             "pages": len(re.findall(r"Page \d+ of \d+", raw)), "segments_checked": n, "all_found": not failures, "method": METHOD + ((" " + describe_options(opts)) if describe_options(opts) else ""), "also_in": also}
    if opts:
        check["options"] = opts
    return {"ok": not failures and also_ok, "failures": failures, "check": check}


def cmd_verify(a):
    doc = load()
    cache = a.cache or os.path.join(tempfile.gettempdir(), "sci_anchors_cache")
    bad = 0
    for an in doc["anchors"]:
        r = verify_anchor(an, cache, a.index)
        print(f"{an['id']}: {'OK' if r['ok'] else 'FAILED'}  ({r['check']['segments_checked']} segments, pdf sha256 {r['check']['sha256'][:12]}, {r['check']['pages']} pages)")
        for f in r["failures"]:
            print(f"   NOT FOUND in {f['excerpt']}: {f['segment']!r}\n      nearest ({f['ratio']}): {f['nearest']!r}")
        for x in r["check"]["also_in"]:
            print(f"   also in doc {x['doc']} ({x['name']}): {x['found']}")
        if r["ok"] and a.write:
            an["check"] = r["check"]
        bad += 0 if r["ok"] else 1
    if a.write:
        save(doc)
        print("verification records written to", DATA)
    return 1 if bad else 0


def cmd_export(a):
    doc = load()
    keep = [x for x in doc["anchors"] if (x.get("check") or {}).get("all_found") and x.get("status") in STATUSES]
    skipped = [x["id"] for x in doc["anchors"] if x not in keep]
    os.makedirs(a.out, exist_ok=True)
    out = {"schema": SCHEMA, "exported_on": datetime.date.today().isoformat(), "anchors": keep}
    path = os.path.join(a.out, "excerpt_anchors.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write("\n")
    print(f"{len(keep)} anchor(s) written to {path}" + (f"; NOT exported (no passing check): {skipped}" if skipped else ""))
    return 0 if keep and not skipped else 1


def cmd_confirm(a):
    doc = load()
    an = next((x for x in doc["anchors"] if x["id"] == a.id), None)
    if not an:
        print("no such anchor:", a.id)
        return 1
    an["status"] = "reviewer_confirmed"
    an["confirmed_by"] = a.by
    an["confirmed_on"] = a.on or datetime.date.today().isoformat()
    save(doc)
    print(f"{a.id}: confirmed by {a.by} on {an['confirmed_on']}. Re-run `export` and copy the file into the Recourse checkout.")
    return 0


TOPIC_INFO = {
    "trial_delay": {
        "file": "ANCHOR_SHEET_trial_delay",
        "kind": "a question about delay in a criminal trial",
        "authority_question": "is it the right authority for a person asking whether a long-pending criminal trial can be closed because of delay?",
        "extra": ("- Is there a better or additional authority you would put in front of a person asking this question? (Our collection also holds, for example, Pankaj Kumar v. State of Maharashtra, Vakil Prasad Singh v. State of Bihar "
                  "and Mohd. Hussain v. State (Govt. of NCT) Delhi; we have not read them for this purpose and have not added them.)"),
    },
    "circumstantial_evidence": {
        "file": "ANCHOR_SHEET_circumstantial_evidence",
        "kind": "a question about whether, and on what proof, a criminal conviction can rest on circumstantial evidence",
        "authority_question": "is it the right authority for a person asking whether a conviction can rest on circumstantial evidence alone and what the prosecution must then prove?",
        "extra": ("- Is there a better or additional authority you would put in front of a person asking this question? (We have not added, for example, later three-judge or larger Bench decisions on the subject; we have not read any "
                  "for this purpose.)\n- The site would say that the five conditions are the Court's own, taken from Hanumant (1952) and called by the Court \"the panchsheel of the proof of a case based on circumstantial evidence\". "
                  "Do you agree that the passage is fairly described that way?"),
    },
}


def _para_words(ex):
    p = str(ex.get("paragraph") or "")
    plural = bool(re.search(r"[-,&]|\bto\b|\band\b", p))
    return ("paragraphs " if plural else "paragraph ") + p + " of the " + (ex.get("report") or "SCC") + " report"


def sheet_markdown(doc, topic=None):
    """The sheet for the reviewer, as Markdown (also the source of the Word file). `topic`: one topic's anchors; None: a sheet for every topic present, one after the other."""
    if topic is None:
        seen = []
        for x in doc["anchors"]:
            t = x.get("topic") or "trial_delay"
            if t not in seen:
                seen.append(t)
        return "\n\n---\n\n".join(sheet_markdown(doc, t) for t in seen)
    info = TOPIC_INFO[topic]
    anchors = [x for x in doc["anchors"] if (x.get("topic") or "trial_delay") == topic]
    L = ["# Verbatim excerpts for the site: please confirm", "",
         f"**What this is.** For {info['kind']} the site would hand the answer-writing model the passages below, and the answer may quote them. Each passage was checked by a program "
         "against the Supreme Court's own archive PDF: the words are exactly the judgment's words (quotation marks, dashes and capitals aside). A program cannot tell whether the passage is the right "
         "authority, whether the law has since moved, or whether our one-line description is fair. That is the part we ask you to confirm.", "",
         "**Until you confirm,** the site words each as *verbatim excerpt, not yet confirmed by our reviewer*: the words are put in brackets after the case name in the answer, and the passages are listed in a box under the answer "
         "(and, in the lawyers' module, in the table of authorities) with the link to the court's own copy. **After you confirm,** that wording is dropped for the anchor you confirmed.", ""]
    for i, an in enumerate(anchors, 1):
        L += [f"## {i}. {an['case_name']} ({an['year']})", "",
              f"- Court and bench we would show: Supreme Court of India, {an.get('bench_note') or 'bench not stated'}" + (f"  *(source: {an['bench_source']})*" if an.get("bench_source") else ""),
              f"- Citation we would show: {an.get('citation') or 'none (not shown)'}" + (f"  *(source of the citation: {an['citation_source']})*" if an.get("citation") and an.get("citation_source") else ""),
              f"- Date of judgment: {an.get('date') or 'not shown'}",
              f"- Official copy: {an['link']}", f"- Status now: {an['status'].replace('_', ' ')}"]
        if describe_options(an.get("verify_options")):
            L.append(f"- How the words were checked: {describe_options(an.get('verify_options'))}")
        if an.get("note_for_reviewer"):
            L.append(f"- Note: {an['note_for_reviewer']}")
        L.append("")
        for j, ex in enumerate(an["excerpts"], 1):
            L += [f"**Excerpt {i}.{j}: {ex['label']}**" + (f" ({_para_words(ex)}: {ex['paragraph_source']})" if ex.get("paragraph") else ""), "", "> " + excerpt_text(ex).replace("\n", " "), ""]
            if ex.get("note"):
                L += [f"*{ex['note']}*", ""]
        L += ["Please answer (tick or write):", "",
              "1. Are these exactly the judgment's words?   Yes / No: ____________",
              f"2. Is this a fair statement of the point the judgment is relied on for, and {info['authority_question']}   Yes / No: ____________",
              "3. Has the Supreme Court, or a statute, since changed or limited this? (If you know of one, name it.) ____________",
              "4. Anything to add or cut? ____________", ""]
    L += ["## For the whole sheet", "", info["extra"], "- Signed / date: ____________", ""]
    return "\n".join(L)


def _write_sheet(doc, topic, out_dir):
    md = sheet_markdown(doc, topic)
    base = TOPIC_INFO[topic]["file"]
    mp = os.path.join(out_dir, base + ".md")
    with open(mp, "w", encoding="utf-8") as fh:
        fh.write(md)
    print("wrote", mp)
    try:
        import docx
        d = docx.Document()
        for line in md.split("\n"):
            if line.startswith("# "):
                d.add_heading(line[2:], 0)
            elif line.startswith("## "):
                d.add_heading(line[3:], 1)
            elif line.startswith("> "):
                p = d.add_paragraph(line[2:])
                p.paragraph_format.left_indent = docx.shared.Cm(1)
            elif line.startswith("- "):
                d.add_paragraph(re.sub(r"\*+", "", line[2:]), style="List Bullet")
            elif re.match(r"\d+\. ", line):
                d.add_paragraph(re.sub(r"\*+", "", line))
            elif line.strip():
                d.add_paragraph(re.sub(r"\*+", "", line))
        dp = os.path.join(out_dir, base + ".docx")
        d.save(dp)
        print("wrote", dp)
    except ImportError:
        print("python-docx is not installed: only the Markdown sheet was written")


def cmd_sheet(a):
    doc = load()
    os.makedirs(a.out, exist_ok=True)
    topics = [a.topic] if getattr(a, "topic", None) else []
    if not topics:
        for x in doc["anchors"]:
            t = x.get("topic") or "trial_delay"
            if t not in topics:
                topics.append(t)
    for t in topics:
        if t not in TOPIC_INFO:
            print("no such topic:", t)
            return 1
        _write_sheet(doc, t, a.out)
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    v = sub.add_parser("verify")
    v.add_argument("--write", action="store_true")
    v.add_argument("--cache")
    v.add_argument("--index", help="the claude/sci-index worktree's index/ folder, to check phrases quoted in later judgments")
    e = sub.add_parser("export")
    e.add_argument("--out", required=True)
    s = sub.add_parser("sheet")
    s.add_argument("--out", required=True)
    s.add_argument("--topic", help="one topic only (trial_delay, circumstantial_evidence); default: a sheet for every topic present")
    c = sub.add_parser("confirm")
    c.add_argument("id")
    c.add_argument("--by", required=True)
    c.add_argument("--on")
    a = ap.parse_args(argv)
    return {"verify": cmd_verify, "export": cmd_export, "sheet": cmd_sheet, "confirm": cmd_confirm}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main())
