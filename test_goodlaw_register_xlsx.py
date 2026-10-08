"""Tests for the register's Excel workbooks and the weekly step (goodlaw/register_xlsx.py, register_weekly.py, register_review.sheet_records).
Run: python test_goodlaw_register_xlsx.py   (needs openpyxl). Writes only to a temporary folder."""
import copy
import hashlib
import json
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from goodlaw import register as R
from goodlaw import register_review as RV
from goodlaw import register_weekly as W
from goodlaw import register_xlsx as X

FAILS = []


def check(name, cond, detail=""):
    print(("[PASS] " if cond else "[FAIL] ") + name + ("" if cond else f"  [{detail}]"))
    if not cond:
        FAILS.append(name)


def digest(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


reg = R.Register.load()
TODAY = "2026-10-08"
with tempfile.TemporaryDirectory() as tmp:
    full = X.export(os.path.join(tmp, "full.xlsx"), "full", TODAY)
    wb = openpyxl.load_workbook(full)
    names = wb.sheetnames
    for need in ("Read me", "Summary", "Findings", "Register", "Post-2015 list (59)", "Post-2015 siblings (18)", "Post-2015 other hits (60)", "Sweep proposals", "Checker flags", "Review sheet",
                 "Review backlog", "Chains", "Known cases", "Library cards (76)", "Index gaps", "Where in repo"):
        check(f"the full workbook has the sheet {need!r}", need in names, names)
    check("every sheet name fits Excel's 31 characters", all(len(n) <= 31 for n in names))
    ws = wb["Register"]
    ids = [ws.cell(i, 1).value for i in range(2, ws.max_row + 1)]
    check("the Register sheet has every row of the register, once, in order", ids == [r["id"] for r in reg.rows], (len(ids), len(reg.rows)))
    hdr = [c.value for c in ws[1]]
    check("the Summary formulas point at the right Register columns", hdr[1] == "List" and hdr[3] == "Trust" and hdr[4] == "Treatment" and hdr[13] == "Overruling year" and hdr[27] == "Recognised how" and hdr[28].startswith("Chain"), hdr)
    s = wb["Summary"]
    rows = {s.cell(i, 1).value: (s.cell(i, 2).value, s.cell(i, 3).value) for i in range(5, s.max_row + 1) if s.cell(i, 1).value}
    check("the Summary counts the whole register", rows["All rows"][0] == len(reg.rows) and rows["All rows"][1].startswith("=COUNTA(Register!"))
    check("the Summary counts the 2016-or-later overruling judgments (list 2 and its siblings)", rows["Rows whose overruling judgment is dated 2016 or later"][0] == sum(1 for r in reg.rows if r["source"] in ("list2", "list2_sibling")))
    check("every Summary count has a live formula beside it", all(v[1] and v[1].startswith("=") for k, v in rows.items() if v[0] is not None))
    p = wb["Post-2015 list (59)"]
    check("the post-2015 sheet has the 59 entries with their verdicts", sum(1 for i in range(4, p.max_row + 1) if p.cell(i, 10).value) == 59)
    check("the post-2015 sheet carries the verdict meaning", {p.cell(i, 11).value for i in range(4, p.max_row + 1) if p.cell(i, 11).value} >= {"Confirmed as stated", "Cannot be checked from our collection"})
    f = wb["Findings"]
    nf = len(json.load(open(os.path.join(os.path.dirname(R.DEFAULT_PATH), "findings.json")))["findings"])
    check("the Findings sheet carries every finding in findings.json", f.max_row - 3 == nf, (f.max_row, nf))
    wk = wb["Review sheet"]
    bl = wb["Review backlog"]
    marked = sum(1 for i in range(2, bl.max_row + 1) if bl.cell(i, 19).value == "yes")
    check("the backlog marks exactly this week's items", marked == wk.max_row - 1, (marked, wk.max_row))
    check("the backlog is at least as big as the week", bl.max_row >= wk.max_row)
    check("the review sheet has decision and treatment drop-downs", len(wk.data_validations.dataValidation) == 2)
    check("no cell begins with a formula except the Summary's counts", all(not (isinstance(c.value, str) and c.value.startswith("=")) for n in names if n != "Summary" for row in wb[n].iter_rows(max_row=60) for c in row))

    # weekly workbook and the round trip of a filled sheet
    wkx = X.export(os.path.join(tmp, "weekly.xlsx"), "weekly", TODAY)
    wb2 = openpyxl.load_workbook(wkx)
    check("the weekly workbook is a read-me and the review sheet", wb2.sheetnames == ["Read me", "Review sheet"], wb2.sheetnames)
    sh = wb2["Review sheet"]
    head = [c.value for c in sh[1]]
    ci = head.index("decision")
    target = next(i for i in range(2, sh.max_row + 1) if sh.cell(i, head.index("kind") + 1).value == "candidate")
    sh.cell(target, ci + 1, "keep")
    sh.cell(target, head.index("note") + 1, "looked at it")
    filled = os.path.join(tmp, "filled.xlsx")
    wb2.save(filled)
    recs = RV.sheet_records(filled)
    got = [r for r in recs if r["decision"]]
    check("a filled Excel sheet is read back with its decision and note", len(got) == 1 and got[0]["decision"] == "keep" and got[0]["note"] == "looked at it", got)
    csvp = os.path.join(tmp, "q.csv")
    RV.write_queue(RV.build_queue(reg, RV.load_flags(), RV.load_proposals(), TODAY, 20, 20, 60), csvp)
    check("the Excel sheet and the CSV sheet list the same items", [r["item_id"] for r in RV.sheet_records(csvp)] == [r["item_id"] for r in recs])
    rows2, flags2, props2 = copy.deepcopy(reg.rows), copy.deepcopy(RV.load_flags()), copy.deepcopy(RV.load_proposals())
    res = RV.apply_decisions(rows2, flags2, props2, filled, "Mr. Advocate", TODAY)
    check("recording the filled Excel sheet applies the one decision", res["applied"] == [(got[0]["item_id"], "keep")] and not res["errors"], res)
    row = next(r for r in rows2 if r["id"] == got[0]["item_id"])
    check("the kept row is due again in seven days and keeps its trust", row["review"]["next_due"] == "2026-10-15" and row["trust"] == "candidate")
    check("recording a sheet does not touch the on-disk register", R.Register.load().rows == reg.rows)

    # the weekly step
    before = {p: digest(p) for p in (R.DEFAULT_PATH, RV.FLAGS_PATH, RV.PROPOSALS_PATH)}
    out = W.run(os.path.join(tmp, "register_review"), TODAY)
    d = os.path.join(tmp, "register_review", TODAY)
    check("the weekly step writes the CSV, the Excel sheet and the summary into DATE/", sorted(os.listdir(d)) == [f"Review_Sheet_{TODAY}.xlsx", "SUMMARY.md", "review_queue.csv"], os.listdir(d))
    check("the weekly step reports 60 items for the week and the backlog size", out["week"] == 60 and out["backlog"] >= 60 and out["weeks"] >= 1, out)
    check("the weekly step changes no row, flag or proposal on disk", before == {p: digest(p) for p in before})
    first_csv, first_md = open(os.path.join(d, "review_queue.csv")).read(), open(os.path.join(d, "SUMMARY.md")).read()
    out2 = W.run(os.path.join(tmp, "register_review"), TODAY)
    check("running the weekly step again on the same day writes the same sheet and summary", open(os.path.join(d, "review_queue.csv")).read() == first_csv and open(os.path.join(d, "SUMMARY.md")).read() == first_md and out2["week"] == out["week"])
    log = open(os.path.join(tmp, "register_review", "RUN_LOG.md"), encoding="utf-8").read()
    check("every run adds one line to the run log, with the time and the counts", log.count("\n- ") == 2 and "UTC: sheet 2026-10-08: 60 items" in log, log)
    check("the run log sits beside the dated folders, not in them", "RUN_LOG.md" not in os.listdir(d))
    check("the summary says it changes nothing", "changes nothing in the register" in open(os.path.join(d, "SUMMARY.md")).read())

print()
print("FAILED: " + ", ".join(FAILS) if FAILS else "ALL PASSED")
sys.exit(1 if FAILS else 0)
