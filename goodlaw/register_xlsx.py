"""Excel workbooks of the Overruled Register, for a person to read and fill in.

    python sci_goodlaw.py register-export-xlsx --out FILE.xlsx                 # the whole register, every finding, the review backlog
    python sci_goodlaw.py register-export-xlsx --out FILE.xlsx --mode weekly   # the week's review sheet only (what the weekly routine writes)

A filled 'Review sheet' (this workbook or the CSV) is recorded with `register-review --sheet FILE`. Only the decision, new_treatment, scope, overruling_case, overruling_year,
overruled_cites, overruled_year, note, reviewer and review_date columns are read. Nothing in this module changes the register.
Source label on everything: the register's own words; rows other than the verified list-1 rows are not confirmed by a person.
"""
import csv
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import register as R
from . import register_review as RV

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="E7EEF7")
LIST_NAME = {"list1": "List 1 (to ~2015)", "list2": "List 2 (after 2015)", "list2_sibling": "List 2 sibling", "checker_flag": "Checker flag", "sweep": "Sweep"}
VERDICT = {"C": "Confirmed as stated", "CN": "Confirmed, but needs a caveat or is incomplete", "P": "Substance right, label stronger than the court's words",
           "W": "Wrong in a material detail", "U": "Cannot be checked from our collection"}
EVIDENCE = {"B": "the judgment's own sentence in its body", "H": "the law report's headnote only", "H+B": "headnote and body"}
DECISIONS = "keep,confirm,change,reject,supply,add_row,keep_open,clear,confirm_overruled,confirm_partial,confirm_not_good_law,confirm_per_incuriam,confirm_not_binding"


def _read_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _sheet(wb, title, headers, rows, widths=None, wrap=(), note=None, freeze=True, filter_=True):
    """A sheet with a header row (row 1, or row 3 under a one-line note), Arial, filters and frozen headers. Returns (ws, first data row)."""
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=FONT, size=10, italic=True)
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(headers), 8))
        ws.row_dimensions[1].height = 44
        r0 = 3
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=headers[j - 1] in wrap, vertical="top")
    for j, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(h, 16)
    if freeze:
        ws.freeze_panes = ws.cell(r0 + 1, 2)
    if filter_ and rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws, r0 + 1


def _cites(side):
    return "; ".join(side.get("cites") or [])


def register_rows(reg):
    out = []
    for r in reg.rows:
        o, b = r["overruled"], r["overruling"]
        ev, ver = r.get("evidence") or {}, r.get("verification") or {}
        chain = reg.chain(r)
        out.append([r["id"], LIST_NAME.get(r["source"], r["source"]), r["source"], r["trust"], r["treatment"], o["name"], "; ".join(o.get("alt_names") or []), o.get("court_detail") or o.get("court"),
                    _cites(o), o.get("year"), b["name"], b.get("court_detail") or b.get("court"), _cites(b), b.get("year"), b.get("date"), b.get("bench"), b.get("doc_key"),
                    r.get("court_words"), r.get("scope"), "yes" if r.get("same_parties") else "", ", ".join(r.get("flags") or []), ev.get("kind"), ev.get("basis"), ev.get("checked_by"),
                    ev.get("list_verdict"), ver.get("by"), ver.get("on"), reg.match_class(r), "; ".join(f"{c['id']}: {reg.phrase(c)}" for c in chain), reg.clause(r) or "(no wording: not confirmed enough)",
                    r.get("note"), (r.get("history") or [{}])[-1].get("event")])
    return out


REG_HEAD = ["ID", "List", "Source code", "Trust", "Treatment", "Overruled case", "Other names", "Overruled court", "Overruled citations", "Overruled year", "Overruling case",
            "Overruling court", "Overruling citations", "Overruling year", "Overruling date", "Bench", "Overruling judgment key", "Court's own words", "Scope", "Same parties (appeal)",
            "Flags", "Evidence kind", "Evidence basis", "Checked by", "List verdict", "Verified by", "Verified on", "Recognised how", "Chain (overruling case later overruled)",
            "Wording the site would use", "Note", "Last event"]
REG_WIDTH = {"ID": 10, "List": 18, "Source code": 12, "Trust": 12, "Treatment": 16, "Overruled case": 44, "Other names": 24, "Overruled citations": 32, "Overruling case": 44, "Overruling citations": 30,
             "Court's own words": 50, "Scope": 36, "Flags": 24, "Chain (overruling case later overruled)": 48, "Wording the site would use": 70, "Note": 50, "Last event": 30, "Checked by": 28}
REG_WRAP = ("Overruled case", "Overruling case", "Court's own words", "Scope", "Chain (overruling case later overruled)", "Wording the site would use", "Note", "Overruled citations", "Overruling citations")


def readme_sheet(wb, reg, mode, today):
    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("THE OVERRULED REGISTER: workbook for checking", True),
        (f"Built {today} from the repository (branch claude/sharp-euler-unyfv0). Source label on everything: the register's own words; only the 1,688 list-1 rows are recorded as verified (on the owner's statement that Mr. Advocate verified that list on 2026-10-07; Claude did not see the verification). Nothing else is confirmed by a person. The register only subtracts trust or warns; it never says a case is good law.", False),
        ("", False),
    ]
    if mode == "full":
        lines += [
            ("HOW TO CHECK", True),
            ("1. Start with 'Summary' for the counts, then 'Findings' (every finding made so far, with where the evidence is).", False),
            ("2. 'Register' is every row (1,765) with its trust, treatment, both cases, the court's own words and the wording the site would use. Filter on 'List' = 'List 2 (after 2015)' for the post-2015 rows.", False),
            ("3. 'Review sheet' holds this week's 60 items and 'Review backlog' all items. Fill the yellow columns (decision, new_treatment, scope, overruling_case, overruling_year, overruled_cites, overruled_year, note) and send the file back; Claude records it with register-review.", False),
            ("4. Decisions: confirm / change / reject / keep for rows; confirm_overruled / confirm_partial / confirm_not_good_law / confirm_per_incuriam / confirm_not_binding / clear / keep_open for the checker's flags; add_row / reject / keep for sweep proposals; supply for a row that needs a citation or year. A verified row is never rejected here (use change). Nothing is deleted. Anything left blank or 'keep' returns in 7 days.", False),
            ("", False),
            ("DOES IT CONTAIN JUDGMENTS FROM AFTER 2015?", True),
            ("Yes. List 1 (the 1,688-entry list) stops around 2015: 2 entries from 2015, none later. The post-2015 material is: List 2 (59 entries whose OVERRULING judgment is from 2016 on; the overruled cases are mostly older), 18 siblings found in the same judgments, the 60 other post-2015 hits not on the list, and the 329 sweep proposals (judgments of 2000 on, 232 open). Sheets: 'Post-2015 list (59)', 'Post-2015 siblings (18)', 'Post-2015 other hits (60)', 'Sweep proposals', and the checker's two flags (Mohan Lal 2018, Mhetre 2011).", False),
            ("", False),
            ("WHAT THE REGISTER IDs MEAN (column 'ID' / 'Register ID')", True),
            ("L1-0006 = row 6 of the 1,688-entry list (list 1, to about 2015). L2-004 = entry 4 of the 59-entry post-2015 list (list 2). D2-NAME (for example D2-KADER, D2-PRABHULAL-2017) = a decision that the same post-2015 judgments overruled but that list 2 leaves out: found by Claude while checking list 2, so 'derived from list 2'; the text after D2- is a short name of the overruled case, with a year where the same case was overruled more than once. (The D2 label is Claude's own naming and was not written down before; there are 18 such rows.) Other prefixes: F- = a checker warning (flag), P- = a sweep proposal, RC- and RS- = rows that will be created only when a reviewer accepts a flag or a proposal (none yet).", False),
            ("", False),
            ("WHERE ARE THE CHECKER'S LABELS (NO_ADVERSE_FOUND, INSUFFICIENT ...)?", True),
            ("On the sheet 'Library cards (76)', column 'Checker status'. They exist only for the 76 cases Recourse can show, not for the register's 1,765 rows and not for the judgments in the search index. ADVERSE_SIGNAL (2): a later bench appears to overrule or reject the case (Mohan Lal, Mhetre; also on 'Checker flags'). QUESTIONED (0 now): later judgments doubt or limit it. NO_ADVERSE_FOUND (46): at least three later judgments we hold cite it and nothing adverse turned up; this is silence in what we hold, NOT a finding that the law is good. INSUFFICIENT (18): fewer than three later judgments we hold cite it, so too little to say. NOT_CHECKABLE_HIGH_COURT (10): a High Court decision; the collection is Supreme Court only. The two checks are separate: the register lists decisions known to be overruled; the checker looks for signs in later judgments.", False),
            ("", False),
            ("SHEETS", True),
        ]
        for name, what in [("Summary", "counts (formulas on the Register sheet)"), ("Findings", "every finding so far, with source and status"), ("Register", "all 1,765 rows"),
                           ("Post-2015 list (59)", "list 2 with Claude's verdict on each entry"), ("Post-2015 siblings (18)", "further decisions overruled in the same judgments"),
                           ("Post-2015 other hits (60)", "machine hits not on list 2, unread"), ("Sweep proposals", "329 proposals from the sweep (2000 on)"), ("Checker flags", "the 2 open warnings"),
                           ("Review sheet", "this week's items, to fill in"), ("Review backlog", "every item that is or will be due"), ("Chains", "rows whose overruling case was itself later overruled"),
                           ("Listing issues", "list 1 rows the list itself contradicts"), ("Needs citation or year", "list 1 rows that cannot be recognised"),
                           ("List1 parse (1688)", "how each list-1 row was classified (court, appeal, in the collection)"), ("List1 checker test (308)", "the checker's test on list-1 pairs"),
                           ("Library cards (76)", "the checker's card for each case the site can show"), ("Known cases", "cases whose later history is known, and whether the register holds them"),
                           ("Currency pass", "the register run over the first search test"), ("Index gaps", "cases the search cannot show, and why"), ("Where in repo", "file locations")]:
            lines.append((f"   {name}: {what}", False))
    else:
        lines += [("THIS WEEK'S REVIEW SHEET", True),
                  ("Fill the yellow columns on 'Review sheet' and send the file back. Decisions: confirm / change / reject / keep for rows; confirm_overruled / confirm_partial / confirm_not_good_law / confirm_per_incuriam / confirm_not_binding / clear / keep_open for flags; add_row / reject / keep for sweep proposals; supply for a row needing a citation or year. A verified row is never rejected here (use change). Anything blank or 'keep' returns in 7 days.", False)]
    lines += [("", False), ("TRUST LEVELS", True),
              ("verified: list 1, on the owner's statement. reviewer_confirmed: a person confirmed it at a weekly review. candidate: reported, not yet confirmed (worded 'reported as ... not yet confirmed by our reviewer'). unchecked: no wording at all. rejected: kept, never worded.", False)]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=FONT, size=12 if i == 1 else 10, bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150
    return ws


def summary_sheet(wb, reg):
    """Counts twice: column B is counted by the export (always visible); column C is a live Excel formula on the Register sheet (Excel computes it when the file is opened).
    The cloud session that builds this file cannot run a spreadsheet engine, so the formulas are not pre-computed: if B and C ever differ in Excel, tell Claude."""
    n = len(reg.rows) + 1
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Counts"
    ws["A1"].font = Font(name=FONT, size=12, bold=True)
    ws["A2"] = "Column B was counted when the file was built. Column C is a live formula on the Register sheet that Excel computes when you open the file; the two should agree."
    ws["A2"].font = Font(name=FONT, size=10, italic=True)
    rows = [("By list and trust", None, None)]
    for s, t in sorted({(r["source"], r["trust"]) for r in reg.rows}):
        k = sum(1 for r in reg.rows if r["source"] == s and r["trust"] == t)
        rows.append((f"{LIST_NAME[s]} / {t}", k, f'=COUNTIFS(Register!$B$2:$B${n},"{LIST_NAME[s]}",Register!$D$2:$D${n},"{t}")'))
    rows += [("All rows", len(reg.rows), f"=COUNTA(Register!$A$2:$A${n})"), ("", None, None), ("By treatment", None, None)]
    for t in sorted({r["treatment"] for r in reg.rows}):
        rows.append((t, sum(1 for r in reg.rows if r["treatment"] == t), f'=COUNTIF(Register!$E$2:$E${n},"{t}")'))
    rows += [("", None, None), ("How a row can be recognised in an answer", None, None)]
    mc = {}
    for r in reg.rows:
        k = reg.match_class(r)
        mc[k] = mc.get(k, 0) + 1
    for k, label in (("name_year", "strongly, by its own names and a year"), ("citation", "strongly, by a reported citation only"), ("weak", "weakly (names only: never worded as a fact)"), ("none", "not at all")):
        rows.append((label, mc.get(k, 0), f'=COUNTIF(Register!$AB$2:$AB${n},"{k}")'))
    rows += [("", None, None), ("Chains (the overruling case was itself later overruled)", sum(1 for r in reg.rows if reg.chain(r)), f'=COUNTIF(Register!$AC$2:$AC${n},"?*")'),
             ("Rows whose overruling judgment is dated 2016 or later", sum(1 for r in reg.rows if (r["overruling"].get("year") or 0) >= 2016), f'=COUNTIF(Register!$N$2:$N${n},">=2016")')]
    ws.cell(4, 1, "Item")
    ws.cell(4, 2, "Count at export")
    ws.cell(4, 3, "Live formula (Excel)")
    for i, (a_, b_, c_) in enumerate(rows, 5):
        ws.cell(i, 1, a_)
        ws.cell(i, 2, b_)
        ws.cell(i, 3, c_)
    headings = {"Item", "By list and trust", "By treatment", "How a row can be recognised in an answer", "All rows"}
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.font = Font(name=FONT, size=10, bold=(c.row == 4 or row[0].value in headings), color="000000")
    for j in (1, 2, 3):
        ws.cell(4, j).fill = HEAD_FILL
        ws.cell(4, j).font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 22
    return ws


def _queue(reg, today, big):
    flags, props = RV.load_flags(), RV.load_proposals()
    if big:
        return RV.build_queue(reg, flags, props, today, proposal_limit=10 ** 6, match_limit=10 ** 6, limit=None)
    return RV.build_queue(reg, flags, props, today, 20, 20, 60)


def review_sheet(wb, items, title="Review sheet", week=None):
    ws, first = _sheet(wb, title, RV.COLUMNS + (["in this week's sheet"] if week is not None else []), [[it.get(c, "") for c in RV.COLUMNS] + ([("yes" if it["item_id"] in week else "") ] if week is not None else []) for it in items],
                       widths={"item_id": 24, "kind": 11, "priority": 8, "case": 40, "later_case": 40, "evidence": 70, "ask": 46, "decision": 18, "new_treatment": 18, "scope": 28, "overruling_case": 30,
                               "note": 30, "next_review_if_kept": 14}, wrap=("case", "later_case", "evidence", "ask", "note", "scope"))
    cols = {h: i + 1 for i, h in enumerate(RV.COLUMNS)}
    last = first + len(items) - 1
    if items:
        for name in ("decision", "new_treatment", "scope", "overruling_case", "overruling_year", "overruled_cites", "overruled_year", "note", "reviewer", "review_date"):
            for r in range(first, last + 1):
                ws.cell(r, cols[name]).fill = INPUT_FILL
        dv = DataValidation(type="list", formula1=f'"{DECISIONS}"', allow_blank=True)
        dv2 = DataValidation(type="list", formula1=f'"{",".join(R.TREATMENTS)}"', allow_blank=True)
        ws.add_data_validation(dv)
        ws.add_data_validation(dv2)
        dv.add(f"{get_column_letter(cols['decision'])}{first}:{get_column_letter(cols['decision'])}{last}")
        dv2.add(f"{get_column_letter(cols['new_treatment'])}{first}:{get_column_letter(cols['new_treatment'])}{last}")
    return ws


def post2015_sheets(wb, reg, root):
    res = {r["no"]: r for r in _read_csv(os.path.join(root, "external_lists/2026-10-07_overruled_after_2015_pdf/results.csv"))}
    rows = []
    for r in reg.rows:
        if r["source"] != "list2":
            continue
        no = str(int(r["id"].split("-")[1]))
        x = res.get(no, {})
        b = r["overruling"]
        rows.append([int(no), r["id"], x.get("earlier_case"), x.get("earlier_year_in_list"), x.get("later_case"), x.get("later_year_in_list"), b.get("date"), b.get("bench"), x.get("later_judgment_in_collection"),
                     x.get("verdict"), VERDICT.get(x.get("verdict"), ""), EVIDENCE.get(x.get("evidence"), x.get("evidence")), x.get("court_words"), x.get("note"), r["trust"], r["treatment"], r.get("scope"),
                     "; ".join(f"{c['id']}: {reg.phrase(c)}" for c in reg.chain(r)), reg.clause(r) or "(no wording)"])
    _sheet(wb, "Post-2015 list (59)", ["List no.", "Register ID", "Earlier case (as listed)", "Year in list", "Later case (as listed)", "Year in list (later)", "Decision date", "Bench", "Later judgment key", "Verdict", "Verdict meaning",
                                       "Evidence", "Court's words", "Claude's note", "Trust", "Treatment", "Scope", "Chain", "Wording the site would use"], rows,
           widths={"Earlier case (as listed)": 40, "Later case (as listed)": 40, "Verdict meaning": 36, "Court's words": 50, "Claude's note": 50, "Scope": 36, "Chain": 40, "Wording the site would use": 70},
           wrap=("Earlier case (as listed)", "Later case (as listed)", "Verdict meaning", "Court's words", "Claude's note", "Scope", "Chain", "Wording the site would use", "Evidence"),
           note="List 2: 'Judgments Overruled After 2015' (59 entries). The OVERRULING judgment is from 2016 on; the overruled cases are mostly older. Verdicts are Claude's reading of the judgments in the collection, not a lawyer's: C confirmed, CN confirmed with a caveat, P label stronger than the court's words, W wrong in a detail, U not checkable. None of these rows is verified.")
    sib = [[r["id"], r["treatment"], r["overruled"]["name"], _cites(r["overruled"]), r["overruled"].get("year"), r["overruling"]["name"], r["overruling"].get("date"), r["overruling"].get("bench"),
            r.get("court_words"), r.get("scope"), r["trust"], r.get("note"), reg.clause(r) or ""] for r in reg.rows if r["source"] == "list2_sibling"]
    _sheet(wb, "Post-2015 siblings (18)", ["Register ID", "Treatment", "Overruled case", "Citations", "Year of overruled case", "Overruled by", "Date of overruling judgment", "Bench (judges) of overruling judgment", "Court's words", "Scope", "Trust", "Note", "Wording the site would use"], sib,
           widths={"Overruled case": 44, "Overruled by": 44, "Court's words": 54, "Scope": 36, "Note": 50, "Wording the site would use": 70, "Year of overruled case": 12, "Date of overruling judgment": 16, "Bench (judges) of overruling judgment": 16}, wrap=("Overruled case", "Overruled by", "Court's words", "Scope", "Note", "Wording the site would use"),
           note="Further decisions found overruled in the same post-2015 judgments (for example the six companions of India Cement). IDs start with D2 ('derived from list 2'): list 2 leaves them out. Candidates: Claude's reading, not confirmed by a person.")
    oth = [[x["doc_key"], x["date"], x["bench_size"], x["judgment"], x["first_operative_looking_sentence"], x["sentences_found"], x["my_reading"]] for x in
           _read_csv(os.path.join(root, "external_lists/2026-10-07_overruled_after_2015_pdf/other_candidates_not_on_the_list.csv"))]
    _sheet(wb, "Post-2015 other hits (60)", ["Judgment key", "Date", "Bench", "Judgment", "First overruling-looking sentence", "Sentences found", "Claude's reading"], oth,
           widths={"Judgment": 44, "First overruling-looking sentence": 90, "Claude's reading": 44}, wrap=("Judgment", "First overruling-looking sentence", "Claude's reading"),
           note="60 judgments of three or more judges (2016 on) with an overruling-looking sentence that are NOT on list 2. Machine hits, not individually read; about 30 look like real fresh overrulings. The scan missed Tofan Singh (2020), so 60 is a floor.")


def proposals_sheet(wb):
    props = RV.load_proposals()["proposals"]
    rows = [[p["id"], p["state"], p["kind"], p.get("date"), p.get("bench"), p["title"], p.get("overruled_name"), p.get("overruled_cites"), p["evidence"], "yes" if p.get("already_in_register") else "",
             p.get("doc_key"), p.get("first_seen"), p.get("next_due"), (p.get("history") or [{}])[-1].get("event")] for p in props]
    _sheet(wb, "Sweep proposals", ["ID", "State", "Kind", "Date", "Bench", "Later judgment", "Overruled case named", "Citations printed", "Evidence (tag or sentence)", "Already in register", "Judgment key", "First seen", "Next due", "Last event"],
           rows, widths={"Later judgment": 40, "Overruled case named": 40, "Citations printed": 30, "Evidence (tag or sentence)": 90}, wrap=("Later judgment", "Overruled case named", "Citations printed", "Evidence (tag or sentence)"),
           note="Proposals from the sweep of judgments of three or more judges, 2000 on (widened from 2016 on 2026-10-08). Proposals only: nothing enters the register until a person confirms it. 'tag' = the law report's own case-law list; 'sentence' = the court's own sentence. None has been read by anyone.")


def flags_sheet(wb, reg):
    flags = RV.load_flags()["flags"]
    rows = []
    for f in flags:
        c = f["case"]
        held = [h["row"]["id"] for h in reg.lookup_case(c["name"], c.get("year"), c.get("citation")) if h["strength"] == "strong"]
        for e in f["evidence"] or [{}]:
            rows.append([f["id"], f["state"], c["name"], c.get("citation"), c.get("year"), f["status"], f.get("note"), ", ".join(held), e.get("title"), e.get("date"), e.get("bench"), e.get("cue"), e.get("label"), e.get("quote")])
    _sheet(wb, "Checker flags", ["Flag", "State", "Case", "Citation", "Year", "Checker status", "Note the site would show", "Register rows", "Later judgment", "Date", "Bench", "Cue", "Reading", "Quoted evidence"], rows,
           widths={"Case": 36, "Note the site would show": 60, "Later judgment": 40, "Quoted evidence": 90}, wrap=("Case", "Note the site would show", "Later judgment", "Quoted evidence"),
           note="The checker's open warnings (cases the site can show whose later history looks adverse). Both are already candidate rows in the register. A person confirms or clears them at the weekly review.")


def subset_sheets(wb, reg):
    ch = []
    for r in reg.rows:
        for c in reg.chain(r):
            ch.append([r["id"], LIST_NAME.get(r["source"], r["source"]), r["trust"], r["overruled"]["name"], r["treatment"], r["overruling"]["name"], c["id"], c["treatment"], c["overruling"]["name"], reg.clause(r) or ""])
    _sheet(wb, "Chains", ["Row", "List", "Trust", "Overruled case", "Treatment", "Overruling case", "Later row", "Later treatment", "Overruling case of the later row", "Wording the site would use"], ch,
           widths={"Overruled case": 44, "Overruling case": 44, "Overruling case of the later row": 44, "Wording the site would use": 80}, wrap=("Overruled case", "Overruling case", "Overruling case of the later row", "Wording the site would use"),
           note="Rows whose overruling case was itself later overruled or disapproved. The recorded status of the earlier case cannot be copied forward without a re-read; each is a 'chain' item for the reviewer.")
    li = [[r["id"], r["treatment"], r["overruled"]["name"], _cites(r["overruled"]), r["overruling"]["name"], ", ".join(r["flags"]), r.get("note")] for r in reg.rows
          if r["source"] == "list1" and any(f in r["flags"] for f in ("list_mutual_pair", "list_no_overruling_decision", "year_order_check", "disputed_listing")) or r["treatment"] == "disputed_listing"]
    _sheet(wb, "Listing issues", ["Row", "Treatment", "Overruled case", "Citations", "Overruling case", "Flags", "Note"], li, widths={"Overruled case": 50, "Overruling case": 50, "Flags": 34, "Note": 50},
           wrap=("Overruled case", "Overruling case", "Note"), note="List 1 rows the list itself contradicts (each case listed as overruled by the other, no overruling decision named, or the overruling decision dated before the overruled one), including the 13 recorded as disputed_listing.")
    nm = [[r["id"], r["overruled"]["name"], _cites(r["overruled"]), r["overruled"].get("year"), reg.match_class(r), r["overruling"]["name"]] for r in reg.rows if "needs_citation_or_year" in r["flags"]]
    _sheet(wb, "Needs citation or year", ["Row", "Overruled case", "Citations", "Year", "Recognised how", "Overruling case"], nm, widths={"Overruled case": 56, "Overruling case": 56}, wrap=("Overruled case", "Overruling case"),
           note="List 1 rows with no citation or year: the register cannot recognise them in an answer (weak or not at all) until the reviewer supplies one ('supply' on the weekly sheet).")


def list1_sheets(wb, root):
    d = os.path.join(root, "external_lists/2026-10-07_overruled_cases_pdf")
    parsed = _read_csv(os.path.join(d, "rows_parsed.csv"))
    _sheet(wb, "List1 parse (1688)", list(parsed[0].keys()), [[r[k] for k in parsed[0]] for r in parsed], widths={k: 18 for k in parsed[0]},
           note="How each of the 1,688 list-1 rows was classified (no case text is kept here: the list's own text stays in the PDF in inbox/). overruled_is: whether the overruled decision is a Supreme Court one; same_parties_appeal: an appeal reversing the lower court; the *_in_collection columns are our collection's document keys.")
    pairs = {r["pdf_row"]: r for r in _read_csv(os.path.join(d, "test_pairs.csv"))}
    res = {r["pdf_row"]: r for r in _read_csv(os.path.join(d, "checker_results_tuning_half.csv"))}
    rows = []
    for k, p in pairs.items():
        x = res.get(k, {})
        rows.append([int(k), p["half"], p["overruled_doc_key"], p["overruling_doc_key"], p["overruled_date"], p["overruling_date"], x.get("reading_before_fixes"), x.get("reading_after_fixes"), x.get("adverse_after_fixes"),
                     x.get("cue"), x.get("voice"), x.get("confident_and_court_or_headnote_voice"), x.get("match_how"), x.get("sentence_read")])
    rows.sort(key=lambda r: r[0])
    _sheet(wb, "List1 checker test (308)", ["List row", "Half", "Overruled key", "Overruling key", "Overruled date", "Overruling date", "Reading before fixes", "Reading after fixes", "Adverse after fixes", "Cue", "Voice", "Confident, court/headnote voice", "Matched how", "Sentence read"],
           rows, widths={"Sentence read": 90}, wrap=("Sentence read",),
           note="The 308 usable pairs. Half 'tune' (even rows, 146 pairs) was run: 38 of 146 (26%) read as adverse before the fixes of 2026-10-07 and 42 (29%) after (in-sample). Half 'held' (odd rows, 162 pairs) has NOT been run and must be run once, at a frozen commit. Every pair is a positive: recall only, not false alarms.")


def cards_sheet(wb, reg, root):
    d = _json(os.path.join(root, "cards/currency_cards.json"))
    rows = []
    for c in d["cards"]:
        cs, co, sp, cv = c["case"], c.get("corpus") or {}, c.get("support") or {}, c.get("coverage") or {}
        sr = c.get("second_reader") or {}
        held = [h["row"]["id"] for h in reg.lookup_case(cs["name"], cs.get("year"), cs.get("citation")) if h["strength"] == "strong"]
        rows.append([cs["name"], cs.get("citation"), cs.get("year"), cs.get("court_class"), c["status"], c.get("statute_flag"), c.get("user_facing_note"), co.get("doc_key"), co.get("title"), co.get("date"), co.get("bench_size"),
                     (co.get("match") or {}).get("method"), (co.get("match") or {}).get("confidence"), sp.get("citing_judgments"), sp.get("latest_citing_year"), sp.get("citing_in_last_10_years"),
                     cv.get("share_of_later_judgments_we_hold"), sr.get("label") if isinstance(sr, dict) else sr, ", ".join(held), c.get("internal_note"), c.get("report_file")])
    _sheet(wb, "Library cards (76)", ["Case", "Citation", "Year", "Court", "Checker status", "Statute flag", "Note the site would show", "Judgment key", "Matched judgment", "Date", "Bench", "Match method", "Match confidence",
                                       "Later judgments citing it", "Latest citing year", "Citing in last 10 years", "Share of later judgments we hold", "Second reader", "Register rows", "Internal note", "Report file"], rows,
           widths={"Case": 40, "Note the site would show": 60, "Matched judgment": 40, "Internal note": 60, "Second reader": 40, "Report file": 40}, wrap=("Case", "Note the site would show", "Matched judgment", "Internal note", "Second reader"),
           note=f"The checker's card for each case Recourse can show ({d['generated']}). Pattern-rule reading of Supreme Court texts; not lawyer-validated. 'No adverse found' is silence in what we hold, not a finding about the law. Counts: " + ", ".join(f"{k} {v}" for k, v in d["counts"].items()) + ".")


def known_sheets(wb, reg, root):
    from review_tools import currency_pass as CP
    kc = CP.known_cases()
    view = {v["name"]: v for v in CP.register_view(reg, kc)}
    def held(c):
        v = view[c["name"]]
        if c["expect"] == "not_adverse":
            return "YES: false alarm" if v["in_register"] else "no (as it should be)"
        return "yes" if v["in_register"] else "NO: a gap"
    rows = [[c["name"], c["expect"], c.get("doc_key"), held(c), ", ".join(view[c["name"]]["trust"]), "FALSE ALARM" if view[c["name"]]["false_alarm"] else "", c.get("basis"), c.get("source", "goodlaw/targets/known.json")] for c in kc]
    _sheet(wb, "Known cases", ["Case", "Expected", "Judgment key", "In the register", "Trust", "False alarm", "What the later court did / why it is listed", "Source"], rows,
           widths={"Case": 44, "What the later court did / why it is listed": 90, "Source": 40}, wrap=("Case", "What the later court did / why it is listed", "Source"),
           note="Cases whose later history is known from outside (7 overruled/doubted/overcome and 13 settled, labelled by an AI assistant from general knowledge, plus 3 added 2026-10-08). 'In the register' says whether the register would word them. Not lawyer-checked.")
    pass_path = os.path.join(root, "review_results/2026-10-08_currency_pass_first_pool/currency_pass.json")
    if os.path.exists(pass_path):
        p = _json(pass_path)
        rr = [[v, s.get("questions", 0), s.get("slots", 0), s.get("own_firm", 0), s.get("own_candidate", 0), s.get("names_flagged_case", 0), s.get("stale_slots", 0), s.get("stale_slots_flagged", 0),
               s.get("after_slots_changed", 0), s.get("after_stale_slots", 0)] for v, s in p["versions"].items()]
        ws, first = _sheet(wb, "Currency pass", ["Search version", "Questions", "Top slots", "Own case held (firm)", "Own case held (candidate)", "Names a held case", "Known-stale slots", "of which register flagged",
                                                 "After demotion: slots changed", "After demotion: stale slots"], rows=rr, widths={"Search version": 28},
                           note=f"The register run over the first review pool ({p['passages']} passages), top {p['top_k']} per version. Private; never sent to reviewers. Stale passages are few (0 to 2 per version): too few to rank versions.")
        r0 = first + len(rr) + 2
        ws.cell(r0, 1, "Stale passages found").font = Font(name=FONT, size=10, bold=True)
        for j, h in enumerate(["Question", "Passage", "Version", "Case", "Register action"], 1):
            ws.cell(r0 + 1, j, h).font = Font(name=FONT, size=10, bold=True)
        for i, x in enumerate(p["stale_slots"], r0 + 2):
            for j, v in enumerate([x["qid"], x["n"], x["version"], x["case"], x["register_action"] or "none"], 1):
                ws.cell(i, j, v).font = Font(name=FONT, size=10)


def gaps_sheet(wb, root):
    g = _json(os.path.join(root, "review_data/ruler/topics_v2.json"))["index_gaps"]
    rows = [[x["case"], x["recourse_use"], x["finding"]] for x in g]
    rows.append(["Mohan Lal v State of Punjab (2018) 17 SCC 627", "overruled by Mukesh Singh (2020)", "NOT a gap: in the collection and the index as 'MOHANLAL versus THE STATE OF PUNJAB' (doc 1000250931). An earlier note said it was missing; that was an error."])
    _sheet(wb, "Index gaps", ["Case", "Why it matters", "Finding"], rows, widths={"Case": 60, "Why it matters": 44, "Finding": 100}, wrap=("Case", "Why it matters", "Finding"),
           note="Cases Recourse relies on that the search over the big index cannot show, and why. Checked with spacing-insensitive title matching and neutral citations.")


def findings_sheet(wb):
    f = _json(os.path.join(R.DEFAULT_PATH.rsplit("/", 1)[0], "findings.json"))["findings"]
    _sheet(wb, "Findings", ["No.", "Area", "Finding", "Where the evidence is", "Confidence", "Status", "Who acts"], [[x["no"], x["area"], x["finding"], x["evidence"], x["confidence"], x["status"], x.get("who_acts", "")] for x in f],
           widths={"No.": 6, "Area": 26, "Finding": 110, "Where the evidence is": 46, "Confidence": 24, "Status": 34, "Who acts": 34}, wrap=("Finding", "Where the evidence is", "Confidence", "Status", "Who acts", "Area"),
           note="Every finding made so far on the lists, the register, the checker and the index, with where the evidence is. Source file: goodlaw/data/register/findings.json. Nothing here is a lawyer's finding.")


def where_sheet(wb):
    rows = [["The register (1,765 rows)", "goodlaw/data/register/register.json", "the data; edited only by register-build and register-review"],
            ["Checker flags", "goodlaw/data/register/flags.json", "2 open"], ["Sweep proposals", "goodlaw/data/register/proposals.json", "329 (232 open)"],
            ["Findings log", "goodlaw/data/register/findings.json", "source of the 'Findings' sheet"],
            ["Description, rules, limits", "REGISTER.md", ""], ["Patch for the Recourse owner (not applied)", "REGISTER_SITE_PATCH.md", ""],
            ["Code", "goodlaw/register.py, register_build.py, register_guard.py, register_review.py, register_sweep.py, register_exam.py, register_xlsx.py", ""],
            ["Tests", "test_goodlaw_register.py", ""], ["Trap-question exam (not run)", "register_exam/", "50 questions and key"],
            ["Post-2015 list: analysis", "external_lists/2026-10-07_overruled_after_2015_pdf/ (README.md, results.csv, other_candidates_not_on_the_list.csv)", ""],
            ["List 1: analysis", "external_lists/2026-10-07_overruled_cases_pdf/ (README.md, rows_parsed.csv, test_pairs.csv, checker_results_tuning_half.csv)", ""],
            ["Checker cards", "cards/ (currency_cards.json, REVIEW_SHEET.csv, MATCH_CHECK.csv, reports/)", ""],
            ["Currency pass", "review_tools/currency_pass.py; review_results/2026-10-08_currency_pass_first_pool/", ""],
            ["Weekly sheets", "register_review/YYYY-MM-DD/ on branch claude/sharp-euler-unyfv0", "written by the weekly routine"],
            ["Weekly routine instructions", "REGISTER_ROUTINE_RUNBOOK.md", ""], ["This workbook", "register_review/2026-10-08/Overruled_Register_Full_Review.xlsx", "rebuild: python sci_goodlaw.py register-export-xlsx --out FILE.xlsx"],
            ["Repository", "github.com/Tanbir-Ahm/recourse-sci-pool (private), branch claude/sharp-euler-unyfv0; not merged to main", ""]]
    _sheet(wb, "Where in repo", ["What", "Path", "Note"], rows, widths={"What": 44, "Path": 110, "Note": 50}, wrap=("Path", "Note"), filter_=False)


def export(path, mode="full", today=None, root=ROOT):
    today = R.today_str(today)
    reg = R.Register.load()
    wb = Workbook()
    readme_sheet(wb, reg, mode, today)
    if mode == "weekly":
        review_sheet(wb, _queue(reg, today, big=False))
    else:
        summary_sheet(wb, reg)
        findings_sheet(wb)
        _sheet(wb, "Register", REG_HEAD, register_rows(reg), widths=REG_WIDTH, wrap=REG_WRAP, note=None)
        post2015_sheets(wb, reg, root)
        proposals_sheet(wb)
        flags_sheet(wb, reg)
        week = _queue(reg, today, big=False)
        review_sheet(wb, week, "Review sheet")
        review_sheet(wb, _queue(reg, today, big=True), "Review backlog", week={it["item_id"] for it in week})
        subset_sheets(wb, reg)
        list1_sheets(wb, root)
        cards_sheet(wb, reg, root)
        import sys
        sys.path.insert(0, root)
        known_sheets(wb, reg, root)
        gaps_sheet(wb, root)
        where_sheet(wb)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
