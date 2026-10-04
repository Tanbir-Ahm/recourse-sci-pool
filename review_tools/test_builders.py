"""Tests for the review builders (first AI-reviewer pack, top-up pack, human review sheet, authority list) using the saved pool.
Run: python review_tools/test_builders.py   (needs openpyxl). Builds into a temporary folder; changes nothing in the repo."""
import os
import re
import shutil
import sys
import tempfile
import zipfile

import openpyxl
from openpyxl.utils import range_boundaries

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_authority_list as bal
import build_llm_pack as blp
import build_review_sheet as brs
import build_topup_pack as btp
import review_common as rc

FAILS = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  [{detail}]"))
    if not cond:
        FAILS.append(name)


pool = rc.load_pool()
first = {q["qid"]: rc.rows_of(q, False) for q in pool}
extra = {q["qid"]: rc.rows_of(q, True) for q in pool}
n_first, n_extra = sum(map(len, first.values())), sum(map(len, extra.values()))
VARIANT_WORDS = sorted({v for q in pool for r in q["rows"] for v in r["variants"]}) + ["landmark", "rewrite", "lane a", "variants", "answer key", "verified paragraph"]
REWRITES = [q["rewrite"] for q in pool if q.get("rewrite")]

# ---------------------------------------------------------------- the pool itself
check("12 questions, 175 first-batch passages and 24 top-up passages", len(pool) == 12 and n_first == 175 and n_extra == 24, (len(pool), n_first, n_extra))
check("passage numbers run 1..N without gaps or repeats within each question, across both batches",
      all(sorted(r["n"] for r in q["rows"]) == list(range(1, len(q["rows"]) + 1)) for q in pool))
check("no passage text appears twice within a question (the pool was de-duplicated)", all(len({r["h"] for r in q["rows"]}) == len(q["rows"]) for q in pool))
check("every row has what the builders need", all({"h", "doc", "pn", "page", "title", "text", "also", "variants", "n"} <= set(r) for q in pool for r in q["rows"]))
check("top-up rows carry only the two 'what a method showed' tags", all(set(r["variants"]) <= {"S site shown", "M mine top 6"} and r["variants"] for q in pool for r in rc.rows_of(q, True)))


def removed_by_index_fixes():
    """Judgment ids that an applied index fix removed (index_fixes/*/plan.json: every group's 'remove' list)."""
    import glob
    import json
    root = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "index_fixes")
    gone = set()
    for f in glob.glob(os.path.join(root, "*", "plan.json")):
        for g in json.load(open(f, encoding="utf-8")).get("groups", []):
            gone |= {int(x) for x in g.get("remove", [])}
    return gone


GONE = removed_by_index_fixes()
check("the index fixes on file remove judgments (so the next check is not vacuous)", len(GONE) >= 100, len(GONE))
check("no pool passage points at a judgment that an index fix removed (the pool was once built before a fix)",
      not [(q["qid"], r["n"], r["doc"]) for q in pool for r in q["rows"] if r["doc"] in GONE], [(q["qid"], r["n"], r["doc"]) for q in pool for r in q["rows"] if r["doc"] in GONE])
fixed_rows = [(q["qid"], r["n"]) for q in pool for r in q["rows"] if r.get("was_doc")]
check("the 4 passages repointed on 2026-10-04 are recorded (judgment 18070 -> 20481) and show the right case name",
      fixed_rows == [("Q03", 4), ("Q04", 1), ("Q04", 7), ("Q04", 12)]
      and all(r["doc"] == 20481 and r["was_doc"] == 18070 and "KAMALADEVI" in r["title"] and "INCOME TAX" not in r["title"].upper() and "INCOME TAX" in (r["was_title"] or "").upper()
              for q in pool for r in q["rows"] if r.get("was_doc")), fixed_rows)


def passages_in(prompt_text):
    body = prompt_text[prompt_text.index("QUESTION: "):]
    parts = re.split(r"--- PASSAGE (\d+) ---\n", body)
    out = []
    for n, b in zip(parts[1::2], parts[2::2]):
        out.append((int(n), b[b.index("\nText:\n") + len("\nText:\n"):].split("\n" + "=" * 78)[0].strip()))
    return out


tmp = tempfile.mkdtemp(prefix="builders_")
try:
    # ---------------------------------------------------------------- first AI-reviewer pack
    print("--- first pack")
    zp, n_files = blp.build(rc.DEFAULT_POOL, tmp)
    z = zipfile.ZipFile(zp)
    names = z.namelist()
    check("24 prompt files, a README and the empty replies folder note", n_files == 24 and sum(n.startswith("LLM_Review_Pack/prompts/") for n in names) == 24
          and "LLM_Review_Pack/README.txt" in names and "LLM_Review_Pack/replies/PUT_SAVED_REPLIES_HERE.txt" in names, len(names))
    ok_text = ok_order = True
    for q in pool:
        want = {r["n"]: r["text"].strip() for r in first[q["qid"]]}
        for order in "AB":
            got = passages_in(z.read(f"LLM_Review_Pack/prompts/{q['qid']}_{order}.txt").decode())
            ok_order &= [n for n, _ in got] == sorted(want, reverse=(order == "B"))
            ok_text &= all(want[n] == t for n, t in got)
    check("every prompt holds exactly the first-batch passages, with the pool's numbers and text; A in number order, B reversed", ok_text and ok_order)
    allp = "\n".join(z.read(n).decode() for n in names).lower()
    check("no answer-key wording anywhere in the pack (search version names, landmark, rewrites)", not [w for w in VARIANT_WORDS if w.lower() in allp] and not any(r.lower()[:60] in allp for r in REWRITES),
          [w for w in VARIANT_WORDS if w.lower() in allp])
    check("no top-up passage leaked into the first pack", not any(r["text"].strip()[:80] in allp for q in pool for r in extra[q["qid"]] if r["h"] not in {x["h"] for x in first[q["qid"]]}))
    q3 = z.read("LLM_Review_Pack/prompts/Q03_A.txt").decode()
    check("each prompt states the passage count and numbers it must cover", f"exactly {len(first['Q03'])} entries" in q3 and f"Passage numbers to cover: {', '.join(str(n) for n in sorted(r['n'] for r in first['Q03']))}" in q3)
    zp2, _ = blp.build(rc.DEFAULT_POOL, os.path.join(tmp, "again"))
    check("building twice gives identical prompts", all(zipfile.ZipFile(zp2).read(n) == z.read(n) for n in names))

    # ---------------------------------------------------------------- top-up pack
    print("--- top-up pack")
    with_extra = [q["qid"] for q in pool if extra[q["qid"]]]
    zt, files, n_pass, n_q = btp.build(rc.DEFAULT_POOL, tmp)                              # default: TA files only
    zt_ = zipfile.ZipFile(zt)
    tn = zt_.namelist()
    check("10 prompts by default: one TA per question with new passages, and no TB", files == len(with_extra) == 10 and n_q == 10 and n_pass == 24
          and not any("_TB" in n for n in tn) and sum(n.startswith("Top_Up_Pack/prompts/") for n in tn) == 10, (files, n_q, n_pass))
    ok = True
    for qid in with_extra:
        want = {r["n"]: r["text"].strip() for r in extra[qid]}
        got = passages_in(zt_.read(f"Top_Up_Pack/prompts/{qid}_TA.txt").decode())
        ok &= [n for n, _ in got] == sorted(want) and all(want[n] == t for n, t in got)
    check("each top-up prompt holds exactly its question's top-up passages, in number order, numbered after the first batch", ok)
    check("top-up numbers never collide with first-batch numbers", all(not ({r["n"] for r in extra[q["qid"]]} & {r["n"] for r in first[q["qid"]]}) for q in pool))
    allt = "\n".join(zt_.read(n).decode() for n in tn if n.endswith(".txt")).lower()
    check("no answer-key wording in the top-up pack", not [w for w in VARIANT_WORDS if w.lower() in allt], [w for w in VARIANT_WORDS if w.lower() in allt])
    rd = zt_.read("Top_Up_Pack/README.txt").decode()
    check("the README counts are right, names the questions with nothing new, says 10 chats per tool and never mentions TB or reversing",
          "24 passages" in rd and "10 of the project's 12 test questions (Q01 and Q06 have none)" in rd and "Q03 has 1 to 16 in the larger set, so the passages here are numbered 17 to 19" in rd
          and "That is 10 chats per tool." in rd and "TB" not in rd and "reverse" not in rd.lower())
    check("the README stands alone: it never assumes the reader has the first pack, and tells them to add initials if someone else uses the same tool",
          "first pack" not in rd.lower() and "first review" not in rd.lower() and "-S.txt" in rd and "You do not need the rest of the set" in rd)
    one = zt_.read("Top_Up_Pack/prompts/Q12_TA.txt").decode()
    check("a one-passage prompt says 'exactly 1 entry' and its template uses the real number", "exactly 1 entry," in one and f'"no": {extra["Q12"][0]["n"]},' in one)
    ws = openpyxl.load_workbook(os.path.join(tmp, "Top_Up_Review_Sheet.xlsx"))["Review"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if isinstance(r[1], int)]
    check("the top-up sheet has the 24 passages, each with its dropdown cell, and unique (Q, No.)", len(rows) == 24 and len({(r[0], r[1]) for r in rows}) == 24 and ws.data_validations.dataValidation)
    cells = " ".join(str(c.value) for row in openpyxl.load_workbook(os.path.join(tmp, "Top_Up_Review_Sheet.xlsx")).worksheets[0:2] for r in row.iter_rows() for c in r if c.value).lower()
    check("no answer-key wording in the top-up sheet", not [w for w in VARIANT_WORDS if w.lower() in cells])
    # the optional reverse-order files
    zr, files_r, _, _ = btp.build(rc.DEFAULT_POOL, os.path.join(tmp, "rev"), with_reverse=True)
    zr_ = zipfile.ZipFile(zr)
    expect_r = sum(1 if len(extra[qid]) < 2 else 2 for qid in with_extra)
    ok = True
    for qid in with_extra:
        fn = f"Top_Up_Pack/prompts/{qid}_TB.txt"
        if len(extra[qid]) < 2:
            ok &= fn not in zr_.namelist()
        else:
            got = [n for n, _ in passages_in(zr_.read(fn).decode())]
            ok &= got == sorted((r["n"] for r in extra[qid]), reverse=True)
    check("--with-reverse adds a TB file (the passages reversed) only where a question has two or more new passages: 17 prompts", files_r == expect_r == 17 and ok, (files_r, expect_r))
    rdr = zr_.read("Top_Up_Pack/README.txt").decode()
    check("the --with-reverse README explains TA and TB and says 17 chats per tool", "_TB.txt" in rdr and "(17 chats per tool)" in rdr and "(10 chats per tool)" in rdr)
    check("the TA prompts are byte-identical with and without --with-reverse", all(zt_.read(f"Top_Up_Pack/prompts/{qid}_TA.txt") == zr_.read(f"Top_Up_Pack/prompts/{qid}_TA.txt") for qid in with_extra))

    # ---------------------------------------------------------------- human review sheet
    print("--- human review sheet")
    path, n_rows, last = brs.build(rc.DEFAULT_POOL, tmp)
    wb = openpyxl.load_workbook(path)
    check("sheets: Read me, Review, Question verdicts, Reviewer summary", wb.sheetnames == ["Read me", "Review", "Question verdicts", "Reviewer summary"])
    rv = wb["Review"]
    data = [r for r in rv.iter_rows(min_row=2, values_only=True) if isinstance(r[1], int)]
    check("175 passages, one per row, matching the pool's text and numbers", n_rows == 175 and len(data) == 175 and all(r[3] == rc.clean(next(x for x in first[r[0]] if x["n"] == r[1])["text"]) for r in data))
    dvs = rv.data_validations.dataValidation[0]
    covered = sum(bb[3] - bb[1] + 1 for bb in (range_boundaries(str(r_)) for r_ in dvs.sqref.ranges))
    check("the mark dropdown covers every passage row and nothing else", covered == 175, covered)
    cells = " ".join(str(c.value) for sh in wb.worksheets for row in sh.iter_rows() for c in row if c.value).lower()
    check("no answer-key wording anywhere in the human sheet", not [w for w in VARIANT_WORDS if w.lower() in cells], [w for w in VARIANT_WORDS if w.lower() in cells])
    os.makedirs(os.path.join(tmp, "n"))
    other_path = brs.build(rc.DEFAULT_POOL, os.path.join(tmp, "n"), judgments=9999)[0]
    readme_txt = lambda book: " ".join(str(c.value) for row in book["Read me"].iter_rows() for c in row if c.value)
    check("the Read me describes 8,236 judgments by default and takes other numbers", "8,236 Supreme Court" in readme_txt(wb) and "9,999 Supreme Court" in readme_txt(openpyxl.load_workbook(other_path)))

    def evaluate(book, sheet, formula):
        """Tiny evaluator for the formulas these sheets use: COUNTIF, COUNTIFS, COUNT, COUNTA, IF(COUNT()=0,"",AVERAGE()), and A-SUM(range) arithmetic."""
        def rng(ref, default_sheet):
            sh, _, cells_ = ref.rpartition("!") if "!" in ref else (default_sheet, "", ref)
            sh = sh.strip("'") or default_sheet
            c1, r1, c2, r2 = range_boundaries(cells_)
            return [book[sh].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        m = re.fullmatch(r"=COUNTIF\((.+?),\"(.*)\"\)", formula)
        if m: return sum(v == m.group(2) for v in rng(m.group(1), sheet))
        m = re.fullmatch(r"=COUNTIFS\((.+?),\"(.*?)\",(.+?),\"<>\"\)", formula)
        if m: return sum(a == m.group(2) and b not in (None, "") for a, b in zip(rng(m.group(1), sheet), rng(m.group(3), sheet)))
        m = re.fullmatch(r"=COUNTA\((.+?)\)", formula)
        if m: return sum(v not in (None, "") for v in rng(m.group(1), sheet))
        m = re.fullmatch(r"=COUNT\((.+?)\)", formula)
        if m: return sum(isinstance(v, (int, float)) for v in rng(m.group(1), sheet))
        m = re.fullmatch(r"=IF\(COUNT\((.+?)\)=0,\"\",AVERAGE\((.+?)\)\)", formula)
        if m:
            vals = [v for v in rng(m.group(1), sheet) if isinstance(v, (int, float))]
            return sum(vals) / len(vals) if vals else ""
        m = re.fullmatch(r"=B(\d+)-SUM\(B(\d+):B(\d+)\)", formula)
        if m: return evaluate_cell(book, sheet, f"B{m.group(1)}") - sum(evaluate_cell(book, sheet, f"B{r}") for r in range(int(m.group(2)), int(m.group(3)) + 1))
        raise ValueError(formula)

    def evaluate_cell(book, sheet, coord):
        v = book[sheet][coord].value
        return evaluate(book, sheet, v) if isinstance(v, str) and v.startswith("=") else v

    filled = openpyxl.load_workbook(path)
    marks = ["Yes", "Partly", "No", "Misleading", None]
    k = 0
    for row in filled["Review"].iter_rows(min_row=2):
        if isinstance(row[1].value, int):
            row[4].value = marks[k % 5]; k += 1
    for i in range(2, 14):
        filled["Question verdicts"].cell(i, 4).value = "Yes" if i % 3 else "No"
    counts = {m: sum(1 for r in filled["Review"].iter_rows(min_row=2, values_only=True) if r[4] == m) for m in marks[:4]}
    got = {lab: evaluate_cell(filled, "Reviewer summary", f"B{i}") for i, lab in zip(range(4, 11), ["good", "Yes", "Partly", "No", "Misleading", "total", "unmarked"])}
    check("summary counts add up on a filled copy (independent calculation)", got["Yes"] == counts["Yes"] and got["Partly"] == counts["Partly"] and got["No"] == counts["No"] and got["Misleading"] == counts["Misleading"]
          and got["total"] == 175 and got["unmarked"] == 175 - sum(counts.values()) and got["good"] == sum(1 for i in range(2, 14) if i % 3), got)

    # ---------------------------------------------------------------- authority list
    print("--- authority list")
    pa, n_j, n_a, n_r = bal.build(rc.DEFAULT_POOL, tmp)
    wa = openpyxl.load_workbook(pa)
    check("sheets: Read me, Judgments, Questions, Our results (open last), Summary", wa.sheetnames == ["Read me", "Judgments", "Questions", "Our results (open last)", "Summary"])
    jd = [r for r in wa["Judgments"].iter_rows(min_row=2, values_only=True)]
    check("33 judgments, 11 of them priority A, each id once", n_j == 33 and n_a == 11 and len(jd) == 33 and len({r[2] for r in jd}) == 33 and sum(r[1] == "A" for r in jd) == 11, (n_j, n_a))
    check("all six landmarks are on the list and in the first rows", {41736, 35436, 42630, 36322, 37060, 31821} <= {r[2] for r in jd[:6]})
    check("the flagged cases are on the list (mislabelled 18070, Prabha Shankar Dubey 25607, Baldev Singh 30317, the two Lalita Kumari entries)", {18070, 25607, 30317, 39112, 40960} <= {r[2] for r in jd})
    qs = [r for r in wa["Questions"].iter_rows(min_row=2, values_only=True)]
    check("the 12 questions are listed with the pool's exact wording", len(qs) == 12 and all(r[2] == q["question"] for r, q in zip(qs, pool)))
    ours = [r for r in wa["Our results (open last)"].iter_rows(min_row=2, values_only=True) if isinstance(r[1], int)]
    # 87 rows until the pool correction of 2026-10-04: judgment 18070 (removed as mislabelled) had one row each for Q03 and Q04, which now merge into 20481's
    check("'our results' lists the judgments behind each question's top passages, 85 rows in all", n_r == 85 == len(ours) and all(r[3] for r in ours))
    check("'Our results (open last)' comes after 'Questions', so the tool's own results are filled in first", wa.sheetnames.index("Our results (open last)") > wa.sheetnames.index("Questions"))
    filled = openpyxl.load_workbook(pa)
    statuses = ["Good law", "Overruled", "Partly overruled", "Doubted", "Distinguished", "Not found", "Unsure"]
    given = []
    for i in range(2, 35):
        if i % 4:
            filled["Judgments"].cell(i, 11).value = statuses[i % 7]
            filled["Judgments"].cell(i, 17).value = "Lawyer" if i % 3 == 0 else "Non-lawyer using a tool"
            given.append(i)
    for i in range(2, 10):
        filled["Questions"].cell(i, 9).value = ["Ours better", "Similar", "Ours worse", "Can't tell"][i % 4]
        filled["Questions"].cell(i, 8).value = i % 6
    sm = filled["Summary"]
    labels = {sm.cell(r, 1).value: evaluate_cell(filled, "Summary", f"B{r}") for r in range(3, 25) if sm.cell(r, 1).value and isinstance(sm.cell(r, 2).value, str) and str(sm.cell(r, 2).value).startswith("=")}
    check("summary counts add up on a filled copy (independent calculation)", labels["Judgments to check"] == 33 and labels["Judgments checked"] == len(given)
          and labels["Priority A judgments checked"] == sum(1 for i in given if filled["Judgments"].cell(i, 2).value == "A")
          and all(labels[f"Status: {s}"] == sum(1 for i in given if statuses[i % 7] == s) for s in statuses)
          and labels["Rows filled by a lawyer"] == sum(1 for i in given if i % 3 == 0) and labels["Questions with a verdict"] == 8
          and abs(labels["Average overlap with our list (tool top 5 found in our results)"] - sum(i % 6 for i in range(2, 10)) / 8) < 1e-9, labels)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

print("\nALL PASSED" if not FAILS else f"\nFAILED: {FAILS}")
sys.exit(1 if FAILS else 0)
