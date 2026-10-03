#!/usr/bin/env python3
"""Builds Top_Up_Pack.zip from the pool's top-up passages: AI-reviewer prompts (TA = number order, TB = reversed), a human review sheet and a README.
Top-up passage numbers continue after the first batch's numbers for the same question, so (question, number) is unique across both batches.
Nothing about which search found a passage is included.

  python review_tools/build_topup_pack.py --out DIR [--pool review_data/pool_all.json]
"""
import argparse
import math
import os
import shutil
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import review_common as rc

HEADER = """You are helping to test a legal search tool for Indian Supreme Court criminal-law judgments. Below is ONE question that a person typed into the tool, followed by {n} passages the tool showed. This is a short second batch for a question that has been reviewed before; judge each passage as an answer to that question, on its own.

RULES
1. Judge only the text of each passage. Do NOT use outside knowledge about whether the case is famous, important or still good law. If the passage text itself says a view was later overruled, doubted or rejected, say so in "caveats".
2. Judge each passage on its own, as if the other passages did not exist. The order of the passages carries no meaning.
3. Use ONE mark per passage, exactly one of: Yes, Partly, No, Misleading.
{meanings}
4. For every passage give "evidence": a quote copied EXACTLY (character for character, no ellipses, no paraphrase) from that passage, at most 30 words, that best supports your mark. For "No", quote the words that show it is off-topic.
5. For every passage give "reason": one plain sentence, at most 30 words.
6. For every passage give "confidence": "high", "medium" or "low".
7. Do not search the web and do not use tools. If you cannot tell, choose the safer lower mark and set confidence to "low".
8. Reply with ONE JSON object only. No text before or after it and no markdown code fences.

REQUIRED JSON SHAPE (a template: use your own values)
{{
  "question_id": "{qid}",
  "marks": [
    {{"no": {n1}, "mark": "Yes", "evidence": "exact quote from passage {n1}", "reason": "one sentence", "confidence": "high"}}{more}
  ],
  "good_answer_present": "Yes",
  "best_passage_no": {n1},
  "caveats": ""
}}
Notes: "marks" must have exactly {n} {entries}, one for each passage number listed below. "good_answer_present" is "Yes" if at least one passage in this batch answers the question well, otherwise "No". "best_passage_no" is the number of the single best passage in this batch, or null if none is good. "caveats" is an empty string unless a careful reader should know something.

Passage numbers to cover: {nums}
"""


def render(q, order):
    """The prompt for one question's top-up passages. q["rows"] must hold only the top-up passages. order 'A' = number order, 'B' = reversed."""
    rows = sorted(q["rows"], key=lambda r: r["n"], reverse=(order == "B"))
    asc = sorted(r["n"] for r in q["rows"])
    nums = ", ".join(map(str, asc))
    more = (",\n    {" + f'"no": {asc[1]}, "mark": "No", "evidence": "exact quote from passage {asc[1]}", "reason": "one sentence", "confidence": "medium"' + "}") if len(asc) > 1 else ""
    out = [f"QUESTION ID: {q['qid']}   (top-up batch, prompt version: order {order})", "",
           HEADER.format(n=len(rows), qid=q["qid"], n1=asc[0], more=more, nums=nums, meanings=rc.meaning_lines(), entries="entry" if len(rows) == 1 else "entries"),
           "=" * 78, f"QUESTION: {rc.clean(q['question'])}", "=" * 78, ""]
    out += rc.passage_blocks(rows)
    out += ["=" * 78, f"END OF PASSAGES. Now reply with the single JSON object for {q['qid']} (passages {nums})."]
    return "\n".join(out) + "\n"


def _names(qids):
    return ", ".join(qids[:-1]) + (" and " if len(qids) > 1 else "") + qids[-1]


def readme(top, n_orig, all_q, example):
    n_pass = sum(len(q["rows"]) for q in top)
    without = [q["qid"] for q in all_q if not rc.rows_of(q, True)]
    none_phrase = f" ({_names(without)} {'has' if len(without) == 1 else 'have'} none)" if without else ""
    ex_old = [r["n"] for r in rc.rows_of(example["full"], False)]
    ex_new = [r["n"] for r in example["new"]]
    chats_a, chats_ab = len(top), len(top) + sum(1 for q in top if len(q["rows"]) > 1)
    return f"""TOP-UP PACK: a second, small batch of passages to mark

WHAT THIS IS
The first review pack (Q01 to Q12) covered {n_orig} passages. Since then the project tested a second way of searching, and it showed {n_pass} passages that nobody has marked yet. This top-up asks for marks on those {n_pass} passages only. They belong to {len(top)} of the {len(all_q)} questions{none_phrase}. The passage numbers CONTINUE after the first pack's numbers for the same question (for example {example['full']['qid']} had {ex_old[0]} to {ex_old[-1]}, so the top-up for {example['full']['qid']} is {ex_new[0]} to {ex_new[-1]}).

If you have not done the first pack yet, please do it first or in the same sitting: the same question appears in both. If you already did the first pack, you only need this.

OPTION 1: AI reviewers (same method as the first pack)
1. Open the folder 'prompts'. Files named {example['full']['qid']}_TA.txt hold the passages in number order; {example['full']['qid']}_TB.txt hold the same passages in reverse order (only for questions with 2 or more passages).
2. Copy one whole file into a NEW chat of an AI tool (switch off web search and tools), send it, and wait.
3. Save the model's whole reply as plain text, named like  {example['full']['qid']}_TA_ChatGPT5.txt  or  {example['full']['qid']}_TB_Gemini.txt  (same model name every time, no spaces). Put the files in the folder 'replies'.
4. If a reply is not JSON or is cut off, send "Reply again with only the JSON object, nothing else." and save the new reply.
5. Zip the 'replies' folder and send it back.
Please use the same two or more AI tools you used for the first pack. Minimum: the TA files ({chats_a} chats per tool). Better: TA and TB ({chats_ab} chats per tool).

OPTION 2: a human reviewer
Open Top_Up_Review_Sheet.xlsx, read each passage as if you had asked that question, and pick Yes, Partly, No or Misleading from the dropdown in the yellow cell. Save under a new name with your initials and send it back.

PLEASE DO NOT
- Edit the prompt text or add your own opinion to the chat.
- Fix or tidy a model's JSON by hand. Save what it returned.
- Ask the model whether a case is still good law. That is checked separately.

The marks mean:
   Yes        - answers the question directly
   Partly     - on the topic, useful, but incomplete
   No         - not helpful
   Misleading - could give a wrong impression (for example a lawyer's argument that the court rejects)
Results will be labelled 'LLM-judged' (or 'expert-checked' if a lawyer marks them). Nothing is described as lawyer-validated unless a lawyer did it.
"""


def build_sheet(top, n_questions, path):
    """The human review sheet for the top-up passages, in the same layout as the first workbook's Review sheet. Returns the last Review row."""
    F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    HEAD, QROW, INPUT, GREY = (PatternFill("solid", fgColor=c) for c in ("1F3864", "D9E2F3", "FFF2CC", "F2F2F2"))
    wrap_top = Alignment(wrap_text=True, vertical="top")
    center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)
    n_pass = sum(len(q["rows"]) for q in top)
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Top-up Review Sheet: Criminal-Law Judgments"
    ws["A1"].font = F(bold=True, size=14)
    ws["A2"] = "A second, small batch for questions that were already in the first review sheet."
    ws["A2"].font = F(italic=True)
    rows_ = [("What this is", f"{n_pass} passages from {len(top)} of the {n_questions} questions that were shown by a second way of searching and are not in the first sheet. Passage numbers continue after the first sheet's numbers for the same question."),
             ("What to do", "Open the sheet 'Review'. Under each blue question banner read each passage as if you had asked that question, and pick one mark from the dropdown in the yellow 'Your mark' cell. Add a short comment if useful. The order of the passages says nothing about quality."),
             ("Save", "Save under a new name (add your initials) and send it back. Do not rename the sheets and do not insert, delete or sort rows."),
             ("Colour guide", "Yellow cells are for you. Please do not edit white or grey cells. 'Ref (do not edit)' lets the project team match your marks to the right passage."),
             ("Blind by design", "The sheet does not show which search returned a passage. Please judge on the text alone."),
             ("Repeated text", "If the same text appears in more than one judgment it is shown once; 'Same text also in judgment id(s)' lists the others. Mark it once."),
             ("Time", f"About {n_pass} passages: roughly an hour.")]
    r = 4
    for a, b in rows_:
        ws.cell(r, 1, a).font = F(bold=True)
        ws.cell(r, 1).alignment = wrap_top
        ws.cell(r, 2, b).font = F()
        ws.cell(r, 2).alignment = wrap_top
        ws.row_dimensions[r].height = max(18, 14 * math.ceil(len(b) / 100) + 4)
        r += 1
    r += 1
    ws.cell(r, 1, "Your mark").font = F(bold=True)
    ws.cell(r, 1).fill = GREY
    ws.cell(r, 2).fill = GREY
    ws.cell(r, 2, "Choose one per passage").font = F(bold=True)
    r += 1
    fills = {"Yes": "C6EFCE", "Partly": "FFEB9C", "No": "D9D9D9", "Misleading": "FFC7CE"}
    for k in rc.MARKS:
        ws.cell(r, 1, k).font = F(bold=True)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=fills[k])
        ws.cell(r, 2, rc.MEANING[k]).font = F()
        ws.cell(r, 2).alignment = wrap_top
        ws.row_dimensions[r].height = 28 if k == "Misleading" else 16
        r += 1

    rv = wb.create_sheet("Review")
    cols = [("Q", 7), ("No.", 6), ("Judgment (opening header, id, paragraph, page)", 40), ("Passage text", 100), ("Your mark", 13), ("Comment", 40), ("Same text also in judgment id(s)", 16), ("Ref (do not edit)", 26)]
    for i, (name, w) in enumerate(cols, 1):
        c = rv.cell(1, i, name)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = box
        rv.column_dimensions[chr(64 + i)].width = w
    rv.row_dimensions[1].height = 32
    row, ranges = 2, []
    for q in top:
        kind = "Plain English" if q["kind"] == "lay" else "Legal wording"
        for ci in range(1, 9):
            rv.cell(row, ci).fill = QROW
            rv.cell(row, ci).border = box
        rv.cell(row, 1, q["qid"]).font = F(bold=True)
        rv.cell(row, 3, kind).font = F(bold=True)
        rv.cell(row, 3).alignment = wrap_top
        rv.cell(row, 1).alignment = wrap_top
        c = rv.cell(row, 4, "QUESTION: " + rc.clean(q["question"]))
        c.font = F(bold=True, size=11)
        c.alignment = wrap_top
        rv.row_dimensions[row].height = max(34, 15 * math.ceil(len(q["question"]) / 95) + 8)
        row += 1
        first = row
        for p in sorted(q["rows"], key=lambda x: x["n"]):
            para = f"paragraph {p['pn']}" if p["pn"] else "unnumbered paragraph"
            j = f"{rc.clean(p['title']) if p['title'] else 'Opening header not available'}\nJudgment id {p['doc']} · {para} · page {p['page']}"
            vals = [q["qid"], p["n"], j, rc.clean(p["text"]), None, None, ", ".join(str(d) for d in p["also"]) or None, f"J{p['doc']}|p{p['page']}|{p['pn'] or '-'}|{p['h']}"]
            for ci, v in enumerate(vals, 1):
                c = rv.cell(row, ci, v)
                c.border = box
                c.font = F(color="7F7F7F" if ci == 8 else "000000")
                c.alignment = center_top if ci in (1, 2, 5) else wrap_top
                if isinstance(v, str):
                    c.data_type = "s"
            for ci in (5, 6):
                rv.cell(row, ci).fill = INPUT
            rv.row_dimensions[row].height = min(409, 13 * max(math.ceil(len(vals[3]) / 100), math.ceil(len(j) / 38) + 1, 3) + 6)
            row += 1
        ranges.append((first, row - 1))
    last = row - 1
    dv = DataValidation(type="list", formula1='"' + ",".join(rc.MARKS) + '"', allow_blank=True, showErrorMessage=True, errorTitle="Pick from the list", error="Please choose Yes, Partly, No or Misleading.",
                        showInputMessage=True, promptTitle="Your mark", prompt="Yes = answers directly | Partly = useful, not complete | No = not helpful | Misleading = could give a wrong impression")
    rv.add_data_validation(dv)
    for a, b in ranges:
        dv.add(f"E{a}:E{b}")
    rng = f"E2:E{last}"
    for m, col in fills.items():
        kw = dict(font=Font(name="Arial", color="9C0006", bold=True)) if m == "Misleading" else {}
        rv.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{m}"'], fill=PatternFill("solid", bgColor=col, fgColor=col), **kw))
    rv.freeze_panes = "C2"
    rv.auto_filter.ref = f"A1:H{last}"
    rv.sheet_view.zoomScale = 90
    rv.page_setup.orientation = "landscape"
    rv.page_setup.fitToWidth = 1
    rv.page_setup.fitToHeight = 0
    rv.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    rv.print_title_rows = "1:1"
    wb.properties.title = "Top-up Review Sheet"
    wb.properties.creator = "Recourse SCI search project"
    wb.save(path)
    return last


def build(pool_path, out_dir):
    """Writes out_dir/Top_Up_Pack/, out_dir/Top_Up_Pack.zip and out_dir/Top_Up_Review_Sheet.xlsx. Returns (zip path, prompt files, passages, questions)."""
    pool = rc.load_pool(pool_path)
    top = [dict(q, rows=rc.rows_of(q, True)) for q in pool if rc.rows_of(q, True)]
    if not top:
        raise ValueError("the pool has no top-up passages")
    n_orig = sum(len(rc.rows_of(q, False)) for q in pool)
    # the README's worked example: the first question with 2 or more top-up passages
    ex_q = next((q for q in top if len(q["rows"]) > 1), top[0])
    example = dict(full=next(q for q in pool if q["qid"] == ex_q["qid"]), new=ex_q["rows"])
    folder, zip_path = os.path.join(out_dir, "Top_Up_Pack"), os.path.join(out_dir, "Top_Up_Pack.zip")
    xlsx = os.path.join(out_dir, "Top_Up_Review_Sheet.xlsx")
    if os.path.isdir(folder):
        shutil.rmtree(folder)
    os.makedirs(os.path.join(folder, "prompts"))
    os.makedirs(os.path.join(folder, "replies"))
    files = 0
    for q in top:
        for order in ("A", "B"):
            if order == "B" and len(q["rows"]) < 2:
                continue
            with open(os.path.join(folder, "prompts", f"{q['qid']}_T{order}.txt"), "w", encoding="utf-8") as fh:
                fh.write(render(q, order))
            files += 1
    with open(os.path.join(folder, "README.txt"), "w", encoding="utf-8") as fh:
        fh.write(readme(top, n_orig, pool, example))
    with open(os.path.join(folder, "replies", "PUT_SAVED_REPLIES_HERE.txt"), "w", encoding="utf-8") as fh:
        fh.write(f"Save each model reply here, named like {example['full']['qid']}_TA_ChatGPT5.txt\n")
    build_sheet(top, len(pool), xlsx)
    shutil.copy(xlsx, os.path.join(folder, "Top_Up_Review_Sheet.xlsx"))
    rc.write_zip(folder, zip_path, "Top_Up_Pack")
    return zip_path, files, sum(len(q["rows"]) for q in top), len(top)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", required=True, help="folder to write into")
    ap.add_argument("--pool", default=rc.DEFAULT_POOL)
    a = ap.parse_args(argv)
    os.makedirs(a.out, exist_ok=True)
    zip_path, files, n, nq = build(a.pool, a.out)
    print(f"built {zip_path} ({os.path.getsize(zip_path):,} bytes): {files} prompt files, {n} passages, {nq} questions")


if __name__ == "__main__":
    main()
