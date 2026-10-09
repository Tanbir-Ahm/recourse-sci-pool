"""Tests for the intake of the owner's real questions (review_tools/real_questions.py): reading the Excel template and CSV, validation, the plan's counts, the sealing rule,
the freeze hash and the once-only guard for the sealed questions. All questions here are invented for the test.
Run: python review_tools/test_real_questions.py"""
import copy
import datetime
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import real_questions as RQ

FAILS = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  [{detail}]"))
    if not cond:
        FAILS.append(name)


def synthetic(counts=None):
    """30 invented, anonymised rows that satisfy every number in the plan."""
    counts = counts or {d: n for d, (_, n) in RQ.DOMAINS.items()}
    rows, i = [], 0
    for d, n in counts.items():
        for j in range(n):
            ctl = RQ.CONTROLS[j % 3] if d == "K" else ""
            rows.append({"id": "", "question": f"Invented test question number {i:02d} for domain {d} about a made-up situation", "kind": RQ.KINDS[i % 3],
                         "numbering": RQ.NUMBERING[(i // 3) % 3] if i % 3 else "none", "domain": d, "trap": "yes" if i in (4, 13, 22) else "no", "trap_note": "an older ruling" if i in (4, 13, 22) else "",
                         "best_law_era": "2017 or later" if i % 3 == 0 else "unsure", "language": "Hinglish" if i % 9 == 8 else "English", "control_type": ctl, "anonymised_ok": "yes"})
            i += 1
    return rows


def with_numbering(rows):
    """make sure at least 5 IPC/CrPC and 5 BNS/BNSS rows, whatever the rotation gave"""
    rows = copy.deepcopy(rows)
    for k, r in enumerate(rows):
        r["numbering"] = ("IPC/CrPC", "BNS/BNSS", "none")[k % 3]
    return rows


rows = with_numbering(synthetic())

# ---- validation -------------------------------------------------------------------------------------------------------------------------------------------
problems, warnings = RQ.validate(rows)
check("a complete invented set has no problems and no warnings", not problems and not warnings, (problems, warnings))
check("and meets every number in the plan", RQ.plan_gaps(rows) == [], RQ.plan_gaps(rows))

bad = copy.deepcopy(rows)
bad[0]["question"] = "too short"
bad[1]["kind"] = "story"
bad[2]["domain"] = "Z"
bad[3]["anonymised_ok"] = ""
bad[4]["numbering"] = "Penal Code"
bad[5]["trap"] = "maybe"
bad[6]["control_type"] = "other"
bad[6]["domain"] = "K"
bad[7]["question"] = bad[8]["question"]
p, _ = RQ.validate(bad)
for needle in ("row 1:", "row 2:", "row 3:", "row 4:", "row 5:", "row 6:", "row 7:"):
    check(f"validate reports a problem for {needle}", any(x.startswith(needle) for x in p), p)
check("a duplicate question is reported against the earlier row", any(x.startswith("row 9:") and "same question as row 8" in x for x in p), p)
check("an empty intake is a problem", RQ.validate([])[0] != [])

ident = copy.deepcopy(rows)
ident[0]["question"] = "My client Mr Sharma, phone 9876543210, says police in Ghaziabad arrested him"
ident[1]["question"] = "FIR No. 123/2023 was registered against my brother, what can we do about it"
ident[2]["question"] = "Please mail me at someone@example.com about the bail for my brother"
ident[3]["question"] = "Rakesh S/o Mohan R/o Delhi has been called by the police, what should he do"
ident[4]["question"] = "My PAN is ABCDE1234F and the bank froze my account after a complaint was filed"
ident[5]["question"] = "Under section 438 can my brother be protected before the police come to our house"
ident[5]["numbering"] = "none"
_, w = RQ.validate(ident)
for k, what in enumerate(("phone number", "case or FIR number", "email", "S/o", "PAN-style", "names a section number")):
    check(f"a warning for {what} on row {k + 1}", any(x.startswith(f"row {k + 1}:") for x in w), w)
check("warnings do not stop a build", RQ.validate(ident)[0] == [])

# ---- normalising what the owner types -------------------------------------------------------------------------------------------------------------
r = RQ.normalise(dict(rows[0], kind="Everyday", numbering="bns", domain="b - bail", trap="Y", best_law_era="2017+", language="hinglish", anonymised_ok="YES", control_type=""))
check("the owner's case and short forms are put into the allowed spelling", (r["kind"], r["numbering"], r["domain"], r["trap"], r["best_law_era"], r["language"], r["anonymised_ok"]) ==
      ("everyday", "BNS/BNSS", "B", "yes", "2017 or later", "Hinglish", "yes"), r)
check("a value that cannot be understood is left as typed, so validate can report it", RQ.normalise(dict(rows[0], kind="story"))["kind"] == "story")

# ---- the plan's counts -----------------------------------------------------------------------------------------------------------------------------
short = [r for r in rows if r["domain"] != "C"]
g = RQ.plan_gaps(short)
check("a missing domain and a short total are named", any("domain C" in x for x in g) and any("27 questions in all" in x for x in g), g)
few = copy.deepcopy(rows)
for r in few:
    r["trap"], r["best_law_era"], r["numbering"] = "no", "unsure", "none"
g = RQ.plan_gaps(few)
check("too few traps, BNS, IPC and 2017 questions are named", all(any(t in x for x in g) for t in ("currency trap", "BNS/BNSS", "IPC/CrPC", "2017")), g)

# ---- build, sealing and freeze -----------------------------------------------------------------------------------------------------------------
real, warn, gaps = RQ.build(rows, today="2026-10-08")
qs = real["questions"]
sealed = [q for q in qs if q["half"] == "sealed"]
check("30 questions get ids R01 to R30 in row order", [q["id"] for q in qs] == [f"R{i:02d}" for i in range(1, 31)])
check("with the plan's domain counts exactly 11 are sealed, one per domain", len(sealed) == 11 and sorted(q["domain"] for q in sealed) == sorted(RQ.DOMAINS), [q["domain"] for q in sealed])
check("the sealing rule is reproduced from the seed alone", RQ.assign_halves(qs) == {q["id"]: q["half"] for q in qs})
check("sealing does not depend on the question wording", RQ.assign_halves([dict(q, text="changed") for q in qs]) == RQ.assign_halves(qs))
check("a different seed gives a different split (the seed matters)", RQ.assign_halves(qs, "another-seed") != RQ.assign_halves(qs))
check("check() finds nothing wrong with a fresh build", RQ.check(real) == [], RQ.check(real))

uneven = RQ.assign_halves([{"id": f"X{i}", "domain": "A"} for i in range(7)])
check("a domain of 7 seals ceil(7/3) = 3", sum(1 for v in uneven.values() if v == "sealed") == 3)
one = RQ.assign_halves([{"id": "X0", "domain": "A"}])
check("a domain of one is sealed", one == {"X0": "sealed"})

tamper = copy.deepcopy(real)
sq = next(q for q in tamper["questions"] if q["half"] == "sealed")
sq["text"] = sq["text"] + " (edited)"
check("editing a sealed question after the freeze is detected", any("hash differs" in x for x in RQ.check(tamper)), RQ.check(tamper))
tamper2 = copy.deepcopy(real)
dq = next(q for q in tamper2["questions"] if q["half"] == "dev")
dq["text"] = dq["text"] + " (edited)"
check("editing a tuning question does not change the freeze", RQ.check(tamper2) == [] and RQ.freeze_digest(tamper2) == real["freeze"]["sha256"])
tamper3 = copy.deepcopy(real)
dq3 = next(q for q in tamper3["questions"] if q["half"] == "dev")
dq3["half"] = "sealed"
check("moving a question across the line is detected", RQ.check(tamper3) != [])

try:
    RQ.build(bad)
    ok = False
except ValueError:
    ok = True
check("build refuses an intake with problems", ok)

# ---- link-first sealing --------------------------------------------------------------------------------------------------------------------------
sides = {"T01": "dev", "T02": "dev", "T08": "sealed", "T13": "sealed"}
links = [{"id": "R10", "topic": "T13", "basis": "shared case", "note": "x"}, {"id": "R11", "topic": "T02", "basis": "doctrine group", "note": "x"},
         {"id": "R05", "topic": "T08", "basis": "doctrine group", "note": "x"}, {"id": "R01", "topic": "T01", "basis": "shared case", "note": "x"}]
lreal, _, _ = RQ.build(rows, today="2026-10-09", links=links, ruler=(sides, "abc"))
lh = {q["id"]: q["half"] for q in lreal["questions"]}
check("a question linked to a sealed ruler topic is sealed and one linked to a dev topic is dev", lh["R10"] == "sealed" and lh["R05"] == "sealed" and lh["R11"] == "dev" and lh["R01"] == "dev", lh)
dq = {q["id"]: q["domain"] for q in lreal["questions"]}
by_dom = {d: [i for i, x in dq.items() if x == d] for d in RQ.DOMAINS}
check("a domain whose forced questions already reach the target seals none of the others", sum(lh[i] == "sealed" for i in by_dom[dq["R10"]]) == 1 and sum(lh[i] == "sealed" for i in by_dom[dq["R05"]]) == 1)
check("every domain still has its ceil(n/3) sealed or more", all(sum(lh[i] == "sealed" for i in ids) >= -(-len(ids) // 3) for ids in by_dom.values()))
check("forced questions say why", lreal["questions"][9]["half_basis"] == "linked to ruler topic T13" and lreal["questions"][0]["half_basis"] == "linked to ruler topic T01" and lreal["questions"][2]["half_basis"] == "hash")
check("the question file records the ruler's freeze hash and a split hash", lreal["ruler_freeze_sha256"] == "abc" and len(lreal["freeze"]["split_sha256"]) == 64 and RQ.check(lreal) == [], RQ.check(lreal))
nolink, _, _ = RQ.build(rows, today="2026-10-09")
check("without links the plain hash rule gives the same split as before", {q["id"]: q["half"] for q in nolink["questions"]} == {q["id"]: q["half"] for q in real["questions"]})
t2 = copy.deepcopy(lreal)
t2["links"][0]["topic"] = "T01"
check("changing a link after the freeze is detected", any("split" in x or "sealing rule" in x for x in RQ.check(t2)), RQ.check(t2))
try:
    RQ.forced_sides([{"id": "R01", "topic": "T01"}, {"id": "R01", "topic": "T08"}], sides)
    ok = False
except ValueError:
    ok = True
check("a question linked to topics on both sides of the seal is refused", ok)
try:
    RQ.build(rows, links=[{"id": "R99", "topic": "T01", "basis": "x", "note": "x"}], ruler=(sides, "abc"))
    ok = False
except ValueError:
    ok = True
check("a link to a question that is not in the intake is refused", ok)
leak = copy.deepcopy(lreal)
leak["questions"][0]["expected_answer"] = "x"
check("a question file that holds an answer-key field fails the check", any("key fields" in x for x in RQ.check(leak)))
check("the built question file carries no trap notes", all("trap_note" not in q for q in lreal["questions"]))
cut = copy.deepcopy(rows)
for r in cut:
    if r["domain"] == "K":
        r["control_type"] = RQ.CUTOFF_TRAP
        break
check("a cut-off trap is a valid control type but does not cover the clean controls", RQ.validate(cut)[0] == [] and any("cut-off trap" in x for x in RQ.plan_gaps(cut)), RQ.plan_gaps(cut))
rsides, rhash = RQ.ruler_sides()
check("the ruler's frozen topics are read as sealed", [t for t, v in rsides.items() if v == "sealed"] == ["T08", "T12", "T13", "T16"] and len(rhash) == 64, rsides)

# ---- the guard -----------------------------------------------------------------------------------------------------------------------------------------
tmp = tempfile.mkdtemp()
runs = os.path.join(tmp, "runs.jsonl")
dev = RQ.select(real, "dev", runs_path=runs)
check("the tuning half is handed out freely and is the other 19", len(dev) == 19 and all(q["half"] == "dev" for q in dev) and not os.path.exists(runs))


def refused(**kw):
    try:
        RQ.select(real, "sealed", runs_path=runs, **kw)
        return False
    except RQ.SealedError:
        return True


check("sealed questions are refused without final=True", refused())
check("sealed questions are refused without a label", refused(final=True, label="  "))
check("a refusal writes nothing", not os.path.exists(runs))
now = datetime.datetime(2026, 11, 1, 10, 0, 0, tzinfo=datetime.timezone.utc)
got = RQ.select(real, "sealed", final=True, label="end of step 6", runs_path=runs, now=now)
log = RQ.read_runs(runs)
check("the first final run hands out the 11 and logs it", len(got) == 11 and len(log) == 1 and log[0]["label"] == "end of step 6" and log[0]["ts"] == "2026-11-01T10:00:00Z" and not log[0]["repeat"], log)
check("the log records the freeze hash and the ids", log[0]["freeze_sha256"] == real["freeze"]["sha256"] and log[0]["qids"] == [q["id"] for q in got])
check("a second final run is refused", refused(final=True, label="again"))
RQ.select(real, "sealed", final=True, label="again", runs_path=runs, allow_repeat=True, now=now)
check("a repeat needs allow_repeat and is logged as a repeat", [x["repeat"] for x in RQ.read_runs(runs)] == [False, True])
runs2 = os.path.join(tmp, "runs2.jsonl")
try:
    RQ.select(tamper, "sealed", final=True, label="x", runs_path=runs2)
    ok = False
except RQ.SealedError:
    ok = not os.path.exists(runs2)
check("a tampered sealed set is refused even for a first final run", ok)

# ---- the Excel template and the round trip -------------------------------------------------------------------------------------------------
import openpyxl

tpl = RQ.write_template(os.path.join(tmp, "template.xlsx"))
wb = openpyxl.load_workbook(tpl)
check("the template has a Read me sheet and a Questions sheet", wb.sheetnames == ["Read me", "Questions"], wb.sheetnames)
check("the example row is on the sheet, grey, and ignored on reading", wb["Questions"]["A2"].value == "EXAMPLE" and RQ.read_intake(tpl) == [])
check("the question columns carry dropdowns for kind, numbering, domain, trap, era, language, control and anonymised", len(wb["Questions"].data_validations.dataValidation) == 8)
check("every domain letter and its name are on the Read me sheet", all(any(str(c.value or "").startswith(f"{d}:") for c in wb["Read me"]["A"]) for d in RQ.DOMAINS))

for dv in wb["Questions"].data_validations.dataValidation:
    items = dv.formula1.strip('"').split(",")
    check(f"dropdown {dv.sqref} has no value with a comma in it and fits Excel's 255-character limit", len(dv.formula1) <= 257 and all(i.strip() in set(RQ.KINDS + RQ.NUMBERING + RQ.ERAS + RQ.LANGS + RQ.CONTROL_VALUES + tuple(RQ.DOMAINS) + ("yes", "no")) for i in items), dv.formula1)
check("a control typed as 'civil' or in the old wording is understood", RQ.normalise(dict(rows[0], control_type="civil"))["control_type"] == "civil/tax/divorce" and RQ.normalise(dict(rows[0], control_type="Civil, tax or divorce"))["control_type"] == "civil/tax/divorce")

ws = wb["Questions"]
for k, r in enumerate(rows):
    for j, col in enumerate(RQ.COLUMNS, 1):
        if col != "id":
            ws.cell(3 + k, j, r[col])
filled = os.path.join(tmp, "filled.xlsx")
wb.save(filled)
back = RQ.read_intake(filled)
check("a filled template reads back as the same 30 rows", len(back) == 30 and [{c: b[c] for c in RQ.COLUMNS if c != "id"} for b in back] == [{c: r[c] for c in RQ.COLUMNS if c != "id"} for r in rows])
real2, w2, g2 = RQ.build(back, today="2026-10-08")
check("the intake builds to the same seal as the rows did", real2["freeze"] == real["freeze"], (real2["freeze"], real["freeze"]))

csv_path = os.path.join(tmp, "filled.csv")
import csv
with open(csv_path, "w", encoding="utf-8", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=RQ.COLUMNS)
    w.writeheader()
    w.writerows(rows)
check("a CSV with the same columns reads the same", RQ.read_intake(csv_path) == back)

# ---- the command line --------------------------------------------------------------------------------------------------------------------------------
out = os.path.join(tmp, "real.json")
check("build writes the file", RQ.main(["build", "--intake", filled, "--out", out]) == 0 and os.path.exists(out))
check("check passes on it", RQ.main(["check", "--real", out]) == 0)
check("a second build is refused without --force", RQ.main(["build", "--intake", filled, "--out", out]) == 2)
check("select refuses sealed questions without --final", RQ.main(["select", "--real", out, "--half", "sealed", "--runs", runs2]) == 3)
check("select hands out the tuning half", RQ.main(["select", "--real", out, "--half", "dev", "--runs", runs2]) == 0)

print()
print("FAILED: " + ", ".join(FAILS) if FAILS else "all passed")
sys.exit(1 if FAILS else 0)
