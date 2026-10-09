#!/usr/bin/env python3
"""Builds the workbook for Mr. Advocate (the practising advocate who reviews the test): five short tasks on the 30 real questions.

    python review_tools/build_advocate_sheet.py --out DIR

  A  the 16 older rulings the owner's trap questions are set against: does he agree with what we think the later court did?
  B  the 7 answers the owner's own file says need a practising advocate (copied word for word)
  C  for each of the 30 questions: other authorities that would also be right, older or replaced rulings a search might wrongly show, and whether the expected answer is correct and current
  D  rulings whose holding a later statute displaced (the Dashrath Rupsingh Rathod type), as a list
  E  one or two clean control questions: answerable only from a High Court ruling (or a 2026 ruling), so that a correct tool says it found no Supreme Court authority

PRIVATE: the workbook holds the owner's answer key. It goes to Mr. Advocate only; never to an AI reviewer, J or S. Nothing in it is lawyer-checked until he fills it in. His answers go into the key
(never into the search), and are used only to score.
Reads: the owner's workbook (review_data/ruler/inputs/2026-10-09_owner_30_questions.xlsx), real30_key_v1.json, real30_stale_status_v1.json.
"""
import argparse
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
sys.path.insert(0, HERE)
import review_common as rc                    # noqa: E402

OWNER_XLSX = os.path.join(ROOT, "review_data", "ruler", "inputs", "2026-10-09_owner_30_questions.xlsx")
KEY = os.path.join(ROOT, "review_data", "ruler", "real30_key_v1.json")
STALE = os.path.join(ROOT, "review_data", "ruler", "real30_stale_status_v1.json")

KIND = {"overruled": "Overruled", "modified": "Modified (part of it changed)", "put_on_hold": "Put on hold or deferred", "superseded_by_statute": "Replaced by a statute",
        "narrower": "Narrower than today's law (not overruled)", "contested": "Contested (benches disagree)"}
AGREE = ["Yes", "No", "Partly"]                                          # no commas in a dropdown list: Excel splits on them
WARN = ["Yes always", "Yes for the changed part only", "No"]
HEAD = PatternFill("solid", start_color="DDE6F0")
INPUT = PatternFill("solid", start_color="FFF8DC")
NOTE = PatternFill("solid", start_color="F2F2F2")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def table(ws, headers, rows, widths, input_cols, first_row=1, dropdowns=None):
    """Header row, data rows (wrapped), the columns the advocate types in shaded, widths, frozen header, and dropdowns {column index: list}."""
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=first_row, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = Font(bold=True), HEAD, Alignment(wrap_text=True, vertical="center"), BOX
    for i, r in enumerate(rows, first_row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=rc.clean(v) if isinstance(v, str) else v)
            c.alignment, c.border = WRAP, BOX
            if j in input_cols:
                c.fill = INPUT
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    ws.freeze_panes = ws.cell(row=first_row + 1, column=1)
    last = first_row + len(rows)
    for col, values in (dropdowns or {}).items():
        assert not any("," in v for v in values), values
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "Please pick one of the listed answers (or leave blank and write in the comment column).", "Not in the list"
        ws.add_data_validation(dv)
        letter = ws.cell(row=1, column=col).column_letter
        dv.add(f"{letter}{first_row + 1}:{letter}{last}")


def owner_questions(path=OWNER_XLSX):
    """[{q, domain, text, expected, authority, ct, control}] in the owner's order, from the owner's workbook."""
    wb = load_workbook(path, data_only=True)
    ws = wb["30 Questions"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        text = str(r[3] or "").strip()
        if len(text) > 1 and text[0] == '"' and text[-1] == '"':
            text = text[1:-1]
        out.append({"q": str(r[0]), "domain": str(r[2] or ""), "text": text, "expected": str(r[8] or ""), "authority": str(r[9] or ""), "ct": str(r[6] or ""), "control": str(r[7] or "")})
    return out


def to_verify(path=OWNER_XLSX):
    wb = load_workbook(path, data_only=True)
    return [(str(r[0]), str(r[1])) for r in wb["To Verify"].iter_rows(min_row=2, values_only=True) if r[0]]


def older_rulings(key, stale):
    """[(ruling label, owner questions it is set against, kind, basis, source)] in the stale file's order."""
    used = {}
    for q in key["questions"].values():
        for s in q.get("stale_docs", []):
            used.setdefault(s["label"], []).append(q["owner_q"])
    return [(label, ", ".join(sorted(set(used.get(label, [])), key=lambda x: int(x[1:]))), KIND.get(v["kind"], v["kind"]), v["basis"], v["source"]) for label, v in stale["rulings"].items()]


def build(out_dir, owner_xlsx=OWNER_XLSX, key_path=KEY, stale_path=STALE):
    import json
    with open(key_path, encoding="utf-8") as fh:
        key = json.load(fh)
    with open(stale_path, encoding="utf-8") as fh:
        stale = json.load(fh)
    qs, tv, rulings = owner_questions(owner_xlsx), to_verify(owner_xlsx), older_rulings(key, stale)
    wb = Workbook()

    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("Review of a legal-search test for Indian criminal law: five short tasks for a practising advocate", True),
        ("", False),
        ("What this is. We are testing a search tool over Supreme Court criminal-law judgments with 30 questions the owner prepared. The tool can show an older ruling that has since been overruled, changed or replaced by a statute. "
         "We need a practising advocate to check our reading of the law. The sheets below hold our current reading. Nothing in them has been checked by a lawyer yet.", False),
        ("", False),
        ("What we ask (do them in this order; each sheet can be done alone)", True),
        ("A. Older rulings (16). For each older ruling, say whether you agree with what we think the later court did, and correct us if not. About 30 to 45 minutes.", False),
        ("B. Answers that need an advocate (7). The points the owner's own file lists as needing an advocate. About 30 minutes.", False),
        ("C. The 30 questions. For each: other authorities that would also be right, older or replaced rulings a search might wrongly show, and whether the expected answer is correct and current. "
         "If time is short, do the questions in the order they appear and stop where you must; two or three authorities per question is plenty.", False),
        ("D. Rulings replaced by a statute. A short list of rulings whose holding a later statute displaced (one example is filled in). Optional but very useful.", False),
        ("E. A clean control question. One or two questions answerable only from a High Court ruling (or from a very recent Supreme Court ruling that a library may not hold yet), so that a correct tool says it found no direct authority. Optional.", False),
        ("", False),
        ("How to fill it in. Type in the yellow cells only. Where a cell has a dropdown, pick from it; use the comment column for anything else. Please do not delete rows or reorder them.", False),
        ("", False),
        ("What happens to your answers. They are added to the answer key that scores the tool. They are never shown to the search tool and do not change how it searches. "
         "Results that rest on your answers will be described as checked by a practising advocate; results that do not will be described as unchecked.", False),
        ("", False),
        ("Please keep this workbook to yourself. It contains the answer key to the test; do not paste it into an AI tool and do not send it to anyone else who is helping to mark the tool's results.", True),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment, c.font = Alignment(wrap_text=True, vertical="top"), Font(bold=bold, size=12 if i == 1 else 11)
    n = len(lines) + 2
    for k, label in enumerate(("Name of the advocate who filled this in:", "Date:", "Sheets completed (A to E):")):
        ws.cell(row=n + k, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=n + k, column=2)
        c.fill, c.border = INPUT, BOX
    ws.column_dimensions["A"].width = 130
    ws.column_dimensions["B"].width = 40

    ws = wb.create_sheet("A Older rulings (16)")
    rows = [[i, q, label, kind, basis, src, "", "", "", "", ""] for i, (label, q, kind, basis, src) in enumerate(rulings, 1)]
    table(ws, ["#", "Used in question(s)", "Older ruling", "What we think happened (our reading)", "Why we think so (our notes; not lawyer-checked)", "Where our reading comes from",
               "Do you agree?", "What actually happened (your words)", "The later ruling, with citation", "Should a tool warn whenever it shows the older ruling?", "Comments"],
          rows, [4, 12, 34, 24, 60, 30, 12, 40, 34, 22, 36], {7, 8, 9, 10, 11}, dropdowns={7: AGREE, 10: WARN})

    ws = wb.create_sheet("B To verify (7)")
    rows = [[n, text, "", "", "", ""] for n, text in tv]
    table(ws, ["#", "The owner's list: answers needing verification by a practising advocate (word for word)", "Verified as stated?", "The correct position, as you understand it", "Source", "Comments"],
          rows, [4, 90, 14, 60, 34, 36], {3, 4, 5, 6}, dropdowns={3: AGREE})

    ws = wb.create_sheet("C The 30 questions")
    rows = [[x["q"], x["domain"], x["text"], x["authority"], x["expected"], "", "", "", "", "", ""] for x in qs]
    table(ws, ["Q", "Domain", "The question (as asked)", "The authority the owner's file names (older to current, or why it is a control)", "The expected answer (owner's file)",
               "Another authority that would also be a right answer (1)", "(2)", "(3)", "Older or replaced rulings a search might wrongly show for this question",
               "Expected answer correct and current?", "Comments (for instance what a High Court would add)"],
          rows, [5, 24, 60, 46, 60, 34, 30, 30, 40, 16, 40], {6, 7, 8, 9, 10, 11}, dropdowns={10: AGREE})

    ws = wb.create_sheet("D Replaced by statute")
    example = ["Dashrath Rupsingh Rathod v. State of Maharashtra (2014) 9 SCC 129", "Negotiable Instruments (Amendment) Act, 2015, s.142(2)", "A cheque presented through the payee's account is triable where the payee's branch is; the 2014 rule (drawee bank's location only) no longer holds", "Q9",
               "Example filled in by us, not yet confirmed by you", ""]
    table(ws, ["Ruling displaced by a later statute", "The statute and section", "What the statute changed", "Question(s) in this test it affects", "Your notes", "Confirmed? (for the example)"],
          [example] + [["", "", "", "", "", ""] for _ in range(10)], [48, 40, 60, 16, 40, 16], {1, 2, 3, 4, 5, 6}, dropdowns={6: AGREE})
    ws.cell(row=14, column=1, value="Please list only rulings whose holding was displaced by a later statute. The IPC to BNS and CrPC to BNSS renumbering is handled separately and need not be listed.").alignment = WRAP

    ws = wb.create_sheet("E Clean control")
    table(ws, ["A question in everyday words", "Why the answer lies only in a High Court ruling (or a very recent Supreme Court ruling)", "The authority (name, court, date)", "What a correct tool should say", "Comments"],
          [["", "", "", "", ""] for _ in range(3)], [60, 60, 40, 40, 36], {1, 2, 3, 4, 5})
    ws.cell(row=6, column=1, value="Why: the project's only control of this kind (Q29) is answered by a Supreme Court ruling of 27 July 2026. It works as a control only while the library does not hold that ruling (its 2026 coverage is partial and ends in August 2026). A question settled only by a High Court is a cleaner test.").alignment = WRAP
    ws.cell(row=6, column=1).fill = NOTE

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, "Mr_Advocate_Sheet.xlsx")
    wb.save(path)
    return path, {"rulings": len(rulings), "to_verify": len(tv), "questions": len(qs)}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    path, n = build(a.out)
    print(f"wrote {path}: {n['rulings']} older rulings, {n['to_verify']} items to verify, {n['questions']} questions")
    return 0


if __name__ == "__main__":
    sys.exit(main())
