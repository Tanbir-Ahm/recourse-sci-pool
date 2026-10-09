"""Tests for review_tools/build_advocate_sheet.py on the real inputs in the repository (the owner's workbook, the key, the older-rulings file). Run: python review_tools/test_build_advocate_sheet.py"""
import json
import os
import re
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_advocate_sheet as B

FAILS = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  [{detail}]"))
    if not cond:
        FAILS.append(name)


tmp = tempfile.mkdtemp()
path, n = B.build(tmp)
wb = load_workbook(path)
check("six sheets: a read-me and tasks A to E", wb.sheetnames == ["Read me", "A Older rulings (16)", "B To verify (7)", "C The 30 questions", "D Replaced by statute", "E Clean control"], wb.sheetnames)
check("the counts are the ones promised: 16 older rulings, 7 items to verify, 30 questions", n == {"rulings": 16, "to_verify": 7, "questions": 30}, n)

A = wb["A Older rulings (16)"]
rows = [[c.value for c in r] for r in A.iter_rows(min_row=2)]
check("sheet A lists 16 rulings, each with the question(s) it is used in and our reading", len(rows) == 16 and all(r[1] and r[2] and r[3] and r[4] for r in rows), [r[:4] for r in rows if not all(r[1:5])])
check("the kind is in plain words, not the code word", {r[3] for r in rows} <= set(B.KIND.values()) and "Overruled" in {r[3] for r in rows} and not any("superseded_by_statute" in str(r) or "counts_as_stale" in str(r) for r in rows))
check("the Shafhi Mohammad row says Q21 and the Dashrath row says Q9", any("Shafhi" in r[2] and r[1] == "Q21" for r in rows) and any("Dashrath" in r[2] and r[1] == "Q9" for r in rows), [(r[2][:20], r[1]) for r in rows])

Bs = wb["B To verify (7)"]
brows = [[c.value for c in r] for r in Bs.iter_rows(min_row=2)]
check("sheet B copies the owner's seven items word for word", len(brows) == 7 and "Shivangi Bansal" in brows[0][1] and brows[6][1].startswith("Citations confirmed only secondhand") and [r[0] for r in brows] == ["1", "2", "3", "4", "5", "6", "7"], [r[0] for r in brows])

C = wb["C The 30 questions"]
crows = [[c.value for c in r] for r in C.iter_rows(min_row=2)]
check("sheet C has the 30 questions in the owner's order with text, authority and expected answer", [r[0] for r in crows] == [f"Q{i}" for i in range(1, 31)] and all(r[2] and r[3] and r[4] for r in crows), [r[0] for r in crows if not all(r[2:5])])
check("a question's text does not carry the quotation marks of the owner's cell", not any(str(r[2]).startswith('"') for r in crows))
check("the three controls are there", sum(1 for r in crows if str(r[1]).startswith("K")) == 3)

D = wb["D Replaced by statute"]
check("sheet D has the Dashrath example filled in and blank rows after it", "Dashrath" in D["A2"].value and D["A3"].value in (None, "") and D.max_row >= 12)
E = wb["E Clean control"]
check("sheet E explains why Q29 is not a clean control", any("Q29" in str(c.value) for r in E.iter_rows() for c in r))

# dropdowns: no commas inside an option (Excel splits on them), and only on the columns meant for them
bad = []
for ws in wb.worksheets:
    for dv in ws.data_validations.dataValidation:
        opts = dv.formula1.strip('"').split(",")
        if any(o not in B.AGREE + B.WARN for o in opts):
            bad.append((ws.title, dv.formula1))
check("every dropdown holds only the listed options", not bad, bad)
check("the options contain no comma", all("," not in o for o in B.AGREE + B.WARN))
yellow = 0
for ws in wb.worksheets[1:]:
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.fill and c.fill.start_color and c.fill.start_color.rgb in ("00FFF8DC", "FFFFF8DC"):
                yellow += 1
check("the cells to type in are shaded", yellow > 100, yellow)
text = "\n".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value)
leaks = [w for w in ("sealed", "tuning", "freeze", "A today", "rerank", "re-rank", "R01", "R21", "pool.json", "side index", "overlay", "Voyage") if w.lower() in text.lower()]
check("nothing about the search versions, the halves, the question ids or the spending is in the workbook", not leaks, leaks)
check("the read-me says it is private, that nothing is lawyer-checked yet, and that the answers never reach the search", "keep this workbook to yourself" in text and "has been checked by a lawyer yet" in text and "never shown to the search tool" in text)

print()
print("FAILED: " + ", ".join(FAILS) if FAILS else "all passed")
sys.exit(1 if FAILS else 0)
