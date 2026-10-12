"""Tests for review_tools/fix_reading_sheet.py on an invented pool and invented replies. Run: python review_tools/test_fix_reading_sheet.py"""
import json
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fix_reading_sheet as R

FAILS = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  [{detail}]"))
    if not cond:
        FAILS.append(name)


t = tempfile.mkdtemp()
qs = []
for qid in ("N01", "N02", "N03"):
    rows = [{"n": i, "h": f"{qid}h{i}", "doc": 100 + i, "pn": str(i), "page": i, "title": f"Case {i}", "text": f"{i}. Passage {i} of {qid} about the point." + ("\x0c form feed and \x00 null in court text" if i == 3 else ""), "also": [], "variants": {"B baseline (rerank-3-lite)": i}} for i in range(1, 13)]
    qs.append({"qid": qid, "domain": "B", "kind": "legal", "question": f"Question text of {qid}?", "rows": rows, "lists": {}})
json.dump({"run": {}, "questions": qs}, open(os.path.join(t, "pool.json"), "w"))
rep = os.path.join(t, "rep"); os.makedirs(rep)
for q in qs:
    for tool in ("ChatGPT", "Gemini"):
        marks = []
        for r in q["rows"]:
            m = "Misleading" if (q["qid"], r["n"], tool) in {("N01", 3, "ChatGPT"), ("N02", 5, "Gemini"), ("N02", 6, "ChatGPT"), ("N02", 6, "Gemini")} else "Partly"
            marks.append({"no": r["n"], "mark": m, "evidence": r["text"][:15], "reason": "x", "confidence": "high"})
        open(os.path.join(rep, f"{q['qid']}_A_{tool}.txt"), "w").write(json.dumps({"question_id": q["qid"], "marks": marks, "good_answer_present": "Yes", "best_passage_no": None, "caveats": ""}))
info = R.build(os.path.join(t, "pool.json"), rep, os.path.join(t, "out"), n_random=20)
check("every passage that ANY run marked Misleading is listed (3 passages here), plus 20 random others", info["misleading"] == 3 and info["random"] == 20 and info["rows"] == 23, info)
ws = openpyxl.load_workbook(os.path.join(t, "out", "Reading_Sheet_for_Mr_Advocate.xlsx"))["Passages"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
mis_rows = {(r[1], r[3]) for r in rows if r[0].startswith("marked")}
check("the three Misleading passages are N01 no. 3, N02 no. 5 and N02 no. 6", mis_rows == {("N01", 3), ("N02", 5), ("N02", 6)}, mis_rows)
check("a random check passage was never marked Misleading by either run", all("Misleading" not in (r[6], r[7]) for r in rows if r[0] == "random check"))
check("the sheet shows both tools' marks and leaves Mr. Advocate's columns empty", rows[0][6] and rows[0][7] and all(r[8] is None and r[9] is None and r[10] is None for r in rows))
check("the sheet does not say which search version showed a passage", not any("baseline" in str(c) or "rerank" in str(c) for r in rows for c in r))
info2 = R.build(os.path.join(t, "pool.json"), rep, os.path.join(t, "out2"), n_random=20)
rows2 = list(openpyxl.load_workbook(os.path.join(t, "out2", "Reading_Sheet_for_Mr_Advocate.xlsx"))["Passages"].iter_rows(min_row=2, values_only=True))
check("the random draw is reproducible from its seed", [(r[1], r[3]) for r in rows] == [(r[1], r[3]) for r in rows2])
check("control characters in court text (a form feed, a null) do not stop the sheet being written and are replaced by spaces", all("\x0c" not in str(r[5]) and "\x00" not in str(r[5]) for r in rows))
print("\nall passed" if not FAILS else f"\n{len(FAILS)} FAILED: {FAILS}")
sys.exit(1 if FAILS else 0)
