"""The owner's real questions (the plan of 2026-10-08: 30 questions in ten domains plus controls, about a third sealed).

    python review_tools/real_questions.py template --out review_data/ruler/real_questions_template.xlsx
    python review_tools/real_questions.py build --intake FILLED.xlsx_or.csv      # validates, assigns ids R01.., seals by the fixed rule, writes real30_v1.json
    python review_tools/real_questions.py check                                  # structure, the plan's counts, the freeze hash
    python review_tools/real_questions.py select --half sealed --final --label "end of step 6"

The intake has one row per question: the question as the client said it (anonymised), its kind (everyday / legal / narrative), whether it uses IPC/CrPC or BNS/BNSS
section numbers, its domain (A to K, the table in PLAN_30_REAL_QUESTIONS.md), whether it is a currency trap, and optional notes. No answer and no case name goes in.

Sealing (agreed 2026-10-08): inside each domain the questions are ordered by SHA-256 of `seed|question id` and the first ceil(n/3) are sealed (one of three, one of two,
one of one). Written down and hashed before any search; a sealed question is neither searched nor marked until the final run, which `select` hands out once and logs.
PRIVATE, like the rest of review_data/.
"""
import argparse
import csv
import datetime
import hashlib
import json
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
DATA = os.path.join(ROOT, "review_data", "ruler")
DEFAULT_REAL = os.path.join(DATA, "real30_v1.json")
DEFAULT_RUNS = os.path.join(DATA, "real_frozen_runs.jsonl")
DEFAULT_TEMPLATE = os.path.join(DATA, "real_questions_template.xlsx")
SEED = "recourse-real-v1"
SCHEMA = "real-questions-1"

DOMAINS = {
    "A": ("Arrest and custody rights", 3), "B": ("Bail: regular, anticipatory, default", 3), "C": ("Cheque bounce (Section 138)", 3), "D": ("NDPS", 2),
    "E": ("FIR quashing; civil dispute dressed as crime (cheating, forgery, breach of trust)", 3), "F": ("Marriage-related offences: 498A, dowry, 304B/306, domestic violence", 3),
    "G": ("Sexual offences and POCSO", 3), "H": ("Violent crime: murder, hurt, evidence, sentence", 3), "I": ("Special-statute arrest and bail: PMLA, UAPA, SC/ST, corruption", 2),
    "J": ("After conviction: appeal, suspended sentence, remission, parole", 2), "K": ("Controls (no good answer)", 3),
}
KINDS = ("everyday", "legal", "narrative")
NUMBERING = ("none", "IPC/CrPC", "BNS/BNSS")
ERAS = ("before 2017", "2017 or later", "unsure")
LANGS = ("English", "Hinglish", "other")
CONTROLS = ("civil/tax/divorce", "High Court only or 2026", "statute table answers it")  # no commas: Excel splits a dropdown list on them
CUTOFF_TRAP = "cut-off trap"          # a domain K question whose newest ruling is not in the collection: older law answers it, so it is not a clean control
CONTROL_VALUES = CONTROLS + (CUTOFF_TRAP,)
RULE_TEXT = ("a question that shares a case, or sits in the same doctrine group, with a ruler topic takes that topic's side (a frozen topic means sealed); the remaining questions are "
             "ordered inside each domain by SHA-256(seed|id) and sealed until ceil(n/3) of the domain is sealed (forced ones count towards that)")
COLUMNS = ["id", "question", "kind", "numbering", "domain", "trap", "trap_note", "best_law_era", "language", "control_type", "anonymised_ok"]
HEADERS = {"id": "ID (leave blank)", "question": "Question, as the client said it", "kind": "Kind", "numbering": "Section numbers used", "domain": "Domain (A to K)", "trap": "Currency trap?",
           "trap_note": "Which older ruling is the trap (private)", "best_law_era": "Best law is from", "language": "Language", "control_type": "Control type (domain K only)",
           "anonymised_ok": "Anonymised? (yes)"}
# things that identify a person or a case: reported as warnings for the owner to look at, never silently dropped
IDENTIFIERS = [
    ("an email address", re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")),
    ("a phone number", re.compile(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{9}(?!\d)")),
    ("a long number (Aadhaar, account, vehicle?)", re.compile(r"(?<!\d)\d{4}\s?\d{4}\s?\d{4}(?!\d)")),
    ("a case or FIR number", re.compile(r"\b(?:FIR|F\.I\.R\.|Crl\.?\s?(?:A|M\.?P|Appeal)|C\.?C\.?|S\.?L\.?P\.?|W\.?P\.?|Case)\s*(?:No\.?|Number)?\s*\d+\s*(?:/|of)\s*\d{2,4}", re.I)),
    ("a family-relation address form (S/o, D/o, W/o, R/o)", re.compile(r"\b[SDWR]/o\b", re.I)),
    ("a PAN-style code", re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b")),
]


class SealedError(RuntimeError):
    pass


# ---- reading the intake ---------------------------------------------------------------------------------------------------------------------------------

def read_intake(path):
    """Rows of the intake (CSV, or the 'Questions' sheet of an Excel file) as dicts keyed by COLUMNS. Rows with an empty question, or an id starting EXAMPLE, are skipped."""
    if str(path).lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["Questions"] if "Questions" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        head = [str(c).strip() if c is not None else "" for c in rows[0]]
        by_label = {v: k for k, v in HEADERS.items()}
        keys = [by_label.get(h, h) for h in head]
        recs = [{k: ("" if v is None else str(v).strip()) for k, v in zip(keys, r)} for r in rows[1:]]
    else:
        with open(path, encoding="utf-8", newline="") as fh:
            recs = [{k: (v or "").strip() for k, v in r.items()} for r in csv.DictReader(fh)]
    out = []
    for r in recs:
        if not (r.get("question") or "").strip() or (r.get("id") or "").upper().startswith("EXAMPLE"):
            continue
        row = {c: (r.get(c) or "").strip() for c in COLUMNS}
        out.append(normalise(row))
    return out


def _norm(value, allowed, aliases=None):
    """The allowed spelling for a typed value (case-insensitive, with a few short forms); anything else is returned as typed so that validate can report it."""
    v = (value or "").strip()
    if not v:
        return ""
    for a in allowed:
        if v.lower() == a.lower():
            return a
    return (aliases or {}).get(v.lower(), v)


NUMBERING_ALIASES = {"ipc": "IPC/CrPC", "crpc": "IPC/CrPC", "ipc/crpc": "IPC/CrPC", "ipc, crpc": "IPC/CrPC", "ipc and crpc": "IPC/CrPC",
                     "bns": "BNS/BNSS", "bnss": "BNS/BNSS", "bns/bnss": "BNS/BNSS", "bns, bnss": "BNS/BNSS", "bns and bnss": "BNS/BNSS", "no": "none", "-": "none", "nil": "none"}
ERA_ALIASES = {"2017+": "2017 or later", "2017 on": "2017 or later", "after 2017": "2017 or later", "old": "before 2017", "pre-2017": "before 2017", "not sure": "unsure", "?": "unsure"}


def normalise(r):
    """Case and short forms the owner may type ('Yes', 'bns', 'b') are put into the allowed spelling; the domain may be written 'A' or 'A - Arrest'."""
    r = dict(r)
    r["kind"] = _norm(r["kind"], KINDS)
    r["numbering"] = _norm(r["numbering"], NUMBERING, NUMBERING_ALIASES)
    d = r["domain"].strip().upper()
    r["domain"] = d[0] if len(d) > 1 and d[0] in DOMAINS and not d[1].isalnum() else d
    r["trap"] = _norm(r["trap"], ("yes", "no"), {"y": "yes", "n": "no", "true": "yes", "false": "no"})
    r["best_law_era"] = _norm(r["best_law_era"], ERAS, ERA_ALIASES)
    r["language"] = _norm(r["language"], LANGS)
    r["control_type"] = _norm(r["control_type"], CONTROL_VALUES, {"cut-off": CUTOFF_TRAP, "civil, tax or divorce": "civil/tax/divorce", "civil": "civil/tax/divorce", "tax": "civil/tax/divorce", "divorce": "civil/tax/divorce",
                                                            "high court only": "High Court only or 2026", "2026": "High Court only or 2026", "statute table": "statute table answers it"})
    r["anonymised_ok"] = _norm(r["anonymised_ok"], ("yes",), {"y": "yes", "true": "yes"})
    return r


def validate(rows):
    """(problems, warnings). Problems stop a build; warnings are shown to the owner (a possible identifier, a shortfall against the plan)."""
    problems, warnings = [], []
    if not rows:
        return ["no questions found in the intake"], warnings
    for i, r in enumerate(rows, 1):
        tag = f"row {i}"
        q = r["question"]
        if not 15 <= len(q) <= 1500:
            problems.append(f"{tag}: the question is {len(q)} characters (15 to 1,500 allowed)")
        if r["kind"] not in KINDS:
            problems.append(f"{tag}: kind {r['kind']!r} is not one of {', '.join(KINDS)}")
        if r["numbering"] not in NUMBERING:
            problems.append(f"{tag}: section numbers {r['numbering']!r} is not one of {', '.join(NUMBERING)}")
        if r["domain"] not in DOMAINS:
            problems.append(f"{tag}: domain {r['domain']!r} is not one of A to K")
        if r["trap"] not in ("yes", "no", ""):
            problems.append(f"{tag}: trap must be yes, no or blank")
        if r["best_law_era"] and r["best_law_era"] not in ERAS:
            problems.append(f"{tag}: best law era {r['best_law_era']!r} is not one of {', '.join(ERAS)}")
        if r["language"] and r["language"] not in LANGS:
            problems.append(f"{tag}: language {r['language']!r} is not one of {', '.join(LANGS)}")
        if r["domain"] == "K" and r["control_type"] and r["control_type"] not in CONTROL_VALUES:
            problems.append(f"{tag}: control type {r['control_type']!r} is not one of {'; '.join(CONTROL_VALUES)}")
        if r["anonymised_ok"].lower() != "yes":
            problems.append(f"{tag}: the owner has not confirmed 'Anonymised? yes'")
        for what, rx in IDENTIFIERS:
            if rx.search(q) or rx.search(r["trap_note"]):
                warnings.append(f"{tag}: may contain {what}; please check it is anonymised")
        sec = re.findall(r"\b(?:section|sections|s\.|ss\.)\s*\d", q, re.I)
        if sec and r["numbering"] == "none":
            warnings.append(f"{tag}: the question names a section number but 'Section numbers used' says none")
    seen = {}
    for i, r in enumerate(rows, 1):
        key = re.sub(r"\W+", " ", r["question"].lower()).strip()
        if key in seen:
            problems.append(f"row {i}: the same question as row {seen[key]}")
        seen[key] = i
    return problems, warnings


def plan_gaps(rows):
    """Shortfalls against the agreed plan (not errors): the questions per domain, kinds, section-number styles, traps, era."""
    gaps = []
    n = {d: sum(1 for r in rows if r["domain"] == d) for d in DOMAINS}
    for d, (name, want) in DOMAINS.items():
        if n[d] != want:
            gaps.append(f"domain {d} ({name}): {n[d]} question(s), the plan has {want}")
    if len(rows) != 30:
        gaps.append(f"{len(rows)} questions in all, the plan has 30")
    for k in KINDS:
        c = sum(1 for r in rows if r["kind"] == k)
        if c < 5:
            gaps.append(f"only {c} {k} question(s), the plan wants at least 5")
    for k in ("IPC/CrPC", "BNS/BNSS"):
        c = sum(1 for r in rows if r["numbering"] == k)
        if c < 5:
            gaps.append(f"only {c} question(s) using {k} section numbers, the plan wants at least 5")
    c = sum(1 for r in rows if r["trap"] == "yes")
    if c < 3:
        gaps.append(f"only {c} currency trap(s), the plan wants at least 3")
    c = sum(1 for r in rows if r["best_law_era"] == "2017 or later")
    if c < 8:
        gaps.append(f"only {c} question(s) marked 'best law is from 2017 or later' (the owner may leave this blank when unsure), the plan wants at least 8")
    kc = {t: sum(1 for r in rows if r["domain"] == "K" and r["control_type"] == t) for t in CONTROLS}
    if n["K"] and any(v == 0 for v in kc.values()):
        cut = sum(1 for r in rows if r["domain"] == "K" and r["control_type"] == CUTOFF_TRAP)
        gaps.append("the controls should cover: " + "; ".join(f"{t} ({kc[t]})" for t in CONTROLS) + (f"; {cut} question(s) are cut-off traps, which do not count as clean controls" if cut else ""))
    return gaps


# ---- the split ----------------------------------------------------------------------------------------------------------------------------------------

def ruler_sides(path=None):
    """{ruler topic id: 'sealed' | 'dev'} and the ruler's freeze hash, read from the ruler (its frozen topics are the sealed ones)."""
    with open(path or os.path.join(DATA, "ruler_v2.json"), encoding="utf-8") as fh:
        ruler = json.load(fh)
    frozen = set(ruler["freeze"]["topics"])
    return {t["id"]: ("sealed" if t["id"] in frozen else "dev") for t in ruler["topics"]}, ruler["freeze"]["sha256"]


def forced_sides(links, sides):
    """{question id: 'sealed' | 'dev'} from the link table. A question linked to topics on both sides is an error: the rule cannot be kept."""
    forced = {}
    for l in links:
        side = sides[l["topic"]]
        if forced.setdefault(l["id"], side) != side:
            raise ValueError(f"{l['id']} is linked to ruler topics on both sides of the seal")
    return forced


def assign_halves(questions, seed=SEED, forced=None):
    """{id: 'sealed'|'dev'}. Questions in `forced` keep that side. Inside each domain the others are ordered by SHA-256(seed|id) and sealed until ceil(n/3) of the domain is sealed
    (forced ones count; none of the others is sealed once the forced ones reach the target)."""
    forced = forced or {}
    out, by = {}, {}
    for q in questions:
        by.setdefault(q["domain"], []).append(q["id"])
    for d, ids in by.items():
        want = math.ceil(len(ids) / 3)
        have = sum(1 for i in ids if forced.get(i) == "sealed")
        free = sorted((i for i in ids if i not in forced), key=lambda i: hashlib.sha256(f"{seed}|{i}".encode("utf-8")).hexdigest())
        extra = set(free[:max(want - have, 0)])
        for i in ids:
            out[i] = forced.get(i) or ("sealed" if i in extra else "dev")
    return out


def freeze_digest(real):
    body = [{k: q[k] for k in ("id", "domain", "kind", "numbering", "text", "trap", "best_law_era", "language", "control_type")}
            for q in sorted((q for q in real["questions"] if q["half"] == "sealed"), key=lambda q: q["id"])]
    return hashlib.sha256(json.dumps(body, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()


def split_digest(real):
    """Hash of the whole split: the seed, the rule, every question's domain and side, and the link table."""
    body = {"seed": real["seed"], "rule": real["rules"]["seal"], "halves": sorted((q["id"], q["domain"], q["half"]) for q in real["questions"]),
            "links": sorted((l["id"], l["topic"], l["basis"]) for l in real.get("links", []))}
    return hashlib.sha256(json.dumps(body, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()


def build(rows, today=None, links=None, ruler=None, snapshot=None, provenance=None):
    """The question file (no answers, no trap notes: those go to the key). `links` is the list of {id, topic, basis, note}; `ruler` is (sides, freeze hash) from ruler_sides()."""
    problems, warnings = validate(rows)
    if problems:
        raise ValueError("; ".join(problems))
    qs = [{"id": f"R{i:02d}", "owner_q": r.get("owner_q") or "", "text": r["question"], "kind": r["kind"], "numbering": r["numbering"], "domain": r["domain"], "trap": r["trap"] == "yes",
           "best_law_era": r["best_law_era"] or "unsure", "language": r["language"] or "English", "control_type": r["control_type"], "marks": {"state": "none"}} for i, r in enumerate(rows, 1)]
    links = list(links or [])
    known = {q["id"] for q in qs}
    if any(l["id"] not in known for l in links):
        raise ValueError("a link names a question that is not in the intake")
    forced = {}
    if links:
        sides, ruler_hash = ruler or ruler_sides()
        forced = forced_sides(links, sides)
    halves = assign_halves(qs, SEED, forced)
    for q in qs:
        q["half"] = halves[q["id"]]
        q["half_basis"] = "linked to ruler topic " + ", ".join(sorted(l["topic"] for l in links if l["id"] == q["id"])) if q["id"] in forced else "hash"
    real = {"schema": SCHEMA, "built": today or datetime.date.today().isoformat(), "seed": SEED, "domains": {d: v[0] for d, v in DOMAINS.items()}, "questions": qs,
            "links": links, "rules": {"seal": RULE_TEXT, "sealed": "neither searched nor marked until the final run; select hands it out once and logs it"}}
    if links:
        real["ruler_freeze_sha256"] = ruler_hash
    if snapshot:
        real["index_snapshot"] = snapshot
    if provenance:
        real["provenance"] = provenance
    real["freeze"] = {"sha256": freeze_digest(real), "split_sha256": split_digest(real), "questions": sorted(q["id"] for q in qs if q["half"] == "sealed"),
                      "note": "Written before any search was run on these questions. Run once, at the end, with `select --half sealed --final`."}
    return real, warnings, plan_gaps(rows)


def check(real, ruler_path=None):
    """Problems found in a built question file: structure, the sealing rule reproduced from the links, both hashes, and (when the ruler file is given) the link sides."""
    bad = []
    ids = [q["id"] for q in real["questions"]]
    if len(ids) != len(set(ids)):
        bad.append("duplicate ids")
    for q in real["questions"]:
        if q["domain"] not in DOMAINS or q["kind"] not in KINDS or q["half"] not in ("sealed", "dev"):
            bad.append(f"{q['id']}: bad domain, kind or half")
    links = real.get("links", [])
    forced = {}
    if links and ruler_path:
        sides, ruler_hash = ruler_sides(ruler_path)
        forced = forced_sides(links, sides)
        if real.get("ruler_freeze_sha256") != ruler_hash:
            bad.append("the ruler's frozen topics have changed since the links were applied")
    elif links:
        forced = {q["id"]: q["half"] for q in real["questions"] if q.get("half_basis", "").startswith("linked")}
    if assign_halves(real["questions"], real["seed"], forced) != {q["id"]: q["half"] for q in real["questions"]}:
        bad.append("the halves do not follow the sealing rule")
    if real["freeze"]["sha256"] != freeze_digest(real):
        bad.append("the sealed questions have changed since the freeze (hash differs)")
    if "split_sha256" in real["freeze"] and real["freeze"]["split_sha256"] != split_digest(real):
        bad.append("the split (sides or links) has changed since the freeze (hash differs)")
    if sorted(q["id"] for q in real["questions"] if q["half"] == "sealed") != real["freeze"]["questions"]:
        bad.append("freeze list differs from the questions marked sealed")
    leaked = [k for q in real["questions"] for k in q if k in ("trap_note", "expected", "expected_answer", "authority", "authorities", "source_urls")]
    if leaked:
        bad.append("the question file holds key fields: " + ", ".join(sorted(set(leaked))))
    return bad


# ---- the guard ----------------------------------------------------------------------------------------------------------------------------------------

def read_runs(path):
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]


def select(real, half, final=False, runs_path=DEFAULT_RUNS, label="", allow_repeat=False, now=None):
    qs = [q for q in real["questions"] if q["half"] == half]
    if half != "sealed":
        return qs
    if not final:
        raise SealedError("the sealed questions are run once, at the end: ask for them with final=True and a label")
    if not label.strip():
        raise SealedError("a final run needs a label saying what is being judged")
    if freeze_digest(real) != real["freeze"]["sha256"]:
        raise SealedError("the sealed questions have changed since the freeze; do not run them")
    prior = read_runs(runs_path)
    if prior and not allow_repeat:
        raise SealedError(f"the sealed questions were already handed out on {prior[0]['ts']} ({prior[0]['label']}); a repeat needs allow_repeat=True and is logged as a repeat")
    stamp = (now or datetime.datetime.now(datetime.timezone.utc)).strftime("%Y-%m-%dT%H:%M:%SZ")
    os.makedirs(os.path.dirname(runs_path) or ".", exist_ok=True)
    with open(runs_path, "a", encoding="utf-8") as fh:
        fh.write(json.dumps({"ts": stamp, "label": label, "repeat": bool(prior), "qids": [q["id"] for q in qs], "freeze_sha256": real["freeze"]["sha256"]}) + "\n")
    return qs


# ---- the Excel template for the owner -----------------------------------------------------------------------------------------------------------------

def write_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    F = "Arial"
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    lines = [("30 REAL QUESTIONS: how to fill this in", True),
             ("Use the 'Questions' sheet. One row per question; the grey first row is an EXAMPLE and is ignored. Yellow cells are yours to fill. Do not add answers or case names.", False),
             ("1. Question: as the client said it, one to three sentences, in the client's own words (everyday), the lawyer's wording (legal) or a long messy story (narrative).", False),
             ("2. ANONYMISE it: no names, places, case or FIR numbers, phone numbers. The questions go to Voyage (search), to ChatGPT/Gemini (marking by J and S) and into a private repository. Type 'yes' in the last column only when it is anonymised.", False),
             ("3. Kind: everyday, legal or narrative. Section numbers used: none, IPC/CrPC or BNS/BNSS. Domain: the letter from the list below. Currency trap: yes when the first ruling a search will find was later overtaken (say which in 'Which older ruling', private).", False),
             ("4. 'Best law is from' and 'Language' are optional (leave blank if unsure). Domain K (controls) also needs a control type.", False),
             ("5. Aim for 30 questions: " + ", ".join(f"{d} {n}" for d, (_, n) in DOMAINS.items()) + ". At least 5 of each kind, 5 using BNS/BNSS numbers, 5 using IPC/CrPC numbers, 3 traps, 8 whose best law is 2017 or later.", False),
             ("", False), ("DOMAINS", True)] + [(f"{d}: {name} (about {n})", False) for d, (name, n) in DOMAINS.items()] + [
             ("", False), ("CONTROLS (domain K)", True), ("One question from outside criminal law (civil, tax, divorce); one whose answer is High Court-only or from 2026; one that a statute table answers (for example 'what is the punishment under BNS 103').", False),
             ("", False), ("EXAMPLE ROW (also on the Questions sheet, grey)", True),
             ("Question: 'Police came at night and took my brother. They gave no paper. Can they do that?'  Kind: everyday.  Section numbers used: none.  Domain: A.  Trap: no.  Anonymised: yes.", False)]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=F, size=12 if i == 1 else 10, bold=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140
    q = wb.create_sheet("Questions")
    for j, col in enumerate(COLUMNS, 1):
        c = q.cell(1, j, HEADERS[col])
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    example = {"id": "EXAMPLE", "question": "Police came at night and took my brother. They gave no paper. Can they do that?", "kind": "everyday", "numbering": "none", "domain": "A", "trap": "no",
               "trap_note": "", "best_law_era": "unsure", "language": "English", "control_type": "", "anonymised_ok": "yes"}
    for j, col in enumerate(COLUMNS, 1):
        c = q.cell(2, j, example[col])
        c.font = Font(name=F, size=10, italic=True, color="808080")
        c.fill = PatternFill("solid", fgColor="EDEDED")
    for i in range(3, 3 + 40):
        for j in range(1, len(COLUMNS) + 1):
            c = q.cell(i, j)
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if COLUMNS[j - 1] != "id":
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    widths = {"id": 10, "question": 70, "kind": 12, "numbering": 14, "domain": 10, "trap": 10, "trap_note": 36, "best_law_era": 16, "language": 12, "control_type": 26, "anonymised_ok": 14}
    for j, col in enumerate(COLUMNS, 1):
        q.column_dimensions[chr(64 + j)].width = widths[col]
    q.freeze_panes = "A2"
    lists = {"kind": KINDS, "numbering": NUMBERING, "domain": tuple(DOMAINS), "trap": ("yes", "no"), "best_law_era": ERAS, "language": LANGS, "control_type": CONTROL_VALUES, "anonymised_ok": ("yes",)}
    for col, vals in lists.items():
        j = COLUMNS.index(col) + 1
        letter = chr(64 + j)
        dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
        q.add_data_validation(dv)
        dv.add(f"{letter}3:{letter}42")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


# ---- command line -------------------------------------------------------------------------------------------------------------------------------------

def load(path=DEFAULT_REAL):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("template")
    t.add_argument("--out", default=DEFAULT_TEMPLATE)
    b = sub.add_parser("build")
    b.add_argument("--intake", required=True)
    b.add_argument("--out", default=DEFAULT_REAL)
    b.add_argument("--force", action="store_true", help="overwrite an existing file (this changes the seal)")
    c = sub.add_parser("check")
    c.add_argument("--real", default=DEFAULT_REAL)
    s = sub.add_parser("select")
    s.add_argument("--real", default=DEFAULT_REAL)
    s.add_argument("--half", choices=["dev", "sealed"], required=True)
    s.add_argument("--final", action="store_true")
    s.add_argument("--label", default="")
    s.add_argument("--allow-repeat", action="store_true")
    s.add_argument("--runs", default=DEFAULT_RUNS)
    a = ap.parse_args(argv)
    if a.cmd == "template":
        print("wrote", write_template(a.out))
        return 0
    if a.cmd == "build":
        if os.path.exists(a.out) and not a.force:
            print(f"{a.out} exists; a rebuild changes the seal. Use --force only on purpose.")
            return 2
        if os.path.exists(a.out) and read_runs(DEFAULT_RUNS):
            print("the sealed questions were already handed out; the question set can no longer be rebuilt.")
            return 2
        rows = read_intake(a.intake)
        try:
            real, warnings, gaps = build(rows)
        except ValueError as e:
            print("not built:", e)
            return 1
        os.makedirs(os.path.dirname(a.out), exist_ok=True)
        with open(a.out, "w", encoding="utf-8") as fh:
            json.dump(real, fh, indent=1, ensure_ascii=False)
            fh.write("\n")
        print(f"{len(real['questions'])} questions; sealed {len(real['freeze']['questions'])}: {', '.join(real['freeze']['questions'])}; freeze {real['freeze']['sha256'][:16]}")
        for w in warnings:
            print("WARNING:", w)
        for g in gaps:
            print("GAP:", g)
        return 0
    real = load(a.real)
    if a.cmd == "check":
        problems = check(real)
        print(f"{len(real['questions'])} questions, {len(real['freeze']['questions'])} sealed; problems: {problems or 'none'}")
        return 1 if problems else 0
    try:
        qs = select(real, a.half, final=a.final, runs_path=a.runs, label=a.label, allow_repeat=a.allow_repeat)
    except SealedError as e:
        print("refused:", e)
        return 3
    print(" ".join(q["id"] for q in qs))
    return 0


if __name__ == "__main__":
    sys.exit(main())
