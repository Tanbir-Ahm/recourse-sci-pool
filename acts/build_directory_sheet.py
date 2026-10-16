"""
build_directory_sheet.py -- writes acts/ACT_DIRECTORY_SHEET.docx, the page(s) for Mr. Advocate to check the Act Directory (STATUTE_COVERAGE_PLAN.md step 1) from site_patch/acts_directory.json.

Nothing in the drafted columns is shown to a user on the site until a person confirms it (python act_directory.py confirm ID status|appeal|link --by NAME, in the Recourse checkout; Claude records the
reviewer's answers, never on its own). Run: python acts/build_directory_sheet.py [directory.json] [out.docx]
"""
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_IN = os.path.join(os.path.dirname(HERE), "site_patch", "acts_directory.json")
DEFAULT_OUT = os.path.join(HERE, "ACT_DIRECTORY_SHEET.docx")


def _hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "14")
    rpr.append(color)                                   # the schema's order inside w:rPr is colour, size, underline (Word rejects other orders)
    rpr.append(sz)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


def _cell(cell, text, size=8, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    return p


def build(src=DEFAULT_IN, out=DEFAULT_OUT):
    data = json.load(open(src, encoding="utf-8"))
    acts = sorted(data["acts"], key=lambda a: (a.get("priority", 9), a["name"]))
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.4))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)

    doc.add_heading("Act Directory v1: sheet for Mr. Advocate", level=1)
    p = doc.add_paragraph()
    p.add_run("What this is. ").bold = True
    p.add_run("The Recourse site holds the full text of the BNS and BNSS only (plus a few sections of the NI Act and the IT Act). When a question names, or clearly describes, an Act whose text the site does not hold, "
              "the site now says so in one fixed notice and points to the official text. The lines below are Claude's DRAFTS, written from general knowledge and from India Code search results; "
              "Claude could not open India Code from its environment, so no link has been clicked. ")
    p.add_run("Nothing in the draft columns is shown to a user until you confirm it.").bold = True
    p = doc.add_paragraph()
    p.add_run("What we ask. ").bold = True
    p.add_run("For each row, please check against the Act itself (the link), not against the draft: (1) Status: is the line right? (2) Appeal route: is the line right? (3) Link: does it open the right, current Act? "
              "Write OK, or write the correction. A row you leave blank stays hidden. Rows are in the order the owner chose for loading the full text: the first five are the priority Acts.")
    p = doc.add_paragraph()
    p.add_run("Confidence column: ").bold = True
    p.add_run("Claude's own confidence in its draft: certain (read on India Code or beyond doubt), likely (recollection of the law, please read with care), guessing (please rewrite).")

    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    heads = ["No.", "Act", "Status (draft)", "Conf.", "Where appeals go (draft)", "Link (click: right Act?)", "Your check: status / appeal / link", "Corrections"]
    widths = [Cm(0.9), Cm(5.2), Cm(5.0), Cm(1.4), Cm(5.0), Cm(2.8), Cm(3.0), Cm(2.9)]
    for i, h in enumerate(heads):
        _cell(table.rows[0].cells[i], h, size=8, bold=True)
    for n, a in enumerate(acts, 1):
        row = table.add_row().cells
        _cell(row[0], str(n))
        _cell(row[1], f"{a['name']}\nAct {a['act_no']}" + ("\n(not detected by the site: the concordance already maps it)" if a.get("detect") is False else "") + ("\nPRIORITY: first five Acts for loading text" if a.get("priority") == 1 else ""))
        _cell(row[2], a["status"]["text"])
        _cell(row[3], a["status"]["confidence"])
        ap = a.get("appeal") or {}
        _cell(row[4], ap.get("text") or "(not drafted)" + (f" [{ap.get('confidence')}]" if ap.get("confidence") else ""))
        if ap.get("text") and ap.get("confidence"):
            row[4].paragraphs[0].add_run(f" [{ap['confidence']}]").font.size = Pt(8)
        p = _cell(row[5], "")
        _hyperlink(p, a["url"], "open on India Code")
        if a.get("url_alt"):
            p.add_run("\n").font.size = Pt(7)
            _hyperlink(p, a["url_alt"], "(second link)")
        if a.get("note"):
            p.add_run("\n" + a["note"]).font.size = Pt(7)
        _cell(row[6], "S: ____   A: ____   L: ____")
        _cell(row[7], "")
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    doc.add_heading("Phrases that make the site say “looks like it is about the … Act”", level=2)
    doc.add_paragraph("Besides the Act's name or its usual short forms, a few phrases that mean nothing else in practice also trigger the notice (worded “looks like it is about”, never “names”). "
                      "Please strike any phrase that could reasonably mean something else.")
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Table Grid"
    _cell(t2.rows[0].cells[0], "Act", bold=True)
    _cell(t2.rows[0].cells[1], "Phrases (strike any that mislead)", bold=True)
    for a in acts:
        if a.get("phrases"):
            r = t2.add_row().cells
            _cell(r[0], a["short"])
            _cell(r[1], "; ".join(a["phrases"]))
    doc.add_paragraph()
    doc.add_paragraph("Short names the site recognises for an Act (it never uses an ambiguous short form such as CPC, HMA, NIA or “IT Act”): see the aliases in acts_directory.json; the full list is in the drafts folder and can be read out on request.")
    doc.add_paragraph("Please send the sheet back (a photo of a marked printout is fine). Claude records only what you confirm; an unconfirmed row stays hidden.")
    doc.save(out)
    return out


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_IN
    out = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT
    print("wrote", build(src, out))
