"""
build_shelf_sheet.py -- writes acts/ACT_SHELF_SHEET.docx, the sheet for Mr. Advocate for the Act shelf (STATUTE_COVERAGE_PLAN.md step 2): about ten sections of each of the five loaded Acts, in full, as Recourse
will hand them to the answer model, to be checked against the Act; the wording the site will show; and three questions only a lawyer can answer (what the shelf cannot know).

The sections are Claude's choice, not a statistical sample, and the sheet says so: the sections the Act Directory's appeal line for that Act cites, the sections a pointer names, the sections that were the marked
answer to the evaluation questions the Act Directory can detect, and the sections that carry the Act's main offence or protection (Claude's judgment). A column says why each is there.

Nothing here is recorded as confirmed by anyone. Claude records only what Mr. Advocate answers.
Run: python acts/build_shelf_sheet.py --recourse <Recourse checkout> [--out acts/ACT_SHELF_SHEET.docx]
"""
import argparse
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_directory_sheet import _hyperlink  # noqa: E402

# (section, why it is on the sheet).  Reasons that start with "appeal line" or "pointer" are checked against the data when the sheet is built.
CHOICE = {
    "pc_act_1988": [("2", "who is a 'public servant' (definitions; Claude's choice)"), ("7", "the main offence: a public servant taking an undue advantage"), ("8", "the main offence: giving one"),
                    ("13", "criminal misconduct (after the 2018 amendment)"), ("17", "who may investigate (evaluation question PC08)"), ("17A", "prior approval for an enquiry (evaluation question PC04)"),
                    ("19", "previous sanction (evaluation question PC03)"), ("20", "presumption (evaluation question PC06)"), ("27", "appeal line: the High Court's appeal and revision")],
    "pocso_act_2012": [("3", "penetrative sexual assault (the main offence)"), ("4", "its punishment (evaluation question PS01)"), ("19", "reporting (evaluation question PS03)"), ("24", "recording a child's statement (evaluation question PS05)"),
                       ("28", "appeal line: Special Courts"), ("29", "presumption (evaluation question PS06)"), ("31", "appeal line: the Code applies to a Special Court"), ("35", "time for the child's evidence and the trial (PS07)"),
                       ("42", "alternate punishment (PS13)")],
    "uapa_1967": [("2", "appeal line: the definition of 'court' (also the Act's other definitions)"), ("15", "terrorist act (evaluation question UA03)"), ("16", "its punishment"), ("17", "raising funds (UA04)"),
                  ("20", "membership of a terrorist gang (UA11)"), ("26", "appeal line: forfeiture ordered by the court"), ("28", "appeal line: appeal against forfeiture"), ("35", "designating an organisation or an individual (UA06)"),
                  ("43D", "bail and the 180-day extension (UA01, UA02)"), ("45", "previous sanction (UA05)")],
    "pmla_2002": [("2", "definitions, including 'proceeds of crime' (PM08)"), ("3", "the offence of money-laundering"), ("4", "its punishment (PM07)"), ("5", "attachment (PM03)"), ("19", "power to arrest (PM02)"),
                  ("26", "appeal line: to the Appellate Tribunal"), ("42", "appeal line: to the High Court"), ("44", "appeal line: offences triable by Special Courts"), ("45", "bail: the twin conditions (PM01)"),
                  ("46", "appeal line: the Code applies to a Special Court")],
    "sc_st_act_1989": [("2", "definitions (SS13)"), ("3", "the offences (SS01, SS08, SS11)"), ("4", "neglect of duties by a public servant (SS06)"), ("8", "presumption (SS07)"), ("14", "Special Courts (SS10)"),
                       ("14A", "appeal line: appeals"), ("15A", "rights of victims and witnesses (SS05)"), ("18", "pointer: anticipatory bail (section 438 of the Code)"), ("18A", "pointer: no enquiry or approval; section 438 again")],
}
ORDER = ["pc_act_1988", "pocso_act_2012", "uapa_1967", "pmla_2002", "sc_st_act_1989"]


def _small(cell, text, size=7, bold=False, italic=False, grey=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if grey:
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return p


def _check_choice(rec, directory, manifest, sections):
    """The sheet's own consistency: every section exists and is not a stub; 'appeal line' reasons are cited by that Act's appeal line; 'pointer' reasons are pointed to by the manifest."""
    problems = []
    for key, rows in CHOICE.items():
        appeal = (next(a for a in directory["acts"] if a["id"] == key).get("appeal") or {}).get("text") or ""
        pointed = {s for p in next(a for a in manifest["acts"] if a["key"] == key).get("pointers", []) for s in p["sections"]}
        for sec, why in rows:
            c = sections[key].get(sec)
            if c is None or c.get("omitted"):
                problems.append(f"{key} s.{sec}: not on the shelf or a stub")
            if why.startswith("appeal line") and not re.search(r"\bsections?\b[^.]*?(?<![0-9A-Za-z])" + re.escape(sec) + r"(?![0-9A-Za-z])", appeal):
                problems.append(f"{key} s.{sec}: the appeal line does not cite it")
            if why.startswith("pointer") and sec not in pointed:
                problems.append(f"{key} s.{sec}: no pointer names it")
    return problems


def build(recourse, out):
    shelf_dir = os.path.join(recourse, "act_shelf")
    manifest = json.load(open(os.path.join(shelf_dir, "manifest.json"), encoding="utf-8"))
    directory = json.load(open(os.path.join(recourse, "acts_directory.json"), encoding="utf-8"))
    sections = {k: {c["section_number"]: c for c in json.load(open(os.path.join(shelf_dir, f"{k}.json"), encoding="utf-8"))} for k in ORDER}
    load_report = json.load(open(os.path.join(shelf_dir, "LOAD_REPORT.json"), encoding="utf-8"))["reports"]
    problems = _check_choice(None, directory, manifest, sections)
    if problems:
        raise SystemExit("the sheet's choice of sections is inconsistent with the data:\n  " + "\n  ".join(problems))
    sys.path.insert(0, recourse)
    import act_directory as AD
    import act_shelf as S
    held = {k: {"state": "used", "loaded_on": next(a for a in manifest["acts"] if a["key"] == k)["loaded_on"]} for k in ORDER}
    pm = [{"id": "pmla_2002", "short": next(a for a in directory["acts"] if a["id"] == "pmla_2002")["short"], "how": "named"}]
    notice = AD.notice(pm, "answer", held={"pmla_2002": held["pmla_2002"]})["text"]
    notice_unused = AD.notice(pm, "answer", held={"pmla_2002": {"state": "unused", "loaded_on": held["pmla_2002"]["loaded_on"]}})["text"]
    caption = f"Official text from India Code, loaded {S.loaded_label(held['pmla_2002']['loaded_on'])}; not yet read by our reviewer"

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.4))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)

    doc.add_heading("Act shelf v1: sheet for Mr. Advocate", level=1)
    p = doc.add_paragraph()
    p.add_run("What this is. ").bold = True
    p.add_run("Recourse can now keep the full text of an Act that is not in its BNS/BNSS library, and hand the sections that bear on a question to the answer model, when the question names that Act "
              "(for example \"section 45 of the PMLA\", \"POCSO\", \"SC/ST Act\"). Five Acts are loaded: the Prevention of Corruption Act, 1988; POCSO, 2012; the UAPA, 1967; the PMLA, 2002; the SC/ST (Prevention of Atrocities) Act, 1989. "
              "The text comes from India Code's own database. A program then compared it with India Code's PDF of the same Act: the arrangement of sections is identical, and every section was read against the PDF's own text word by word, in order, "
              "in both directions (so a doubled word, a lost apostrophe or a whole lost proviso is found, not only an unknown word). It found, and the loader now corrects against the PDF, among others: a proviso of the PMLA's section 44 that the database had lost, "
              "and the words \"is said to commit aggravated sexual assault\" that POCSO's section 9 had in the wrong place. The differences left (the PDF's stubs of omitted sections, two slips of the PDF's own text layer, one Act number the database adds) are listed against what the PDF reads. "
              "Not on the shelf: the Schedules at the back of POCSO (the armed and security forces), the PMLA (the scheduled offences), the UAPA (the listed organisations and individuals) and the SC/ST Act (the IPC offences), and the State or Union Territory amendments that India Code prints inside three sections (PC Act 17A, POCSO 42, SC/ST 14): "
              "they apply only in that State and each of the three sections ends with a one-line note saying they exist. Both two-way comparisons and the Schedules are explained in the last section of the sheet. "
              "A program cannot tell whether a provision has since been amended, struck down or read down, or whether the text is the one a court would apply. ")
    p.add_run("That is why this sheet exists. Nothing on it is treated as checked by a person until you say so.").bold = True
    p = doc.add_paragraph()
    p.add_run("What we ask. ").bold = True
    p.add_run("(1) For each section below: does the text match the Act as in force today? Write OK, or write the correction. (2) Say whether the wording the site will show (the part headed \"Wording\", near the end) is fair. "
              "(3) Answer the three questions at the end. The sections are Claude's choice, not a statistical sample: a column says why each is here (the Act Directory's appeal line for the Act cites it, a pointer names it, "
              "it answered one of our 75 test questions, or it carries the Act's main offence or protection). If there is a section you would check first and it is not here, name it.")
    p = doc.add_paragraph()
    p.add_run("What your answer changes. ").bold = True
    p.add_run("A sample check is recorded as a sample (which sections, by whom, when). The status of an Act's text stays \"not yet read by our reviewer\" unless you tell us you have read all of it. "
              "Claude records only what you answer, never on your behalf, and corrections to the text are made by reloading the Act with the correction listed against India Code's PDF, so they can be checked.")
    p = doc.add_paragraph()
    p.add_run("What the shelf cannot know. ").bold = True
    p.add_run("It holds the words of the Act, not what the courts have held about them (for example the Supreme Court's reading of the PMLA's section 45 or section 19, or of the UAPA's section 43D). "
              "The answer model is told to use only the sections it is given and not to state a case's holding that is not in the material; nobody has measured how well it obeys with these Acts, because the environment this was built in cannot call the answer model. "
              "The shelf is searched only when the question names the Act or uses a listed phrase (about two in five of our test questions did).")

    for key in ORDER:
        m = next(a for a in manifest["acts"] if a["key"] == key)
        d = next(a for a in directory["acts"] if a["id"] == key)
        doc.add_heading(f"{m['name']} ({m['citation']})", level=2)
        p = doc.add_paragraph()
        _hyperlink(p, d["url"], "open the Act on India Code")
        p.add_run("   |   ")
        _hyperlink(p, sections[key][CHOICE[key][0][0]]["source_url"], "its section list")
        fixes = len(m.get("corrections", []))
        oc = load_report[key]["ordered_check"]
        p2 = doc.add_paragraph()
        r = p2.add_run(f"{m['sections_loaded']} sections loaded on {m['loaded_on']} (some are omitted or repealed and shown as such); compared with India Code's PDF \"as on {m['pdf_as_on']}\": arrangement of sections identical; "
                       f"{oc['sections_compared']} sections ({oc['words_compared']:,} words) read against the PDF word by word in both directions, {oc['differences']} difference(s) left, each explained; "
                       f"{fixes} correction(s) of India Code's database against the PDF, {len(m.get('accepted', []))} spelling variant(s) accepted. India Code's database record is newer than the PDF. Status: {m['text_status']} (not read by a person).")
        r.font.size = Pt(8)
        r.italic = True
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        heads = ["Section", "Why it is on the sheet", "The text as Recourse holds it", "Your check: OK, or the correction"]
        widths = [Cm(1.5), Cm(4.0), Cm(15.5), Cm(5.2)]
        for i, h in enumerate(heads):
            _small(t.rows[0].cells[i], h, size=8, bold=True)
        for sec_no, why in CHOICE[key]:
            c = sections[key][sec_no]
            row = t.add_row().cells
            _small(row[0], f"s.{sec_no}", size=8, bold=True)
            _small(row[1], why, size=7, italic=True, grey=True)
            _small(row[2], c["text"], size=7)
            fn = (c.get("footnotes") or "").strip()
            if fn:
                para = row[2].add_paragraph()
                rr = para.add_run("Amendment notes India Code prints: " + fn.replace("\n", " "))
                rr.font.size = Pt(6.5)
                rr.italic = True
                rr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            _small(row[3], "", size=7)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w

    doc.add_heading("Wording: what the site will show", level=2)
    doc.add_paragraph("When the shelf's text reached the answer, the notice above the answer reads (this is the PMLA; the others are the same with their own names and links):")
    q = doc.add_paragraph(notice)
    q.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph("Under each statute in \"Read the source\", one line:")
    q = doc.add_paragraph(caption + "; shortened here (when the section is longer than the excerpt); a link to the whole Act on India Code.")
    q.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph("When the Act is loaded but nothing from it reached the answer (nothing close enough was found):")
    q = doc.add_paragraph(notice_unused)
    q.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph("To the answer model, in addition to the rules it always has: use only the sections it is given for anything it says about the Act; name each with the Act shown in its label, not as a BNS or BNSS section; do not state another section number, punishment, procedure or route of appeal from memory; "
                      "if the question needs a provision that is not among them, say so in one sentence and say that the full text is on India Code; a provision missing from the material does not mean the Act has none; where a section's words "
                      "and the person's facts do not clearly match, say so instead of stretching the section.")
    p = doc.add_paragraph()
    p.add_run("Is this fair? Anything to add or strike? ").bold = True
    p.add_run("____________________________________________________________________________")

    doc.add_heading("How the text was checked, and what is not on the shelf", level=2)
    doc.add_paragraph("The check is a comparison of two texts that India Code publishes: its database and its PDF of the same Act. It cannot find an error that both share, and it cannot tell whether either is the law today. "
                      "Run on all five Acts it compared about 54,000 words and found nine places where the database was wrong (two doubled words, two lost apostrophes, a lost proviso, a misplaced phrase, three headings), which the loader now corrects against the PDF. "
                      "Everything else it found is listed in the loader's manifest next to what the PDF reads.")
    doc.add_paragraph("What is NOT loaded: the Schedules of POCSO (the list of armed forces and security forces that section 2(c) refers to), the PMLA (the list of scheduled offences, Parts A to C, on which \"scheduled offence\" in section 2 depends), the UAPA (the organisations and individuals listed) and the SC/ST Act (the offences section 3(2)(va) refers to); "
                      "the State or Union Territory amendments inside PC Act 17A (Jammu and Kashmir and Ladakh), POCSO 42 (Arunachal Pradesh) and SC/ST 14 (Karnataka); the Rules made under any of the Acts.")

    doc.add_heading("Three questions only you can answer", level=2)
    qs = ["1. Everyday words that the statute never uses. Example already in: \"anticipatory bail\" means sections 18 and 18A of the SC/ST Act (they say only \"section 438 of the Code shall not apply\"). "
          "What others should we add, for any of the five Acts (up to ten)? Write the everyday phrase and the section(s) it means.",
          "2. Provisions the text above does not show as they stand: an amendment not yet on India Code, a provision struck down or read down by a court, an ordinance. Which, if any, in these five Acts? And: should the Schedules (see above) be loaded, which of them first?",
          "3. The five questions people most often ask about each Act (in their own words). We will use them as the next test of the search."]
    for text in qs:
        doc.add_paragraph(text)
        doc.add_paragraph("____________________________________________________________________________")
        doc.add_paragraph("____________________________________________________________________________")
    doc.add_paragraph("Please send the sheet back (a photo of a marked printout is fine). Claude records only what you confirm; an unconfirmed Act stays \"not yet read by our reviewer\".")
    doc.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--recourse", required=True, help="a Recourse checkout that has act_shelf/")
    ap.add_argument("--out", default=os.path.join(HERE, "ACT_SHELF_SHEET.docx"))
    args = ap.parse_args()
    print("wrote", build(os.path.abspath(args.recourse), args.out))
