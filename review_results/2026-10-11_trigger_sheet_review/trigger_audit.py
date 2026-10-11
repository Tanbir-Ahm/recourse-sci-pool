#!/usr/bin/env python3
"""Audit of the owner's trigger workbook (Indian_Criminal_Law_Search_Trigger_Review_Layperson_Additions.xlsx) against our own index and the 30 used questions. No Voyage calls, no spend. PRIVATE.

  1. How common is each Act's STRONG passage marker in the index (passages, judgments)? A marker found in a large share of passages cannot discriminate.
  2. Which Acts would the question-side triggers (lawyer, layperson high-signal, layperson broad) fire on for each of the 30 used questions, with a simple stated matcher?
  3. Passage-level versus judgment-level labelling for the domestic-violence question R16 (does the right judgment get the Act's label? do the two judgments that crowded it out?).
"""
import collections
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "review_tools"))
sys.path.insert(0, ROOT)
import pool_runner as PR
import sci_index as si

XLSX = os.path.join(ROOT, "review_data", "ruler", "inputs", "2026-10-11_owner_trigger_workbook.xlsx")
wb = openpyxl.load_workbook(XLSX, data_only=True)
main = [r for r in wb["Revised Trigger Map"].iter_rows(min_row=3, values_only=True) if r[0] and r[1] and not str(r[0]).startswith("Act / statute")]
lay = {r[0]: r for r in wb["Layperson Trigger Additions"].iter_rows(min_row=3, values_only=True) if r[0] and r[1] and not str(r[0]).startswith("Act / statute")}
SHORT = ["PWDVA", "NDPS", "POCSO", "SC/ST", "NI Act", "PMLA", "UAPA", "DPA", "JJ Act", "BNS/IPC", "BNSS/CrPC", "BSA/IEA", "PC Act", "Arms Act", "IT Act", "ITPA", "PCMA", "PCR Act", "MV Act", "PDPP", "Companies Act", "CGST", "Muslim Women Act"]
assert len(main) == len(SHORT), (len(main), len(SHORT))

# STRONG passage markers (Tier 1: full title, acronym, section-linked), from the workbook's own STRONG lists; compiled case-insensitively, acronyms case-sensitively
S = lambda p: re.compile(p, re.I)
C = lambda p: re.compile(p)
MARK = {
    "PWDVA": [S(r"protection of women from domestic violence act"), C(r"\bPWDVA\b|\bPWDV Act\b|\bD\.?V\.? Act\b"), S(r"section\s+(12|18|19|20|22|23)\s+of\s+the\s+(D\.?V\.?|domestic violence)")],
    "NDPS": [C(r"\bNDPS\b|\bN\.D\.P\.S"), S(r"narcotic drugs and psychotropic substances act")],
    "POCSO": [C(r"\bPOCSO\b|\bP\.O\.C\.S\.O"), S(r"protection of children from sexual offences")],
    "SC/ST": [C(r"\bSC\s*/\s*ST\b|\bSC\s*&\s*ST\b"), S(r"prevention of atrocities\)? act")],
    "NI Act": [S(r"negotiable instruments act"), C(r"\bN\.?\s?I\.? Act\b")],
    "PMLA": [C(r"\bPMLA\b"), S(r"prevention of money[- ]laundering act")],
    "UAPA": [C(r"\bUAPA\b"), S(r"unlawful activities \(prevention\) act")],
    "DPA": [S(r"dowry prohibition act")],
    "JJ Act": [S(r"juvenile justice \(care and protection of children\) act"), C(r"\bJJ Act\b|\bJJB\b"), S(r"juvenile justice board|child welfare committee")],
    "BNS/IPC": [S(r"bharatiya nyaya sanhita"), C(r"\bBNS\b|\bIPC\b"), S(r"indian penal code")],
    "BNSS/CrPC": [S(r"bharatiya nagarik suraksha sanhita"), C(r"\bBNSS\b|\bCr\.?\s?P\.?\s?C\b"), S(r"code of criminal procedure")],
    "BSA/IEA": [S(r"bharatiya sakshya adhiniyam"), C(r"\bBSA\b"), S(r"evidence act")],
    "PC Act": [S(r"prevention of corruption act"), C(r"\bPC Act\b")],
    "Arms Act": [S(r"\barms act\b")],
    "IT Act": [S(r"information technology act"), C(r"\bIT Act\b")],
    "ITPA": [S(r"immoral traffic \(prevention\) act"), C(r"\bITPA\b")],
    "PCMA": [S(r"prohibition of child marriage act"), C(r"\bPCMA\b")],
    "PCR Act": [S(r"protection of civil rights act"), C(r"\bPCR Act\b")],
    "MV Act": [S(r"motor vehicles act"), C(r"\bMV Act\b")],
    "PDPP": [S(r"prevention of damage to public property act"), C(r"\bPDPP\b")],
    "Companies Act": [S(r"companies act")],
    "CGST": [S(r"central goods and services tax act"), C(r"\bCGST\b")],
    "Muslim Women Act": [S(r"muslim women \(protection of rights on marriage\) act")],
}

head = PR.index_head(PR.DEFAULT_INDEX)
metas, M, _ = si._load_index(PR.DEFAULT_INDEX)
for d in PR.FINAL_SIDE_INDEXES:
    metas, M, _, _ = PR.load_supplement(os.path.join(ROOT, d), metas, M, head)
N = len(metas)
docs_total = len({m["doc"] for m in metas})
per_doc = {a: collections.Counter() for a in MARK}
for m in metas:
    t = m["text"]
    for a, pats in MARK.items():
        if any(p.search(t) for p in pats):
            per_doc[a][m["doc"]] += 1
print(f"INDEX: {N:,} passages, {docs_total:,} judgments\n\n1. HOW COMMON IS EACH ACT'S STRONG PASSAGE MARKER")
print(f"{'Act':18} {'passages':>9} {'% of all':>9} {'judgments':>10} {'% of all':>9} {'judgments with 3+ marked passages':>34}")
audit = {}
for a in SHORT:
    p = sum(per_doc[a].values()); j = len(per_doc[a]); j3 = sum(1 for v in per_doc[a].values() if v >= 3)
    audit[a] = {"passages": p, "judgments": j, "judgments_3plus": j3}
    print(f"{a:18} {p:9,} {100 * p / N:8.2f}% {j:10,} {100 * j / docs_total:8.2f}% {j3:34,}")

# 2. question-side triggers
STOP = set("the a an and or of to in on for is are was were be been can could would should must may might will shall do does did not that this with from by as at it its who what when how which if any all have has had than then them they their there here about into over under after before also only such out up down off again further once".split()) | {"case", "police", "under", "act", "section", "me", "my", "i", "someone", "person"}


def stem(w):
    w = w.lower()
    for suf in ("ing", "ed", "es", "s"):
        if w.endswith(suf) and len(w) - len(suf) >= 3:
            return w[: -len(suf)]
    return w


def toks(s):
    return [stem(w) for w in re.findall(r"[A-Za-z0-9/']+", s) if w.lower() not in STOP and len(w) >= 3]


def phrases(cell):
    out = []
    for p in re.split(r";|\n", cell or ""):
        p = re.sub(r"\(.*?\)", "", p).strip()
        if p:
            out.append(p)
    return out


def fires(phrase, qtoks):
    pt = [t for t in toks(phrase)]
    if not pt:
        return False
    hit = sum(1 for t in set(pt) if t in qtoks)
    if len(set(pt)) == 1:
        return hit == 1
    return hit >= 2 and hit / len(set(pt)) >= 0.6


tun = json.load(open(os.path.join(ROOT, "review_results/2026-10-10_real30_rerank_test_pack_v2/pool_with_rerank_arms.json"), encoding="utf-8"))["questions"]
sealed = json.load(open(os.path.join(ROOT, "review_results/2026-10-10_sealed_run/pool.json"), encoding="utf-8"))["questions"]
qs = {q["qid"]: q["question"] for q in tun + sealed}
trig = {}
for a, row in zip(SHORT, main):
    full = row[0]
    L = phrases(row[1])
    lr = lay.get(full)
    trig[a] = {"lawyer": L, "lay_high": phrases(lr[1]) if lr else [], "lay_broad": phrases(lr[2]) if lr else []}
print("\n2. WHICH ACTS WOULD THE TRIGGERS FIRE ON (matcher: stemmed content words; a phrase fires on >=2 shared words and >=60% of its words, or on its single word)")
fire_counts = collections.Counter()
fired = {}
for qid in sorted(qs):
    qt = set(toks(qs[qid]))
    row = {}
    for a, t in trig.items():
        for tier in ("lawyer", "lay_high", "lay_broad"):
            hits = [p for p in t[tier] if fires(p, qt)]
            if hits:
                row.setdefault(a, {})[tier] = hits[:2]
    fired[qid] = row
    for a in row:
        fire_counts[a] += 1
    print(f"\n{qid}: {qs[qid][:140]}")
    for a, tiers in row.items():
        print("    ", a, {k: [h[:40] for h in v] for k, v in tiers.items()})
print("\nQuestions on which each Act fired (of 30):", dict(fire_counts.most_common()))

# 3. R16: passage-level versus judgment-level labelling
key = json.load(open(os.path.join(ROOT, "review_data/ruler/real30_key_v1.json")))["questions"]["R16"]
kd = set(key["key_docs"])
print("\n3. R16 (domestic violence): which judgments carry the PWDVA label?")
pw = per_doc["PWDVA"]
print("   key judgment(s)", sorted(kd), "-> marked passages:", {d: pw.get(d, 0) for d in kd}, "| key judgment passages total:", sum(1 for m in metas if m["doc"] in kd))
for d in (30017, 7544, 37423, 43646, 34852, 38403, 1000001385):
    print(f"   judgment {d}: PWDVA-marked passages {pw.get(d, 0)} of {sum(1 for m in metas if m['doc'] == d)}")
print("   judgments with 3+ PWDVA-marked passages:", sum(1 for v in pw.values() if v >= 3), "of", len(pw))
json.dump({"index_passages": N, "judgments": docs_total, "marker_audit": audit, "fired": fired}, open(os.path.join(HERE, "trigger_audit.json"), "w"), indent=1, default=str)
