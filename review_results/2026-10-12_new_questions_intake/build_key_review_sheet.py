"""Builds Key_Review_Sheet_for_Mr_Advocate.xlsx: the writer's key, side by side with what Claude found in the index and in the judgments' indexed text, and the questions Mr. Advocate must answer before sealing.
PRIVATE: contains the answer key. It goes to Mr. Advocate only (never to Voyage, ChatGPT, Gemini, J or S). Run from anywhere."""
import os, re, sys, json
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "review_tools")); sys.path.insert(0, ROOT); sys.path.insert(0, HERE)
import pool_runner as PR, sci_index as si
from authorities import A
from adjudication import ADJ, STALE
from notes import CHECK, FLAGS, OPEN_POINTS
SRC = os.path.join(ROOT, "review_data/ruler/inputs/2026-10-12_Fix_Test_New_Questions_N01-N30_as_received.xlsx")
OUT = os.path.join(HERE, "Key_Review_Sheet_for_Mr_Advocate.xlsx")
head = PR.index_head(PR.DEFAULT_INDEX)
metas, M, _ = si._load_index(PR.DEFAULT_INDEX)
for d in PR.FINAL_SIDE_INDEXES:
    metas, M, _, _ = PR.load_supplement(os.path.join(ROOT, d), metas, M, head)
title = PR.titles([PR.DEFAULT_INDEX] + [os.path.join(ROOT, d) for d in PR.FINAL_SIDE_INDEXES], metas)
first = {}
for m in metas:
    first.setdefault(m["doc"], m["text"])
def name_of(d):
    t = title(d) or ""
    if t.startswith(("IN THE SUPREME", "CASE NO", "PETITIONER", "Judgment")) or len(t) < 8:
        t = re.sub(r"\s+", " ", first.get(d, ""))[:150]
    return re.sub(r"\s+", " ", t)[:150]
wb0 = openpyxl.load_workbook(SRC)
Q = {r[0]: dict(zip([c.value for c in wb0["Questions"][1]], r)) for r in wb0["Questions"].iter_rows(min_row=2, values_only=True)}
K = {r[0]: dict(zip([c.value for c in wb0["Key"][1]], r)) for r in wb0["Key"].iter_rows(min_row=2, values_only=True)}
per = {}
for q, role, label, *_ in A:
    st, ds = ADJ[label]
    per.setdefault(q, []).append((role, label, st, ds))
def status_text(q):
    rows = per.get(q, [])
    if not rows:
        return "no authority named (control)"
    out = []
    for role, label, st, ds in rows:
        tag = {"IN": "IN the index", "LEFT_OUT": "in the collection but LEFT OUT of the index", "ABSENT": "NOT in the index"}[st]
        r = {"primary": "PRIMARY", "acceptable": "acceptable", "trap_old": "older ruling", "note": "note"}[role]
        out.append(f"{r}: {label} - {tag}" + (f" (doc {', '.join(map(str, ds))})" if st == "IN" else ""))
    return "\n".join(out)
wb = openpyxl.Workbook()
ws0 = wb.active; ws0.title = "Read me"
thin = Side(style="thin", color="BBBBBB")
hfill = PatternFill("solid", fgColor="DDE6F0"); gfill = PatternFill("solid", fgColor="E2F0D9"); rfill = PatternFill("solid", fgColor="FBE4E4")
wrap = Alignment(wrap_text=True, vertical="top")
readme = [
    ("Key review sheet for Mr. Advocate - new test questions N01 to N30 and reserves NR1 to NR3", True),
    ("PRIVATE. This sheet contains the answer key. It is for Mr. Advocate only. It must never go to Voyage, ChatGPT, Gemini, J or S.", False),
    ("", False),
    ("What this is. The question writer's workbook (11 October 2026) was checked by Claude on 12 October against the rules and the search index. Nothing has been searched and nothing is sealed. This sheet puts the writer's key next to what Claude could verify, so that the legal review can be quick and can concentrate on what only a lawyer can decide.", False),
    ("What you are asked to do. For each row of 'Key review' fill the four green columns: (1) is the expected answer correct and current? (2) is the primary authority the best one, and is the list of other acceptable authorities complete? (3) is the older-ruling / trap entry right (for traps: was the older ruling really overruled, modified or replaced)? (4) any correction. Then answer the 'Open points' sheet.", False),
    ("What the index columns mean. 'IN the index' means the judgment's text is in the collection the search runs over, so a search can return it. 'NOT in the index' means no search can ever return it; the question can then be answered only through the other authorities. 'Document no.' is the index number of the judgment, so that its passages can be read.", False),
    ("What Claude's check means. Claude read passages of the judgments that are in the index and noted whether they contain the words the key relies on. That is evidence that the words are there, NOT a legal opinion that the judgment decides the point. Where Claude found a mismatch it says so.", False),
    ("Rows shaded red are being returned to the writer for replacement (N02, N12, N15): you need not review them. Rows N28 to N30 are controls (no case-law answer is expected). NR1 to NR3 are reserves.", False),
    ("The writer is an independent legal person who used an AI tool to find the citations and cases (owner, 12 October). Treat every citation as unverified until you have checked it exists as cited.", False),
    ("Everything in the project is 'LLM-judged', not lawyer-validated; your review is the first lawyer's check on this key.", False),
]
for i, (t, b) in enumerate(readme, 1):
    c = ws0.cell(i, 1, t); c.alignment = wrap; c.font = Font(bold=b, size=13 if b else 11)
ws0.column_dimensions["A"].width = 130
for i in range(1, len(readme) + 1):
    ws0.row_dimensions[i].height = 18 if i in (1, 3) else 62
ws = wb.create_sheet("Key review")
cols = ["Id", "Domain", "Question (for reference)", "Expected answer (writer)", "Primary authority (writer)", "Other acceptable authorities (writer)", "Older ruling / trap (writer)", "Code that applies (writer)", "Writer's own 'To verify'", "Writer's Conf",
        "Index status of each authority named (Claude)", "Claude's check against the judgments' indexed text", "Claude's flags", "MR. ADVOCATE: expected answer correct and current? (Y / N / correct it)", "MR. ADVOCATE: primary and acceptable authorities right and complete? (Y / N / add or remove)", "MR. ADVOCATE: trap entry right? (Y / N / n.a.)", "MR. ADVOCATE: comments and corrections"]
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True); c.fill = gfill if c.value.startswith("MR.") else hfill; c.alignment = wrap
for qid in list(Q):
    q, k = Q[qid], K[qid]
    ws.append([qid, q["Domain (A to K)"] + (f" ({q['Control type (K1/K2/K3, domain K only)']})" if q.get("Control type (K1/K2/K3, domain K only)") else ""), q["Question text"], k["Expected answer"], k["Primary authority"], k["Other acceptable authorities"],
               k["Older ruling a search may wrongly show (traps) with what changed and the current authority"], k["Code that applies and why"], k["To verify"], k["Conf"], status_text(qid), CHECK.get(qid, ""), FLAGS.get(qid, ""), None, None, None, None])
    r = ws.max_row
    for c in ws[r]:
        c.alignment = wrap
        if c.column >= 14:
            c.fill = gfill
        if qid in ("N02", "N12", "N15"):
            if c.column <= 13:
                c.fill = rfill
widths = [6, 8, 42, 55, 34, 44, 44, 34, 40, 6, 52, 60, 40, 20, 24, 18, 36]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "D2"
wa = wb.create_sheet("Authorities in the index")
wa.append(["Question", "Role in the key", "Authority as named in the key", "Status", "Index document no.", "Title or opening text of that document in the index"])
for c in wa[1]:
    c.font = Font(bold=True); c.fill = hfill
for q, role, label, toks, year, slack in A:
    st, ds = ADJ[label]
    wa.append([q, {"primary": "primary", "acceptable": "other acceptable", "trap_old": "older ruling (trap)", "note": "mentioned, not a counted trap"}[role], label, {"IN": "in the index", "LEFT_OUT": "in the collection, left out of the index", "ABSENT": "not in the index"}[st],
               ", ".join(map(str, ds)) if ds else "", "; ".join(name_of(d) for d in ds[:2])])
for i, w in enumerate([9, 22, 62, 34, 20, 100], 1):
    wa.column_dimensions[get_column_letter(i)].width = w
for row in wa.iter_rows(min_row=2):
    for c in row:
        c.alignment = wrap
wo = wb.create_sheet("Open points")
wo.append(["No.", "Question", "Where", "MR. ADVOCATE: answer"])
for c in wo[1]:
    c.font = Font(bold=True); c.fill = gfill if c.value.startswith("MR.") else hfill
for i, (where, text) in enumerate(OPEN_POINTS, 1):
    wo.append([i, text, where, None])
    for c in wo[wo.max_row]:
        c.alignment = wrap
    wo.cell(wo.max_row, 4).fill = gfill
for i, w in enumerate([5, 110, 14, 60], 1):
    wo.column_dimensions[get_column_letter(i)].width = w
wb.save(OUT)
print("wrote", OUT, "rows", ws.max_row - 1)
