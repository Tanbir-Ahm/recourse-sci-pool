"""Builds Key_Review_Sheet_for_Mr_Advocate_v2.xlsx: the SECOND version of the sheet, after the writer's replacements (round 2, 12 October 2026) and the owner's decisions (N12 dropped, N02 kept as a soft overlap).
Same layout as the first version, plus a status column (NEW / CHANGED / DROPPED / unchanged), the writer's round-2 key, the mechanical citation pre-check, and an updated list of open points.
PRIVATE: contains the answer key. It goes to Mr. Advocate only (never to Voyage, ChatGPT, Gemini, J or S)."""
import json, os, re, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__)); INTAKE = os.path.normpath(os.path.join(HERE, "..")); ROOT = os.path.normpath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "review_tools")); sys.path.insert(0, ROOT); sys.path.insert(0, INTAKE); sys.path.insert(0, HERE)
import pool_runner as PR, sci_index as si
from notes_round2 import CHECK, FLAGS, OPEN_POINTS
SRC1 = os.path.join(ROOT, "review_data/ruler/inputs/2026-10-12_Fix_Test_New_Questions_N01-N30_as_received.xlsx")
SRC2 = os.path.join(ROOT, "review_data/ruler/inputs/2026-10-12_Fix_Test_New_Questions_N01-N30_round2_as_received.xlsx")
OUT = os.path.join(HERE, "Key_Review_Sheet_for_Mr_Advocate_v2.xlsx")
AUTH = json.load(open(os.path.join(HERE, "authorities_final_proposed.json"), encoding="utf-8"))
CITP = os.path.join(HERE, "citations_check_round2.json")
CIT = json.load(open(CITP, encoding="utf-8"))["rows"] if os.path.exists(CITP) else []
NEW, DROPPED = {"N02", "N15"}, {"N12"}

head = PR.index_head(PR.DEFAULT_INDEX)
metas, M, _ = si._load_index(PR.DEFAULT_INDEX)
extra_titles = {}
for d in PR.FINAL_SIDE_INDEXES:
    metas, M, _i, extra = PR.load_supplement(os.path.join(ROOT, d), metas, M, head)
    extra_titles.update(extra)
title = PR.titles([PR.DEFAULT_INDEX] + [os.path.join(ROOT, d) for d in PR.FINAL_SIDE_INDEXES], metas, extra_titles or None)
first = {}
for m in metas:
    first.setdefault(m["doc"], m["text"])


def name_of(d):
    t = title(d) or ""
    if t.startswith(("IN THE SUPREME", "CASE NO", "PETITIONER", "Judgment")) or len(t) < 8:
        t = re.sub(r"\s+", " ", first.get(d, ""))[:150]
    return re.sub(r"\s+", " ", t)[:150]


def sheet(path, name):
    ws = openpyxl.load_workbook(path)[name]
    rows = list(ws.iter_rows(values_only=True)); hdr = list(rows[0])
    return hdr, {r[0]: dict(zip(hdr, r)) for r in rows[1:]}


hq, Q = sheet(SRC2, "Questions"); hk, K = sheet(SRC2, "Key"); _, K1 = sheet(SRC1, "Key")
changed = {}
for qid, k in K.items():
    if qid in K1 and qid not in NEW and qid not in DROPPED:
        diff = [h for h in hk if h != "Id" and (k.get(h) or "") != (K1[qid].get(h) or "")]
        if diff:
            changed[qid] = diff
per = {}
for e in AUTH:
    per.setdefault(e["qid"], []).append(e)


def status_text(qid):
    rows = per.get(qid, [])
    if not rows:
        return "no authority named (control or dropped)"
    out = []
    for e in rows:
        tag = {"IN": "IN the index", "LEFT_OUT": "in the collection but LEFT OUT of the index", "ABSENT": "NOT in the index"}[e["status"]]
        r = {"primary": "PRIMARY", "acceptable": "acceptable", "trap_old": "older ruling", "note": "note"}[e["role"]]
        out.append(f"{r}: {e['label']} - {tag}" + (f" (doc {', '.join(map(str, e['docs']))})" if e["status"] == "IN" else "") + (f"  [Claude: {e['note']}]" if e.get("note") else ""))
    return "\n".join(out)


def version_status(qid):
    if qid in DROPPED:
        return "DROPPED by the owner (12 Oct). Do not review."
    if qid in NEW:
        return "NEW in this version: replaces the first " + qid + ". Please review in full."
    if qid in changed:
        return "CHANGED since the first version: " + "; ".join(changed[qid]) + ". If you already reviewed the first version, re-read only these cells."
    return "Unchanged since the first version."


thin = None
hfill = PatternFill("solid", fgColor="DDE6F0"); gfill = PatternFill("solid", fgColor="E2F0D9"); yfill = PatternFill("solid", fgColor="FFF2CC"); rfill = PatternFill("solid", fgColor="FBE4E4"); grey = PatternFill("solid", fgColor="E7E6E6"); ofill = PatternFill("solid", fgColor="FCE4D6")
wrap = Alignment(wrap_text=True, vertical="top")
wb = openpyxl.Workbook()
ws0 = wb.active; ws0.title = "Read me"
readme = [
    ("Key review sheet for Mr. Advocate - new test questions, SECOND VERSION (12 October 2026, after the writer's replacements)", True),
    ("PRIVATE. This sheet contains the answer key. It is for Mr. Advocate only. It must never go to Voyage, ChatGPT, Gemini, J or S.", False),
    ("", False),
    ("What changed since the first version of this sheet. (1) The writer replaced two questions: N02 and N15 are NEW (yellow rows). (2) N12 was DROPPED by the owner because it repeated an earlier question (grey row; nothing to review). (3) In eight rows the writer changed the key text (orange rows: N10, N14, N17, N19, N20, N24, N29, NR2); the column 'Status in this version' says which cells. (4) A new sheet, 'Citations (pre-check)', lists every citation in the key and what Claude could and could not confirm from other judgments in the index. (5) The open points are updated.", False),
    ("If you already reviewed the first version: you only need to look at the yellow and orange rows, the new open points (marked NEW), and the 'Citations (pre-check)' sheet. The unchanged rows stand.", False),
    ("What this is. The question writer's workbook was checked by Claude against the rules and the search index. Nothing has been searched and nothing is sealed. This sheet puts the writer's key next to what Claude could verify, so that the legal review can concentrate on what only a lawyer can decide.", False),
    ("What you are asked to do. For each row of 'Key review' fill the four green columns: (1) is the expected answer correct and current? (2) is the primary authority the best one, and is the list of other acceptable authorities complete? (3) is the older-ruling / trap entry right (for traps: was the older ruling really overruled, modified or replaced)? (4) any correction. Then answer the 'Open points' sheet and fill the last column of 'Citations (pre-check)'.", False),
    ("What the index columns mean. 'IN the index' means the judgment's text is in the collection the search runs over, so a search can return it. 'NOT in the index' means no search can ever return it; the question can then be answered only through the other authorities. 'Document no.' is the index number of the judgment, so that its passages can be read.", False),
    ("What Claude's check means. Claude read passages of the judgments that are in the index and noted whether they contain the words the key relies on. That is evidence that the words are there, NOT a legal opinion that the judgment decides the point. Where Claude found a mismatch it says so. The 'Citations (pre-check)' sheet is mechanical: it shows whether OTHER judgments in the index repeat the key's citation next to the same case name. A case nobody else in the index cites cannot be checked that way, and a citation other judgments repeat is only as good as they are.", False),
    ("Who wrote what (owner, 12 October): the questions and the key were written by a person; an AI tool (Claude) helped only to find the citations and cases. AI-found citations are the usual place for wrong or invented cases, so please treat every citation as unverified until you have checked that it exists as cited. In the pre-check Claude found citations that look wrong (for example Rajnesh v. Neha, Vinod Katara, Mahmood Farooqui) and one authority that does not say what the key says (Shyamlal Devda); the rows say so.", False),
    ("Everything in the project is 'LLM-judged', not lawyer-validated; your review is the first lawyer's check on this key.", False),
]
for i, (t, b) in enumerate(readme, 1):
    c = ws0.cell(i, 1, t); c.alignment = wrap; c.font = Font(bold=b, size=13 if b else 11)
ws0.column_dimensions["A"].width = 140
for i in range(1, len(readme) + 1):
    ws0.row_dimensions[i].height = 20 if i in (1, 2, 3) else 78
ws = wb.create_sheet("Key review")
cols = ["Id", "Status in this version", "Domain", "Question (for reference)", "Expected answer (writer)", "Primary authority (writer)", "Other acceptable authorities (writer)", "Older ruling / trap (writer)", "Code that applies (writer)", "Writer's own 'To verify'", "Writer's Conf",
        "Index status of each authority named (Claude)", "Claude's check against the judgments' indexed text", "Claude's flags", "MR. ADVOCATE: expected answer correct and current? (Y / N / correct it)", "MR. ADVOCATE: primary and acceptable authorities right and complete? (Y / N / add or remove)", "MR. ADVOCATE: trap entry right? (Y / N / n.a.)", "MR. ADVOCATE: comments and corrections"]
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True); c.fill = gfill if c.value.startswith("MR.") else hfill; c.alignment = wrap
for qid in Q:
    q, k = Q[qid], K[qid]
    ws.append([qid, version_status(qid), q["Domain (A to K)"] + (f" ({q['Control type (K1/K2/K3, domain K only)']})" if q.get("Control type (K1/K2/K3, domain K only)") else ""), q["Question text"], k["Expected answer"], k["Primary authority"], k["Other acceptable authorities"],
               k["Older ruling a search may wrongly show (traps) with what changed and the current authority"], k["Code that applies and why"], k["To verify"], k["Conf"], status_text(qid), CHECK.get(qid, ""), FLAGS.get(qid, ""), None, None, None, None])
    r = ws.max_row
    for c in ws[r]:
        c.alignment = wrap
        if c.column >= 15 and qid not in DROPPED:
            c.fill = gfill
        if qid in NEW and c.column <= 14:
            c.fill = yfill
        elif qid in changed and c.column <= 2:
            c.fill = ofill
        elif qid in DROPPED:
            c.fill = grey
widths = [6, 34, 8, 42, 55, 34, 44, 44, 34, 40, 6, 52, 60, 44, 20, 24, 18, 36]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "E2"
wa = wb.create_sheet("Authorities in the index")
wa.append(["Question", "Role in the key", "Authority as named in the key", "Status", "Index document no.", "Title or opening text of that document in the index", "Claude's note"])
for c in wa[1]:
    c.font = Font(bold=True); c.fill = hfill
for e in AUTH:
    wa.append([e["qid"], {"primary": "primary", "acceptable": "other acceptable", "trap_old": "older ruling (trap)", "note": "mentioned, not a counted trap"}[e["role"]], e["label"], {"IN": "in the index", "LEFT_OUT": "in the collection, left out of the index", "ABSENT": "not in the index"}[e["status"]],
               ", ".join(map(str, e["docs"])) if e["docs"] else "", "; ".join(name_of(d) for d in e["docs"][:2]), e.get("note", "")])
    if e.get("note"):
        for c in wa[wa.max_row]:
            c.fill = ofill
for i, w in enumerate([9, 22, 62, 34, 20, 100, 80], 1):
    wa.column_dimensions[get_column_letter(i)].width = w
for row in wa.iter_rows(min_row=2):
    for c in row:
        c.alignment = wrap
wc = wb.create_sheet("Citations (pre-check)")
wc.append(["Question", "Authority", "Citation as written in the key", "Result", "What other judgments in the index say", "Index document(s) of the judgment itself", "MR. ADVOCATE: exists as cited? (Y / N / correct it)"])
for c in wc[1]:
    c.font = Font(bold=True); c.fill = gfill if c.value.startswith("MR.") else hfill; c.alignment = wrap
SHORT = {"confirmed": "Repeated by other judgments", "confirmed-fn": "Cited in a footnote of a judgment that names the case", "footnote-only": "Citation seen in a footnote; case name not read there", "own": "In the judgment's own text", "parallel": "Parallel citation to check", "alt-differs": "CHECK: other judgments cite it differently", "cite-elsewhere": "CHECK: the citation appears next to other case names",
         "unconfirmed": "Not confirmable here", "not-checkable": "Not checkable here", "no-cite": "No citation in the key", "no-cite-alt": "No citation in the key"}
for r in CIT:
    base = r["code"].split("+")[0]
    label = SHORT.get(base, base) + ("; YEAR DOES NOT FIT" if "+year" in r["code"] else "")
    wc.append([r["q"], r["label"], r["key_cite"] or "(none given)", label, r["verdict"], ", ".join(map(str, r["docs"])), None])
    row = wc[wc.max_row]
    flag = r["code"].startswith(("alt-differs", "cite-elsewhere")) or "+year" in r["code"]
    for c in row:
        c.alignment = wrap
        if flag and c.column <= 6:
            c.fill = rfill
        elif c.column == 7:
            c.fill = gfill
        elif base in ("unconfirmed", "not-checkable", "no-cite", "no-cite-alt", "parallel", "footnote-only") and c.column <= 6:
            c.fill = yfill
for i, w in enumerate([9, 52, 24, 34, 110, 20, 34], 1):
    wc.column_dimensions[get_column_letter(i)].width = w
wc.freeze_panes = "C2"
wo = wb.create_sheet("Open points")
wo.append(["No.", "Question", "Where", "MR. ADVOCATE: answer"])
for c in wo[1]:
    c.font = Font(bold=True); c.fill = gfill if c.value.startswith("MR.") else hfill
for i, (where, text) in enumerate(OPEN_POINTS, 1):
    wo.append([i, text, where, None])
    for c in wo[wo.max_row]:
        c.alignment = wrap
    wo.cell(wo.max_row, 4).fill = gfill
    if "NEW" in where:
        wo.cell(wo.max_row, 3).fill = yfill
for i, w in enumerate([5, 110, 14, 60], 1):
    wo.column_dimensions[get_column_letter(i)].width = w
wb.save(OUT)
print("wrote", OUT, "| key rows", ws.max_row - 1, "| authorities", wa.max_row - 1, "| citations", wc.max_row - 1, "| open points", wo.max_row - 1, "| changed rows:", sorted(changed))
