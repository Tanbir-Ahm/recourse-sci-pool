"""Builds Return_Note_for_the_Question_Writer.docx (python-docx). PRIVATE (it names the earlier test topics; it contains no answer key and no result)."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
HERE = os.path.dirname(os.path.abspath(__file__))
d = Document()
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
def h(t, lvl=1):
    p = d.add_heading(t, level=lvl)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
def para(t, bold=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; p.paragraph_format.space_after = Pt(6); return p
def bullets(items):
    for t in items:
        p = d.add_paragraph(style="List Bullet"); p.add_run(t); p.paragraph_format.space_after = Pt(2)
def table(rows, widths):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(txt); r.font.size = Pt(10); r.bold = (i == 0)
    d.add_paragraph()
h("Note to the question writer: three questions to replace, and what to fix before sealing", 0)
para("12 October 2026. PRIVATE. Sent through the owner. This note contains no answer key and no search result.", bold=True)
h("1. Where things stand")
bullets([
    "Your workbook passed the format and counting checks (32 of 36). Your freeze hash is correct: we recomputed it from your 33 question texts and it matches 7a6963d1159963f5611c4df2fe4d9840ce7f37423d360e656e3afec83e1c99a3.",
    "24 of the 27 scored questions, the three controls and the three reserves stay exactly as you froze them. Please do not change a word of them.",
    "Three questions repeat topics that were already used in earlier tests and must be replaced (the rules, section 11.2: overlaps are returned). A few items in the key and the declaration need fixing.",
])
h("2. Replace these three (rewrite once; keep the ids)")
table([
    ["Id and domain", "Why it is returned", "What the replacement must do"],
    ["N02 (domain A)", "Same legal issue and same primary authority (Vijaysinh Chandubha Jadeja, section 50 NDPS: the duty to tell the person of the right to be searched before a Gazetted Officer or Magistrate) as an earlier test topic.", "A new topic. Do not use section 50 NDPS, the topics of the earlier 30, or the 16 earlier topics listed in section 5. If you can, make it a look-alike under one of the 13 listed Acts that does not name the Act: that keeps the mix rules safe."],
    ["N12 (domain E)", "Same legal issue and same primary authority (Md. Ibrahim v. State of Bihar: a civil dispute over land turned into a forgery case) as an earlier test topic.", "A new topic in domain E (quashing, cheating, fraud, or online fraud), different from the earlier topics."],
    ["N15 (domain F)", "Its primary authority, Satish Chander Ahuja v. Sneha Ahuja, is not in our index and is not to be used. The only in-index authority for the issue (Prabha Tyagi v. Kamlesh Devi) was an earlier test topic (domestic violence after leaving the shared household). The older ruling (S.R. Batra) is not in the index either, so the trap cannot be measured.", "A question under the Protection of Women from Domestic Violence Act, as the twin of N17 (the same surface facts as N17: a wife forced out of her in-laws' house) on a different point from 'can she stay in the shared household after leaving'. Its primary authority must be a Supreme Court judgment in our index (section 4). It need not be a currency trap."],
], [3.0, 7.0, 6.8])
para("Each replacement must follow every rule in the question guidelines: one answerable core, no names or case numbers, the tags, and a complete key (primary authority, every other acceptable authority, code that applies, how verified, Conf, To verify).")
h("3. What the set still needs after the replacements (please do not break these)")
bullets([
    "PWDVA must appear in at least 2 questions. After N15 is removed only N16 carries it, so N15's replacement must be a PWDVA question.",
    "At least 14 Act-concerning questions among the 27: after the removals there are exactly 14.",
    "Leave N08, N10, N23 and N24 untouched: they are the four currency traps whose older ruling is in our index, which is the minimum the rules allow.",
    "Keep the N15/N17 twin pair (both cross-referenced), so the set still has three twin pairs.",
])
h("4. Check the index before you write (names only)")
para("Of the 77 authorities named for N01 to N27, 58 are in our index. A replacement's primary authority must be in it. Send the owner the names you intend to use and Claude will check them against the index within minutes (no question text needs to be sent).")
para("Domestic-violence judgments that are in our index and have not been used in an earlier test: Juveria Abdul Majid Patni v. Atif Iqbal Mansoori (2014); Shyamlal Devda v. Parimala (2020); Rajnesh v. Neha (2020); Kiran Jyot Maini v. Anish Pramod Patel (2024); Kailashben Mahendrabhai Patel v. State of Maharashtra (2024); Geddam Jhansi v. State of Telangana (2025); Shaurabh Kumar Tripathi v. Vidhi Rawal (2025); Dhananjay Rathi v. Ruchika Rathi (2026). Not in the index: Hiral P. Harsora, Indra Sarma, Krishna Bhatacharjee, Satish Chander Ahuja, S. Vanitha, Independent Thought. This list says what can be searched, not what the answer is; the legal content and the key are yours to verify.")
h("5. Earlier test topics to avoid (titles only)")
para("1 Arnesh Kumar (arrest must be necessary); 2 Md. Ibrahim (civil dispute dressed as crime); 3 Nanda Gopalan (compromise, grievous hurt); 4 Rangappa (section 139 presumption); 5 Jadeja (section 50 NDPS compliance); 6 Noor Aga (NDPS reverse burden); 7 grounds of arrest in writing; 8 default bail, oral application and the court's duty to tell the accused; 9 relatives named in a 498A case with general allegations; 10 a failed commercial deal turned into cheating or breach of trust; 11 compounding where police caused custodial violence; 12 NDPS section 37 twin conditions; 13 NDPS informant also investigated; 14 NDPS section 67 statement is not a confession; 15 a signed blank cheque given as security (section 139); 16 domestic violence complaint after leaving the shared household.")
h("6. Fixes on the key and the declaration (all other rows)")
bullets([
    "N29: Conf must be G, because the verification rests on secondary sources only (rules, section 10).",
    "N10, N14, N17, N24 and NR2: say whether any other acceptable authority exists, or confirm that none does. N20: add modern authority on consent if it exists (only Mango Ram, 2000, and Gurmit Singh, 1996, are listed).",
    "N19: the primary authority (P. Yuvaprakash) concerns the age of a victim; the question concerns an accused. Choose the primary again or list both Yuvaprakash and Rishipal Singh Solanki.",
    "Declaration: add the line 'AI tools used and for what'. If an AI tool drafted any question or any part of a key, name the tool and say what it did (rules, section 3). Confirm that no wording was chosen to match or avoid the trigger words you read.",
    "Mr. Advocate will review the key; expect his corrections to come back through the owner.",
])
h("7. How to return it")
bullets([
    "The same workbook, the same three sheets and column names. The replacements keep the ids N02, N12 and N15; every other row must reproduce your existing text exactly.",
    "Recompute your freeze hash over all 33 rows (Id, a tab, the question text, one row per line, rows joined by a newline) and state it on the declaration sheet.",
    "Do not run the new questions through any search tool or chatbot, and do not show them to anyone but the owner.",
])
out = os.path.join(HERE, "Return_Note_for_the_Question_Writer.docx")
d.save(out)
print("wrote", out)
