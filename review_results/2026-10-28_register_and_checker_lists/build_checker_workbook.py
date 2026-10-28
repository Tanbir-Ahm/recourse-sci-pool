"""Builds Checker_List_and_Coverage_2026-10-28.xlsx: the list the invented-case check looks up, and which of the corpus's judgments are in it.
Run: python3 build_checker_workbook.py RECOURSE_DIR COVERAGE.json(from coverage_check.py) OUT.xlsx"""
import bisect, collections, datetime, json, os, re, sqlite3, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

rec, cov_path, out = sys.argv[1:4]
cov = json.load(open(cov_path, encoding="utf-8"))
db = sqlite3.connect("file:" + os.path.join(rec, "case_lookup_index.db") + "?mode=ro", uri=True)
LIST = db.execute("select case_id,title,petitioner,respondent,decision_date,citation,source from cases order by decision_date,title").fetchall()
STOP = set("the and of v vs versus state union india ors anr others another through smt shri sri mr mrs dr m s ltd pvt etc by lrs lr thr".split())
toks = lambda s: {w for w in re.findall(r"[a-z]{3,}", (s or "").lower()) if w not in STOP}
L = []
for cid, t, p, r, dt, cit, src in LIST:
    try:
        L.append((datetime.date.fromisoformat(dt), t, toks(t), cit))
    except Exception:
        pass
L.sort(key=lambda x: x[0]); dates = [x[0] for x in L]

PRESENT_OTHER = "In the list under another spelling or date; the checker's name match misses it"
ABSENT = "NOT in the list (decided before the list ends)"
AFTER = "NOT in the list: decided after the list ends (" + cov["list_end"] + ")"
rows = []
for r in cov["wide"]:
    st, close = r["status"], r.get("list_title")
    if st.startswith("in the list under another spelling"):
        st = PRESENT_OTHER
    elif st.startswith("NOT in the list (decided before"):
        st = ABSENT
        try:
            dd = datetime.date.fromisoformat(r["date"])
            lo = bisect.bisect_left(dates, dd - datetime.timedelta(days=150)); hi = bisect.bisect_right(dates, dd + datetime.timedelta(days=150))
            a = toks(r["name"]); best = (0, None)
            for dt, t, tk, cit in L[lo:hi]:
                if a and tk:
                    j = len(a & tk) / len(a | tk)
                    if j > best[0]:
                        best = (j, t + " (" + str(dt) + ")")
            if best[0] >= 0.6:
                st, close = PRESENT_OTHER, best[1]
        except Exception:
            pass
    elif st.startswith("NOT in the list: decided after"):
        st = AFTER
    elif st.startswith("found"):
        st = "Found by the checker"
    elif st.startswith("not a Supreme Court"):
        st = "Not applicable: a High Court judgment (the list is Supreme Court only)"
    rows.append([r["doc"], r.get("name"), r.get("date"), r.get("court"), r.get("citations"), st, close])
cnt = collections.Counter(x[5] for x in rows)
sc_named = sum(v for k, v in cnt.items() if k in ("Found by the checker", PRESENT_OTHER, ABSENT, AFTER))

HEAD = PatternFill("solid", fgColor="1F3864"); F = "Arial"
wb = Workbook()
def sheet(title, header, data, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in data:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return ws

# ---- library
LIB_NOTE = {
 "Deepa v S. Vijayalakshmi": ("High Court (Madras)", "not applicable: the list is Supreme Court only"),
 "Jatin Narendrabhai Sanghvi & Ors v State of Gujarat & Anr": ("High Court (Gujarat)", "not applicable"),
 "Johnson V.U. v State of Kerala": ("High Court (Kerala)", "not applicable"),
 "Malabar Gold and Diamond Limited v Union of India": ("High Court (Delhi)", "not applicable"),
 "Neelkanth Pharma Logistics Pvt. Ltd. v Union of India": ("High Court (Delhi)", "not applicable"),
 "Prakash Ranjan v State of Bihar": ("High Court (Patna)", "not applicable"),
 "Rakhi Mitra and Anr v State of West Bengal": ("High Court (Calcutta)", "not applicable"),
 "Sri Manjunath M P v State of Karnataka": ("High Court (Karnataka)", "not applicable"),
 "T. Manikadan v State (NCT of Delhi) & Anr": ("High Court (Delhi)", "not applicable"),
 "Viraj Chetan Shah v Union of India & Anr (& Connected Matters)": ("High Court (Bombay)", "not applicable"),
 "Lalita Kumari v Government of Uttar Pradesh": ("Supreme Court", "IN the list as 'LALITA KUMARI versus GOVT. OF U.P. AND ORS.' (2013-11-12) and 'GOVERNMENT OF U.P.' (2012-02-27); the checker's name match misses it ('Uttar Pradesh' vs 'U.P.')"),
 "Union of India v Md Nawaz Khan": ("Supreme Court", "IN the list as 'UNION OF INDIA THROUGH NARCOTICS CONTROL BUREAU, LUCKNOW versus MD. NAWAZ KHAN' (2021-09-22); the match misses it (first party spelt differently)"),
 "S. Vanitha v Deputy Commissioner, Bengaluru Urban District": ("Supreme Court", "IN the list as 'SMT. S. VANITHA versus THE DEPUTY COMMISSIONER, BENGALURU URBAN DISTRICT & ORS.' (2020-12-15); the match misses it ('Smt.')"),
 "Lokesh B.H v State of Karnataka": ("Supreme Court", "NOT in the list: 2026 INSC 784, decided after the list ends"),
 "Mathai Verghese v State of Kerala": ("Supreme Court", "NOT in the list (no title with these words found by hand)"),
 "Prakash Chimanlal Sheth v Jagruti Keyur Rajpopat": ("Supreme Court", "NOT in the list (2025 INSC 897; no title with these words found by hand)"),
 "Usha Chakraborty v State of West Bengal": ("Supreme Court", "NOT in the list (2023; no title with these words found by hand)"),
 "Youth Bar Association v Union of India": ("Supreme Court", "NOT in the list (2016; no title with these words found by hand)"),
}
lib = [[x["name"], x["citation"], "Found by the checker" if x["checker_says"] else "Not found by the checker",
        LIB_NOTE.get(x["name"], ("Supreme Court", ""))[0] if not x["checker_says"] else "Supreme Court (or found in the list)",
        LIB_NOTE.get(x["name"], ("", ""))[1], x["file"]] for x in cov["library"]]
lib.sort(key=lambda r: (r[2] != "Not found by the checker", r[0]))
lib_miss = [r for r in lib if r[2].startswith("Not")]
lib_hc = sum(1 for r in lib_miss if r[3].startswith("High"))
lib_sc_inlist = sum(1 for r in lib_miss if r[4].startswith("IN the list"))
lib_sc_out = len(lib_miss) - lib_hc - lib_sc_inlist

# ---- read me and summary
ws = wb.active; ws.title = "Read me"
txt = [
 "THE LIST THE INVENTED-CASE CHECK LOOKS UP, AND HOW MUCH OF THE CORPUS IS IN IT (built 2026-10-28)",
 "",
 "HOW THE CHECK WORKS (Recourse grounding_checks.py, check_cases; called from chat_assistant.py after the answer is written)",
 "1. It reads every case name the answer mentions (names written like 'X v. Y').",
 "2. A name PASSES if its distinctive party name appears in the MATERIAL the model was given for this answer: the question, the sections and judgments retrieved, the wider-search passages, the Act text, the library case names. This is the main test, and it does not use the list.",
 "3. Only a name that is NOT in the material goes to the LIST. The list is case_lookup_index.db: 38,026 Supreme Court titles, decided 1950-03-14 to " + cov["list_end"] + ". A name counts as 'in the list' only when a title has the same first party and a similar second party (at least half the words shared).",
 "4. Result: 'real_not_supplied' (in the list, but not among the sources) or 'not_found' (not in the list) or 'unknown' (could not look). Either way the answer gets a notice under it (when GROUNDING_CHECKS is warn, as on production): 'not among the sources we read ... do not rely on it'. 'not_found' adds 'we could not find it in our list of Supreme Court judgments'.",
 "5. In enforce mode (not on production) a case that is 'not_found' after a second try makes the answer be withheld. So a real case missing from the list would then block an answer: the list's gaps matter before enforce is ever switched on.",
 "",
 "IS EVERY JUDGMENT OF THE CORPUS IN THE LIST?  NO.",
 f"Wide index d8811a9 (the one on staging): {len(rows):,} judgments = {cnt.get('Found by the checker',0)+cnt.get(PRESENT_OTHER,0)+cnt.get(ABSENT,0)+cnt.get(AFTER,0)+sum(1 for r in rows if r[1] is None and r[3]=='Supreme Court of India'):,} Supreme Court + {sum(1 for r in rows if r[3]!='Supreme Court of India'):,} Gauhati High Court (the list cannot hold High Court judgments).",
 f"Of the {sc_named:,} named Supreme Court judgments: {cnt.get('Found by the checker',0):,} ({cnt.get('Found by the checker',0)/sc_named:.0%}) are found by the checker; {cnt.get(PRESENT_OTHER,0):,} are in the list but under a spelling the checker's name match does not recognise; {cnt.get(AFTER,0):,} were decided after the list ends; {cnt.get(ABSENT,0):,} ({cnt.get(ABSENT,0)/sc_named:.0%}) are not in the list at all (hand checks of five: four were really absent, one was present with a changed spelling and date).",
 f"Library (the 68 judgments Recourse holds as reviewed text): {len(lib)-len(lib_miss)} found; {len(lib_miss)} not found = {lib_hc} High Court judgments (not applicable), {lib_sc_inlist} Supreme Court judgments that ARE in the list but the name match misses, {lib_sc_out} Supreme Court judgments not in the list (one decided in 2026, four that the list lacks).",
 "",
 "WHY THIS MATTERS, SAID PLAINLY",
 "- It does NOT let an invented case through: a case that is not in the material is flagged whatever the list says. The list only decides the wording of the notice and, in enforce mode, whether the answer is withheld.",
 "- It DOES make the notice wrong for real cases: a real Supreme Court case that is missing from the list, or spelt differently, is called 'not found in our list', which is unfair to a genuine citation. High Court cases are always 'not found' because the list is Supreme Court only.",
 "- The list stops at " + cov["list_end"] + "; the corpus runs to October 2026.",
 "- 'Found by the checker' means the name was found IN THE LIST. It does not mean the case supports the point it is cited for; nothing here checks that (the parked support-and-fit plan).",
 "",
 "SHEETS",
 "Summary: the counts. Checker list (38,026): every title the check can find. Corpus judgments (19,255): each judgment of the wide index with its status and the closest list title. Library (68): the reviewed judgments.",
 "Method and limits: 'Found by the checker' is the checker's own function run on each name (grounding_checks.sc_list_has). 'Another spelling' was found by same-day first-party words (840) and by a looser words-overlap within 150 days (197). 'Not in the list' is the rest; a few may be present in a form neither test recognises. No lawyer has read this.",
]
for t in txt:
    ws.append([t])
ws.column_dimensions["A"].width = 190
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name=F, size=10, bold=bool(c.value) and c.value.isupper())

ws = wb.create_sheet("Summary")
ws.append(["Item", "Count"])
S = [("Titles in the checker's list", len(LIST)), ("List covers decisions from", LIST[0][4]), ("List covers decisions to", cov["list_end"]),
     ("Corpus judgments (wide index d8811a9)", len(rows)),
     ("  Supreme Court, named", sc_named), ("    Found by the checker", cnt.get("Found by the checker", 0)),
     ("    In the list under another spelling (checker misses)", cnt.get(PRESENT_OTHER, 0)), ("    Decided after the list ends", cnt.get(AFTER, 0)),
     ("    Not in the list", cnt.get(ABSENT, 0)), ("  Gauhati High Court (list is Supreme Court only)", sum(1 for r in rows if r[3] != "Supreme Court of India")),
     ("  Without a name in the names table", sum(1 for r in rows if r[1] is None)),
     ("Library judgments", len(lib)), ("  Found by the checker", len(lib) - len(lib_miss)), ("  Not found: High Court", lib_hc), ("  Not found: in the list, name match misses", lib_sc_inlist), ("  Not found: not in the list", lib_sc_out)]
for a, b in S:
    ws.append([a, b])
ws.column_dimensions["A"].width = 62; ws.column_dimensions["B"].width = 18
for c in ws[1]:
    c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HEAD

sheet("Checker list (38,026)", ["Case id", "Title", "Petitioner", "Respondent", "Decision date", "Citation", "Source"], [list(r) for r in LIST], [18, 80, 40, 40, 13, 26, 10])
sheet("Corpus judgments (19,255)", ["Doc id", "Name (names table)", "Decision date", "Court", "Citations", "Status in the checker's list", "Closest list title"], rows, [14, 60, 12, 22, 40, 70, 70])
sheet("Library (68)", ["Case (library)", "Citation held", "Checker", "Court (reading)", "Reason", "Chunk file"], lib, [50, 50, 24, 26, 90, 50])
wb.save(out)
print(out, dict(cnt), "library", len(lib), len(lib_miss), lib_hc, lib_sc_inlist, lib_sc_out)
