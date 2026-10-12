#!/usr/bin/env python3
"""Appends 'Machine Spec v2' and 'v2 Change Log' to the 2026-10-11 appended workbook (original five sheets untouched). Patterns come from review_tools/machine_spec_v2.py. PRIVATE."""
import hashlib, json, os, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "review_tools"))
import machine_spec_v2 as S
SRC = os.path.join(ROOT, "review_data/ruler/inputs/2026-10-11_trigger_workbook_with_claude_appends.xlsx")
OUT = os.path.join(ROOT, "review_data/ruler/inputs/2026-10-12_trigger_workbook_with_spec_v2.xlsx")
JSON_OUT = os.path.join(ROOT, "review_data/ruler/machine_spec_v2.json")
json.dump(S.export(), open(JSON_OUT, "w", encoding="utf-8"), indent=1)
sha = hashlib.sha256(open(JSON_OUT, "rb").read()).hexdigest()
wb = openpyxl.load_workbook(SRC)
before = [(w.title, w.dimensions) for w in wb.worksheets[:5]]
wb["Machine Spec (draft)"]["A1"].value = "SUPERSEDED on 2026-10-12 by the sheet 'Machine Spec v2'. " + str(wb["Machine Spec (draft)"]["A1"].value)
HDR = PatternFill("solid", fgColor="DDE6F0")
def sheet(title, intro, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append([intro]); ws["A1"].font = Font(bold=True); ws.append([]); ws.append(header)
    for c in ws[3]:
        c.font = Font(bold=True); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=4):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A4"
    return ws
join = lambda xs: "\n".join(f"{i + 1}. {x}" for i, x in enumerate(xs))
rows = []
for a in S.ACTS:
    s = S.SPEC[a]
    rows.append([s["name"], join(s["N1"]), join(s["N2_extra"]), s.get("N2_exclude") or "(none)", join(s["passage"]), "3 or more distinct passages (by text hash)" + ("; NOT labelled if 3+ passages mention Section 66A (struck down 2015)" if a == "IT Act" else "")])
ws = sheet("Machine Spec v2", f"MACHINE SPEC v2 (DRAFT 2026-10-12; not frozen). Source of truth: review_tools/machine_spec_v2.py; JSON export review_data/ruler/machine_spec_v2.json, sha256 {sha[:16]}... Flags: case-insensitive, DOTALL; acronyms case-sensitive via (?-i:...); near(a,b,n) = a within n characters of b in either order (written out in full below). N1 = precise question anchors (Mr. Advocate's, repaired). N2 = N1 or the plain-language anchors, unless the Act's exclusion pattern matches the question. Passage markers are shared by N1 and N2.",
           ["Act", "N1 question anchors (any)", "N2 extra plain-language anchors (any)", "N2 exclusion (if it matches, the plain-language anchors do not fire)", "Passage-side strong markers (any)", "Judgment label rule"], rows, [34, 70, 70, 36, 70, 36])
log = [
    ["PWDVA passage-side pattern", "Mr. Advocate (Python reports 'missing )': an unclosed group)", "Closed the group; added the bare 'Domestic Violence Act' (105 passages, more than the full title's 60) and sections 17, 21, 31; kept Protection Officer and aggrieved woman."],
    ["N1 'shared household', 'residence order', 'protection order', 'monetary relief' standalone", "Mr. Advocate's own rule (no standalone generic words) versus his own pattern", "Now need a domestic-violence context (not merely husband/wife) within 120 characters."],
    ["NDPS passage 'conscious possession' and 'commercial quantity' standalone", "Code test: 15 of 61 judgments carrying them never name the NDPS Act", "Now need NDPS or a drug word or section 20/21/37 within 200 characters. Labelled NDPS judgments with no passage naming the Act: 3 of 241."],
    ["NDPS supplement (recovery/seizure, no recovery)", "Mr. Advocate's supplement", "Added to N1 with NDPS/contraband context. 'Kingpin', 'cartel' and 'drug network' not added as anchors (not statutory terms; ambiguous); a looser 'drug ... trafficking/peddling' phrase is in N2 only."],
    ["POCSO passage 'special court' near child/minor; 'sexual assault' near child", "Code test and sample reading (ordinary rape cases and unrelated special courts matched)", "Removed. Markers: the Act title, 'penetrative sexual assault', and sections of the POCSO Act."],
    ["SC/ST title", "Code test: 'the Scheduled Tribes' (with 'the') is the official form in 67 of 182 title passages", "Title allows an optional 'the'; SC/ST acronym case-sensitive; bare 'SC/ST' no longer a marker by itself (it matched 'SC & ST communities' outside the Act)."],
    ["NI Act N1 'cheque bounce'", "Mr. Advocate's pattern missed 'bounced' and either word order", "Cheque near bounce/bounced/dishonour/dishonoured, in either order, within 60 characters."],
    ["PMLA title", "Code test: hyphenated 'Money-Laundering' missed (28 of 116 passages)", "Optional hyphen. Added 'money-laundering' near proceeds of crime / scheduled offence as a passage marker (real PMLA language). 'ED' only in N2 and case-sensitive."],
    ["IT Act bare section numbers", "Mr. Advocate's own rule versus his pattern: 'Section 67' fired the IT Act on an NDPS section 67 question", "Section numbers need IT context within 100 characters; 'IT Act' case-sensitive; judgments with 3+ Section 66A passages are not labelled."],
    ["All patterns", "Code test: patterns were one-directional (term order mattered)", "near(a,b,n) works in either order; DOTALL so line breaks do not break proximity."],
    ["Label rule", "Mr. Advocate's protocol", "Distinct passages by text hash, not repeated matches or duplicated headers; 3 or more."],
    ["Plain-language tier (N2)", "Owner's decision 2026-10-12: N1 precise, N2 broader", "Anchors from the owner's layperson sheet with guards: PWDVA lay anchors off when 498A/section 85 BNS/304B/dowry death appears; NDPS off with forensic/medical words; POCSO off with adult ages, adult, boss, employer, colleague, workplace; JJ off with sexual words."],
    ["Not adopted", "Mr. Advocate's supplement terms 'demand', 'acceptance', 'intermediary', 'enquiry', 'official duty', 'sanction' alone", "As he says, only with context: they appear only inside context patterns. Section 17A and 19 with the PC Act name, previous sanction/approval with the PC Act name, known sources of income with disproportionate assets are in N1."],
    ["Tests", "Written by Claude before the first run", "59 question cases (N1 and N2 each), 26 passage cases, 2 label tests: all pass after fixing two patterns and correcting three of my own wrong expectations. Not a lawyer's test."],
    ["Results on the index", "238,762 passages", "Marked judgments (3+ passages): PWDVA 21, NDPS 241, POCSO 58, SC/ST 50, NI Act 199, PMLA 50, UAPA 45, DPA 64, JJ Act 93, PC Act 407, Arms Act 255, IT Act 16, MV Act 16. Thirty-nine passages read by hand were genuine Act mentions; one was a list of laws (hence the 3-passage threshold)."],
    ["Results on the 30 used questions (dev check only)", "Used up", "N1 fires on 7 of 30, N2 on 11 of 30; the same 11 questions I would call Act-concerning; no firing looked wrong to me. The patterns were shaped after seeing these questions, so this is not evidence."],
    ["Open for Mr. Advocate", "", "(1) Are the section lists right (e.g. PWDVA 12,17-23,31; NDPS 20-22,27A,29,37,42,43,50,52A,57,67; UAPA 10,13,15-18,20,38-40,43D,43E)? (2) Any exclusion or anchor that would misfire in a real client's words? (3) Is a bare 'Domestic Violence Act' and 'Juvenile Justice Act' (including the 1986 and 2000 Acts) acceptable as markers?"],
]
sheet("v2 Change Log", "WHAT CHANGED FROM MR. ADVOCATE'S AUDIT AND FROM THE 2026-10-11 DRAFT, AND WHY (2026-10-12).", ["Item", "Source of the problem", "Action in v2"], log, [44, 56, 110])
assert [(w.title, w.dimensions) for w in wb.worksheets[:5]] == before
wb.save(OUT)
print("saved", OUT, "| sheets", wb.sheetnames, "| json sha256", sha)
